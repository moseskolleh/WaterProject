"""The chance of a working borehole (PLAN.md step 3.3).

The worked examples are computed by hand in the comments beside them, so a
reader can check the arithmetic on paper: odds are p / (1 - p), each factor
multiplies them, and the chance is odds / (1 + odds).
"""

from __future__ import annotations

import copy
import math
from types import SimpleNamespace

import pytest

from groundwater.config import Config
from groundwater.siting import odds as odds_module
from groundwater.siting.odds import (
    beta_quantile,
    lognormal_share_above,
    odds_headline,
    odds_point_text,
    odds_rows,
    odds_tables,
    odds_text,
    point_latlon,
    prior_for,
    programme_rate,
    regularised_beta,
    success_odds,
)
from groundwater.ves.model_range import Band


def _interp(rho=80.0, err=4.0, zones=((5.0, 20.0),), sid="P1"):
    """A point whose water zone, 5 to 20 m, is one layer of ``rho``."""
    layers = [SimpleNamespace(top_m=0.0, bottom_m=5.0, rho=300.0, water_bearing=False),
              SimpleNamespace(top_m=5.0, bottom_m=20.0, rho=rho, water_bearing=True),
              SimpleNamespace(top_m=20.0, bottom_m=math.inf, rho=5000.0,
                              water_bearing=False)]
    return SimpleNamespace(sounding_id=sid, water_zones=list(zones), layers=layers,
                           fit_error_percent=err, site_easting=None, site_northing=None)


def _range(p50=22.0, unresolved=0.05):
    basement = None if p50 is None else Band(p50 - 4.0, p50, p50 + 8.0)
    return SimpleNamespace(basement_m=basement, basement_unresolved=unresolved)


@pytest.fixture
def neutral(monkeypatch):
    """The real tables with every ratio set to 1 and every fit weight to 1,
    so that one piece of evidence at a time can be set and checked alone."""
    tables = copy.deepcopy(odds_tables())
    evidence = tables["evidence"]
    for band in evidence["regolith"]:
        band["lr"] = 1.0
    for key in ("resolved", "partly", "unresolved"):
        evidence["basement"][key]["lr"] = 1.0
    evidence["resistivity"]["none"]["lr"] = 1.0
    for band in evidence["resistivity"]["bands"]:
        band["lr"] = 1.0
    for band in evidence["fit"]:
        band["weight"] = 1.0
    monkeypatch.setattr(odds_module, "odds_tables", lambda: tables)
    return evidence


def _factors(o):
    return {e.key: e.factor for e in o.evidence}


# ------------------------------------------------------------------ the maths

def test_the_prior_for_basement_is_worked_out_from_the_bgs_quartiles():
    # B-L: quartiles 0.1 and 0.5 L/s; success is 1 m3/h = 0.27778 L/s.
    # median sqrt(0.1 x 0.5) = 0.22361, ln = -1.49787
    # sigma = ln(5) / (2 x 0.67449) = 1.19305
    # z = (ln 0.27778 + 1.49787) / 1.19305 = (-1.28093 + 1.49787) / 1.19305 = 0.18184
    # 1 - Phi(0.18184) = 0.42785
    assert lognormal_share_above(0.1, 0.5, 1 / 3.6) == pytest.approx(0.42785, abs=1e-5)
    assert prior_for("B-L", "pCm", 1.0)["rate"] == pytest.approx(0.42785, abs=1e-5)
    # a threshold at a quartile leaves a quarter above the upper one
    assert lognormal_share_above(0.1, 0.5, 0.5) == pytest.approx(0.25, abs=2e-7)
    assert lognormal_share_above(0.1, 0.5, 0.1) == pytest.approx(0.75, abs=2e-7)


def test_the_beta_functions_against_known_values():
    # Beta(1, 1) is uniform: its distribution function is x, its quantile q
    for x in (0.1, 0.37, 0.9):
        assert regularised_beta(1.0, 1.0, x) == pytest.approx(x, abs=1e-12)
        assert beta_quantile(1.0, 1.0, x) == pytest.approx(x, abs=1e-12)
    # Beta(2, 1) has I_x = x^2, so its 0.25 quantile is 0.5
    assert regularised_beta(2.0, 1.0, 0.3) == pytest.approx(0.09, abs=1e-12)
    assert beta_quantile(2.0, 1.0, 0.25) == pytest.approx(0.5, abs=1e-12)
    # Beta(1, 3): I_x = 1 - (1 - x)^3, so its median is 1 - 0.5^(1/3)
    assert beta_quantile(1.0, 3.0, 0.5) == pytest.approx(1 - 0.5 ** (1 / 3), abs=1e-12)
    # symmetric: the median of Beta(a, a) is a half
    assert beta_quantile(4.0, 4.0, 0.5) == pytest.approx(0.5, abs=1e-12)


# ----------------------------------------------------- the worked examples

def test_the_prior_alone_with_no_evidence(neutral):
    # No position: the fallback, an even chance held as two boreholes, which
    # is Beta(1, 1), uniform. With every factor 1 the answer is the prior:
    # 50 percent, between its 10th and 90th percentiles, 10 and 90.
    o = success_odds(_interp(), None, None, None)
    assert o.matched == "fallback" and o.prior == 0.5
    assert o.prior_odds == pytest.approx(1.0)
    assert o.posterior_odds == pytest.approx(1.0)
    assert o.probability == pytest.approx(0.5)
    assert (o.low, o.high) == (pytest.approx(0.1), pytest.approx(0.9))
    assert odds_headline(o) == (
        "About 50 percent (between 10 and 90) that a borehole here yields enough "
        "for a handpump through the dry season. Placeholder prior, with no cited "
        "rate behind it.")


def test_the_regolith_ratio_alone(neutral):
    # favourable (15 to 35 m) at 1.5: odds 1 x 1.5 = 1.5, chance 1.5 / 2.5 = 0.6
    neutral["regolith"][2]["lr"] = 1.5
    o = success_odds(_interp(), _range(p50=22.0), None, None)
    assert _factors(o) == {"regolith": 1.5, "basement": 1.0, "resistivity": 1.0,
                           "fit": 1.0}
    assert o.probability == pytest.approx(0.6)
    # the band moves with it: 0.1 x 1.5 / (0.1 x 1.5 + 0.9) = 0.15 / 1.05
    assert o.low == pytest.approx(0.15 / 1.05)
    assert o.high == pytest.approx(1.35 / 1.45)


def test_the_basement_ratio_alone(neutral):
    # partly resolved (a quarter of the models find none) at 0.9:
    # odds 0.9, chance 0.9 / 1.9 = 0.47368
    neutral["basement"]["partly"]["lr"] = 0.9
    o = success_odds(_interp(), _range(unresolved=0.25), None, None)
    assert o.evidence[1].band == "partly"
    assert o.probability == pytest.approx(0.9 / 1.9)


def test_the_resistivity_ratio_alone(neutral):
    # 80 ohm-m is in the productive band, 50 to 300, at 1.4:
    # odds 1.4, chance 1.4 / 2.4 = 0.58333
    neutral["resistivity"]["bands"][2]["lr"] = 1.4
    o = success_odds(_interp(rho=80.0), None, None, None)
    assert o.evidence[2].band == "productive"
    assert o.evidence[2].value == pytest.approx(80.0)
    assert o.probability == pytest.approx(1.4 / 2.4)


def test_the_fit_alone_moves_nothing(neutral):
    # The fit weights the other evidence; with none to weight, 1^w = 1.
    neutral["fit"][3]["weight"] = 0.3
    o = success_odds(_interp(err=27.0), None, None, None)
    assert o.evidence[3].band == "unreliable" and o.evidence[3].weight == 0.3
    assert o.evidence[3].factor == pytest.approx(1.0)
    assert o.probability == pytest.approx(0.5)


def test_every_piece_together_with_the_real_ratios():
    # B-L on pCm, prior 0.42786, odds 0.42786 / 0.57214 = 0.74782.
    # regolith 22 m, favourable:            x 1.5
    # basement found in 95 percent:         x 1.0  (resolved)
    # 80 ohm-m, productive:                 x 1.4
    # fit 12 percent, poor, weight 0.6:     x (1.5 x 1.0 x 1.4)^(0.6 - 1)
    #                                         = 2.1^-0.4 = 0.74321
    # posterior odds 0.74782 x 2.1^0.6 = 0.74782 x 1.56074 = 1.16716
    # chance 1.16716 / 2.16716 = 0.53857
    o = success_odds(_interp(rho=80.0, err=12.0), _range(p50=22.0, unresolved=0.05),
                     "B-L", "pCm")
    assert [e.band for e in o.evidence] == ["favourable", "resolved", "productive", "poor"]
    assert _factors(o)["fit"] == pytest.approx(2.1 ** -0.4)
    assert o.prior_odds == pytest.approx(0.74782, abs=1e-5)
    assert o.posterior_odds == pytest.approx(o.prior_odds * 2.1 ** 0.6)
    assert o.probability == pytest.approx(0.53857, abs=1e-5)
    rows = odds_rows(o)
    assert rows[0] == ["Prior", "the BGS class B-L on the USGS unit pCm", "", "43"]
    assert rows[1] == ["Depth to basement (P50)", "22 m: favourable", "1.50", "53"]
    assert rows[4][2] == "0.74" and rows[4][3] == "54"
    assert odds_text(o)[0].startswith("About 54 percent (between ")
    assert odds_text(o)[0].endswith("through the dry season. Provisional prior.")


def test_the_breakdown_multiplies_back_to_the_posterior():
    """Over every band of every class, the prior odds times the factors in
    the breakdown are the posterior odds, and the last step is the answer."""
    ranges = [None, _range(3.0, 0.0), _range(10.0, 0.05), _range(22.0, 0.1),
              _range(50.0, 0.5), _range(50.0, 0.95), _range(None, 1.0)]
    interps = [_interp(rho=r, err=e) for r in (10.0, 35.0, 80.0, 500.0, 2000.0)
               for e in (2.0, 7.0, 15.0, 30.0, None)]
    interps.append(_interp(zones=()))
    grounds = [("B-L", "pCm"), ("U-M/H", "Qe"), ("CSF-L/M", None), (None, None)]
    seen = set()
    for interp in interps:
        for r in ranges:
            for code, glg in grounds:
                o = success_odds(interp, r, code, glg)
                product = math.prod(e.factor for e in o.evidence)
                assert o.prior_odds * product == pytest.approx(o.posterior_odds, rel=1e-12)
                assert o.evidence[-1].after == pytest.approx(o.probability, rel=1e-12)
                assert o.low <= o.probability <= o.high
                seen |= {(e.key, e.band) for e in o.evidence}
    # every band of the evidence file was reached
    table = odds_tables()["evidence"]
    wanted = {("regolith", b["key"]) for b in table["regolith"]}
    wanted |= {("basement", k) for k in ("resolved", "partly", "unresolved")}
    wanted |= {("resistivity", b["key"]) for b in table["resistivity"]["bands"]}
    wanted |= {("resistivity", "none"), ("fit", "unknown")}
    wanted |= {("fit", b["key"]) for b in table["fit"]}
    assert wanted <= seen, sorted(wanted - seen)


# ------------------------------------------------------------- edge cases

def test_an_unknown_geology_unit_takes_the_class_row():
    o = success_odds(_interp(), None, "B-L", "XX")
    assert o.matched == "class"
    assert o.prior == pytest.approx(prior_for("B-L", "pCm", 1.0)["rate"])
    assert "which the prior table does not list on its own" in odds_text(o)[2]
    no_unit = success_odds(_interp(), None, "B-L", None)
    assert no_unit.matched == "class"
    assert odds_rows(no_unit)[0][1] == "the BGS class B-L, with no USGS unit under the point"


@pytest.mark.parametrize("code,glg", [(None, None), (None, "pCm"), ("n/a", "H2O"),
                                      ("ZZ", "pCm")])
def test_ground_with_no_cited_rate_takes_the_labelled_fallback(code, glg):
    o = success_odds(_interp(), None, code, glg)
    assert o.matched == "fallback" and o.status == "fallback"
    assert o.prior == 0.5 and o.effective_n == 2.0
    assert odds_headline(o).endswith("Placeholder prior, with no cited rate behind it.")
    assert "a placeholder held as weakly as 2 boreholes" in odds_text(o)[2]


def test_a_range_not_yet_sampled_leaves_its_evidence_out():
    o = success_odds(_interp(), None, "B-L", "pCm")
    assert [(e.band, e.factor) for e in o.evidence[:2]] == [("not_sampled", 1.0),
                                                           ("not_sampled", 1.0)]
    assert not o.range_sampled
    assert any(line.startswith("The range of models has not been sampled at P1")
               for line in odds_text(o))
    assert odds_rows(o)[1][1] == "range not sampled"


def test_a_basement_band_not_quoted_gives_no_depth_and_counts_as_unresolved():
    # basement in 5 percent of the models: under the tenth that quotes a band
    o = success_odds(_interp(), _range(p50=45.0, unresolved=0.95), "B-L", "pCm")
    regolith, basement = o.evidence[0], o.evidence[1]
    assert (regolith.band, regolith.value, regolith.factor) == ("not_quoted", None, 1.0)
    assert (basement.band, basement.factor) == ("unresolved", 0.8)
    assert odds_rows(o)[1][1] == "no depth quoted, too few models find basement"
    # and with no basement in any model at all
    never = success_odds(_interp(), _range(p50=None, unresolved=1.0), "B-L", "pCm")
    assert [e.band for e in never.evidence[:2]] == ["not_quoted", "unresolved"]
    # exactly a tenth resolved, as 400 of 4,000 store it, is quoted
    tenth = success_odds(_interp(), _range(p50=45.0, unresolved=1.0 - 400 / 4000),
                         "B-L", "pCm")
    assert tenth.evidence[0].band == "deep" and tenth.evidence[1].band == "partly"


def test_the_water_zone_resistivity_leaves_out_what_rounding_takes_in():
    """A zone is rounded to whole metres. 200 ohm-m over 5,000 ohm-m
    basement, 2.4 or 2.6 m thick, is the zone 6 to 8 m or 6 to 9 m; the
    second took in 0.4 m of basement, read 307 ohm-m, "resistive", and
    dropped the chance from 61 to 47 percent."""
    from groundwater.models import LayeredModel
    from groundwater.ves import interpret_model

    seen = []
    for thickness in (2.4, 2.6):
        model = LayeredModel([300.0, 200.0, 5000.0], [6.0, thickness])
        model.fit_error_percent = 3.0
        interp = interpret_model(None, model)
        o = success_odds(interp, _range(p50=20.0, unresolved=0.0), "B-L", "pCm")
        seen.append((interp.water_zones, o.evidence[2].band, o.evidence[2].value,
                     round(o.probability, 12)))
    assert [z for z, *_ in seen] == [[(6, 8)], [(6, 9)]]
    assert [rest for _z, *rest in seen] == [["productive", 200.0, seen[0][3]]] * 2
    # and a dry layer above a zone is left out too: the water-bearing layer
    # from 8.3 m is the zone 8 to 40 m, which took in 0.3 m of 1,200 ohm-m
    model = LayeredModel([1200.0, 49.97], [8.3])
    model.fit_error_percent = 3.0
    interp = interpret_model(None, model)
    assert interp.water_zones[0][0] == 8
    o = success_odds(interp, None, "U-M/H", "Qe")
    assert (o.evidence[2].band, o.evidence[2].value) == ("clayey", 49.97)


@pytest.mark.parametrize("p50,band,under", [(5.0, "shallow", "thin"),
                                            (15.0, "favourable", "shallow"),
                                            (35.0, "deep", "favourable")])
def test_a_depth_on_a_band_edge_takes_the_band_above_it(p50, band, under):
    # the edges are each band's "below": 5 m is not under 5 m
    def depth(d):
        return success_odds(_interp(), _range(p50=d), None, None).evidence[0].band

    assert (depth(p50), depth(p50 - 1e-9)) == (band, under)


def test_a_basement_share_on_its_edge_counts_as_resolved():
    # "resolved" is at most a tenth of the models finding none
    edge = success_odds(_interp(), _range(unresolved=0.1), None, None)
    over = success_odds(_interp(), _range(unresolved=0.1 + 1e-12), None, None)
    assert (edge.evidence[1].band, over.evidence[1].band) == ("resolved", "partly")


@pytest.mark.parametrize("err,band,under", [(5.0, "acceptable", "excellent"),
                                            (10.0, "poor", "acceptable"),
                                            (20.0, "unreliable", "poor")])
def test_a_misfit_on_a_band_edge_takes_the_band_above_it(err, band, under):
    def fit(e):
        return success_odds(_interp(err=e), None, None, None).evidence[3].band

    assert (fit(err), fit(err - 1e-9)) == (band, under)


@pytest.mark.parametrize("rho,band,under", [(20.0, "clayey", "clay"),
                                            (50.0, "productive", "clayey"),
                                            (300.0, "resistive", "productive"),
                                            (800.0, "fresh", "resistive")])
def test_a_resistivity_on_a_band_edge_takes_the_band_above_it(rho, band, under):
    """A zone of one layer reads that layer's resistivity exactly: the
    geometric mean as exp(mean ln) read 50 ohm-m as 49.99999999999999, and
    so "clayey", where the evidence file says 50 to 300 is productive."""
    def found(r):
        return success_odds(_interp(rho=r), None, None, None).evidence[2]

    assert (found(rho).value, found(rho).band) == (rho, band)
    assert found(rho * (1 - 1e-9)).band == under


def test_no_water_zone_and_no_misfit():
    o = success_odds(_interp(zones=(), err=None), None, "B-L", "pCm")
    assert (o.evidence[2].band, o.evidence[2].factor) == ("none", 0.3)
    assert (o.evidence[3].band, o.evidence[3].weight) == ("unknown", 1.0)
    assert odds_rows(o)[3][1] == "no water-bearing zone in reach"


# ------------------------------------------------- configuration and place

def test_success_is_defined_in_the_configuration():
    assert Config().odds.success_yield_m3_per_h == 1.0
    stricter = Config()
    stricter.odds.success_yield_m3_per_h = 3.6  # 1 L/s
    base = success_odds(_interp(), None, "B-L", "pCm")
    strict = success_odds(_interp(), None, "B-L", "pCm", stricter)
    # 1 L/s is above the class's upper quartile, so under a quarter clear it
    assert strict.prior < 0.25 < base.prior
    assert "at least 3.6 m3/h (1.00 L/s)" in odds_text(strict)[1]


def test_a_low_success_yield_keeps_the_band_finite(tmp_path):
    """On the coastal sands a low success yield puts the prior rate near 1,
    and the 90th percentile of its Beta bisected to exactly 1.0: carried
    through the odds that was a ZeroDivisionError in Python and "nan" in the
    browser, from 0.33 m3/h down. The band is held inside (0, 1) as the rate
    is, so it reads "over 99", not 100 and not nan."""
    path = tmp_path / "config.yaml"
    path.write_text("odds:\n  success_yield_m3_per_h: 0.3\n", encoding="utf-8")
    config = Config.load(path)
    # a point at the Rokel survey's position, on the coastal sands
    here = _interp()
    here.site_easting, here.site_northing = 708958.0, 926355.0
    from groundwater.siting import survey_odds

    (o,) = survey_odds([here], [_range()], 28, None, config)
    assert (o.bgs_code, o.glg) == ("U-M/H", "Qe")
    assert 0.0 < o.prior_low <= o.prior_high < 1.0
    assert all(math.isfinite(v) and 0.0 < v < 1.0 for v in (o.low, o.probability, o.high))
    assert "over 99" in odds_headline(o)
    assert "nan" not in odds_headline(o) and " 100" not in odds_headline(o)
    # over the range a project might set, on every cited class, both ends
    for rate in [k / 100 for k in range(1, 100)] + [3.6, 36.0, 360.0]:
        config.odds.success_yield_m3_per_h = rate
        for code, glg in (("U-M/H", "Qe"), ("B-L", "pCm"), ("I-L", "Pi"),
                          ("CSF-L/M", "pCm")):
            o = success_odds(_interp(), _range(), code, glg, config)
            assert all(0.0 < v < 1.0 for v in (o.low, o.probability, o.high)), (rate, code)


@pytest.mark.parametrize("rate", [0.0, -1.0, float("nan")])
def test_a_success_yield_of_nothing_is_refused(rate):
    with pytest.raises(ValueError, match="above zero"):
        prior_for("B-L", "pCm", rate)


def test_the_config_file_overrides_success(tmp_path):
    path = tmp_path / "config.yaml"
    path.write_text("odds:\n  success_yield_m3_per_h: 2\n", encoding="utf-8")
    assert Config.load(path).odds.success_yield_m3_per_h == 2.0


def test_a_point_is_placed_by_its_own_coordinates_else_the_site():
    here = _interp()
    assert point_latlon(here, None, (8.4, -13.2)) == (8.4, -13.2)
    assert point_latlon(here, None, None) is None
    here.site_easting, here.site_northing = 708958.0, 926355.0
    lat, lon = point_latlon(here, 28, (1.0, 1.0))
    assert lat == pytest.approx(8.3759, abs=1e-4) and lon == pytest.approx(-13.1024, abs=1e-4)


def test_the_programme_rate_never_plans_certainty():
    # 99.8 percent is offered as 99, and 0.2 as 1, the least the estimate takes
    assert programme_rate(SimpleNamespace(probability=0.998)) == 99.0
    assert programme_rate(SimpleNamespace(probability=0.002)) == 1.0
    # in between, the central figure to a whole percent, halves to even
    assert programme_rate(SimpleNamespace(probability=0.425)) == 42.0
    low = success_odds(_interp(zones=()), _range(p50=3.0, unresolved=0.0), "B-L", "pCm")
    assert 1.0 <= programme_rate(low) == float(round(100 * low.probability))


def test_every_band_of_the_evidence_file_has_words():
    from groundwater.text import phrase_table

    table = odds_tables()["evidence"]
    words = phrase_table("odds.bands")
    keys = [b["key"] for b in table["regolith"]] + ["resolved", "partly", "unresolved"]
    keys += [b["key"] for b in table["resistivity"]["bands"]] + [b["key"] for b in table["fit"]]
    assert sorted(keys) == sorted(words)


def test_every_prior_row_is_cited_or_labelled_a_fallback():
    for row in odds_tables()["prior"]:
        if row["status"] == "provisional":
            assert row["yield_q1_l_per_s"] < row["yield_q3_l_per_s"]
            assert "O Dochartaigh 2021" in row["basis"]
        else:
            assert row["status"] == "fallback" and row["rate"] == 0.5
            assert "no cited rate" in row["basis"]
    # the fallback for a point with no class at all is there
    assert any(r["bgs_code"] == "*" and r["glg"] == "*" for r in odds_tables()["prior"])


@pytest.mark.slow
def test_the_report_prints_the_odds_under_the_score(tmp_path):
    from pathlib import Path

    from docx import Document

    from groundwater.ingestion.ves import read_ves_workbook
    from groundwater.reporting import build_geophysical_report
    from groundwater.reporting.geophysical import GeophysicalReportInputs
    from groundwater.siting import survey_odds
    from groundwater.ves import interpret_model
    from groundwater.ves.inversion import invert_sounding

    soundings = read_ves_workbook(Path(__file__).resolve().parents[1] / "examples"
                                  / "data" / "rokel" / "rokel_ves.xlsx")
    inversions = [invert_sounding(s) for s in soundings]
    interps = [interpret_model(s, r.model) for s, r in zip(soundings, inversions, strict=True)]
    path = build_geophysical_report(
        GeophysicalReportInputs(soundings=soundings, inversions=inversions,
                                interpretations=interps, figures_dir=tmp_path),
        tmp_path / "odds.docx")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.index("Chance of a working borehole") > text.index("Drill-target suitability")
    # what success means and the basis, once for the survey
    assert text.count("Success means a yield of at least 1 m3/h") == 1
    assert text.count("Basis: the prior from success_prior.csv") == 1
    site = soundings[0].site
    for o in survey_odds(interps, None, site.utm_zone, site.latlon):
        # each point placed on the map under it: Rokel is on the coastal sands
        assert (o.bgs_code, o.glg) == ("U-M/H", "Qe")
        assert " ".join(odds_point_text(o)) in text
        assert f"How the survey moved the chance of a working borehole at point {o.sounding_id}" in text
    cells = [[c.text for c in row.cells] for t in doc.tables for row in t.rows]
    assert ["Evidence", "What the survey found", "Factor on the odds",
            "Chance after (percent)"] in cells


def test_the_report_pairs_each_point_with_its_own_odds_in_rank_order():
    """Two points with one id (a sheet copied without renumbering): the
    report ordered the odds by looking each id up in the ranking, so both
    took the last row's place and printed in the order listed, the weaker
    point first."""
    from groundwater.models import LayeredModel
    from groundwater.reporting.geophysical import _odds_block
    from groundwater.siting import assess_siting, survey_odds
    from groundwater.ves import interpret_model

    points = []
    for thickness in (3.0, 25.0):
        model = LayeredModel([300.0, 120.0, 5000.0], [6.0, thickness])
        model.fit_error_percent = 3.0
        interp = interpret_model(None, model)
        interp.sounding_id = "VES 1"
        points.append(interp)
    # the point listed second is the better, scored a point at a time
    assert assess_siting(points[1:])[0].weighted > assess_siting(points[:1])[0].weighted
    suit = assess_siting(points)
    odds = survey_odds(points, [None, _range(p50=31.0)], None, None)
    assert odds[0].probability != odds[1].probability

    class Writer:
        def __init__(self):
            self.paragraphs = []

        def heading(self, *args, **kwargs):
            pass

        def table(self, *args, **kwargs):
            pass

        def paragraph(self, text, **kwargs):
            self.paragraphs.append(text)

    rb = Writer()
    _odds_block(rb, suit, odds)
    assert rb.paragraphs[1:] == [" ".join(["Point VES 1."] + odds_point_text(o))
                                 for o in (odds[1], odds[0])]


def test_the_ground_is_what_the_maps_say():
    from groundwater.mapping import aquifer_unit_at, geology_unit_at
    from groundwater.siting.odds import ground_at

    for lat, lon in ((8.3759, -13.1024), (8.6, -11.5), (8.47, -13.24), (9.5, -12.0),
                     (5.0, -20.0)):
        aquifer, unit = aquifer_unit_at(lat, lon), geology_unit_at(lat, lon)
        assert ground_at((lat, lon)) == (aquifer.glg if aquifer else None,
                                         unit.glg if unit else None)
    assert ground_at(None) == (None, None)


def test_a_survey_with_fewer_ranges_than_points_reads_the_rest_unsampled():
    from groundwater.siting import survey_odds

    points = [_interp(sid="A"), _interp(sid="B")]
    out = survey_odds(points, [_range()], None, None)
    assert [o.range_sampled for o in out] == [True, False]
    assert [o.matched for o in out] == ["fallback", "fallback"]
