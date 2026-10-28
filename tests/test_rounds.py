"""Service rounds (PLAN.md step 7.3): what is due at each borehole of the
register on a day, grouped by area mechanic, put in an order, and printed on
a round sheet.

The registry already said, borehole by borehole, when the sanitary
inspection and the water sample were due. It now schedules the routine
inspection of the pump and rising main too, at the interval the design sets
for the rising main (step 7.2), with every interval read from
data/defaults.json; and the rounds group what is due by area mechanic and
order each round by a stated nearest-neighbour walk. The day is an input,
never the clock, and the browser's engine is held to this one on seeded
registers.
"""

from __future__ import annotations

import json
import random
import shutil
import subprocess
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest

from groundwater import rounds as R
from groundwater.config import Config, DesignRules, ServiceConfig
from groundwater.design.materials import CHEMISTRY_STATES
from groundwater.models import SiteMetadata
from groundwater.registry import (
    Asset,
    AssetEvent,
    Schedule,
    asset_from_dict,
    asset_from_project,
    asset_state,
    asset_with_rising_main,
    clean_rising_main,
    mint_asset_id,
    registry_stats,
)
from groundwater.sanitary import form_questions
from groundwater.text import phrase, phrase_table

REPO = Path(__file__).resolve().parents[1]
JS = REPO / "docs" / "js"
DAY = date(2024, 6, 1)

AGGRESSIVE = {"material": "stainless steel", "inspection_months": 6, "chemistry": "aggressive"}
PLAIN = {"material": "galvanised iron", "inspection_months": 12, "chemistry": "not_aggressive"}


def _id(easting: float, northing: float = 938150.0, zone: int = 28,
        district: str = "Western Area Rural") -> str:
    return mint_asset_id(SiteMetadata(district=district, easting=easting,
                                      northing=northing, utm_zone=zone))


def _asset(*events, easting=694912.0, northing=938150.0, **fields) -> Asset:
    return Asset(asset_id=fields.pop("asset_id", None) or _id(easting, northing),
                 community=fields.pop("community", "Dr. Timbo's"),
                 district="Western Area Rural", easting=easting, northing=northing,
                 utm_zone=28, events=list(events), **fields)


def _routine(state):
    return next(item for item in state.due if item.key == "routine_inspection")


# ---------------------------------------------------------------- the schedule

def test_the_schedule_is_the_configuration():
    """Every interval comes from defaults.json: the service section, and the
    design rules' interval for a rising main in water nobody has analysed."""
    plan = Schedule()
    assert plan == Schedule.from_config(Config())
    service = ServiceConfig()
    assert (plan.sanitary_inspection_months, plan.water_sample_months, plan.grace_days) == (
        service.sanitary_inspection_months, service.water_sample_months, service.grace_days)
    assert plan.routine_inspection_months == DesignRules().aggressive_rising_main_inspection_months
    config = Config()
    config.service.water_sample_months = 3
    config.design.aggressive_rising_main_inspection_months = 4
    moved = Schedule.from_config(config)
    assert (moved.water_sample_months, moved.routine_inspection_months) == (3, 4)


_BAD_SCHEDULES = [{"water_sample_months": 1.5}, {"sanitary_inspection_months": 0},
                  {"routine_inspection_months": 1201}, {"water_sample_months": "12"},
                  {"grace_days": -1}, {"grace_days": 2.5}, {"sanitary_inspection_months": True}]


@pytest.mark.parametrize("bad", _BAD_SCHEDULES)
def test_a_schedule_no_due_date_can_be_counted_with_is_refused(bad):
    """A configuration can now set the intervals, and 1.5 months would raise a
    TypeError in Python's date arithmetic and give the browser a due date of
    "2024-02.5-NaN"; both engines refuse it instead."""
    with pytest.raises(ValueError, match="schedule: "):
        Schedule(**{**Schedule().__dict__, **bad})
    assert Schedule(water_sample_months=6.0).water_sample_months == 6


def test_the_browser_refuses_the_same_schedules(tmp_path):
    node = shutil.which("node")
    if node is None:
        pytest.skip("node is not installed; run `node --version` to check")
    driver = tmp_path / "schedules.mjs"
    driver.write_text(_DRIVER.split("const cases")[0] + r"""
const bad = JSON.parse(readFileSync(0, 'utf8'));
const asset = { asset_id: 'SL-WAR-8FEEVKQ-T', events: [] };
process.stdout.write(JSON.stringify(bad.map((s) => {
  const said = [];
  for (const run of [() => C.assetState(asset, '2024-06-01', s),
                     () => C.serviceRounds([asset], '2024-06-01', s, [])]) {
    try { run(); said.push(''); } catch (e) { said.push(String(e.message)); }
  }
  return said;
})));
""", encoding="utf-8")
    done = subprocess.run([node, str(driver), str(JS)], input=json.dumps(_BAD_SCHEDULES),
                          capture_output=True, text=True, encoding="utf-8", check=True)
    for said in json.loads(done.stdout):
        assert all(text.startswith("schedule: ") for text in said), said


def test_the_chemistry_words_cover_every_state_the_design_names():
    assert set(phrase_table("rounds.water")) == set(CHEMISTRY_STATES)


# ------------------------------------------------------- the routine inspection

def test_a_routine_inspection_is_due_at_the_interval_the_design_sets():
    asset = _asset(AssetEvent("2020-01-10", "commissioned"),
                   AssetEvent("2023-12-01", "routine_inspection", by="A. Bangura"),
                   rising_main=AGGRESSIVE)
    on_the_day = _routine(asset_state(asset, DAY))
    assert on_the_day.due_on == date(2024, 6, 1) and on_the_day.state == "scheduled"
    assert on_the_day.months == 6
    assert on_the_day.basis == phrase("rounds.basis_design", months=6,
                                      material="stainless steel",
                                      water=phrase_table("rounds.water")["aggressive"])
    assert _routine(asset_state(asset, date(2024, 6, 2))).state == "due"
    late = _routine(asset_state(asset, date(2024, 7, 2)))
    assert late.state == "overdue" and phrase("rounds.routine_why") in late.detail
    assert asset_state(asset, DAY).last_routine_inspection == date(2023, 12, 1)


def test_water_that_is_not_aggressive_is_inspected_at_the_usual_interval():
    asset = _asset(AssetEvent("2020-01-10", "commissioned"),
                   AssetEvent("2023-12-01", "routine_inspection"), rising_main=PLAIN)
    item = _routine(asset_state(asset, DAY))
    assert (item.due_on, item.state, item.months) == (date(2024, 12, 1), "scheduled", 12)
    assert "in water that is not aggressive" in item.basis


def test_a_record_with_no_rising_main_is_inspected_at_the_shorter_interval():
    item = _routine(asset_state(_asset(AssetEvent("2024-01-15", "commissioned")), DAY))
    shorter = DesignRules().aggressive_rising_main_inspection_months
    assert item.months == shorter and item.due_on == date(2024, 7, 15)
    assert item.basis == phrase("rounds.basis_default", months=shorter)
    bare = _routine(asset_state(_asset(AssetEvent("2024-01-15", "commissioned"),
                                       rising_main={"inspection_months": 9}), DAY))
    assert bare.months == 9 and bare.basis == phrase("rounds.basis_record", months=9)


def test_a_repair_or_a_sanitary_inspection_is_not_a_routine_inspection():
    """A repair records what was mended, not that the rest was looked at, and
    a sanitary inspection looks at the apron and the catchment, not the
    rising main; each visit is counted from its own records."""
    asset = _asset(AssetEvent("2020-01-10", "commissioned"),
                   AssetEvent("2024-05-01", "repair", "new seals"),
                   AssetEvent("2024-05-02", "inspection"))
    state = asset_state(asset, DAY)
    assert _routine(state).state == "overdue" and "has never been done" in _routine(state).detail
    sanitary = next(i for i in state.due if i.key == "inspection")
    assert sanitary.state == "scheduled" and sanitary.due_on == date(2024, 11, 2)
    routine_only = asset_state(_asset(AssetEvent("2020-01-10", "commissioned"),
                                      AssetEvent("2024-05-02", "routine_inspection")), DAY)
    assert next(i for i in routine_only.due if i.key == "inspection").state == "overdue"


def test_the_three_visits_are_in_one_order_and_counted():
    state = asset_state(_asset(AssetEvent("2020-01-10", "commissioned")), DAY)
    assert [i.key for i in state.due] == ["inspection", "routine_inspection", "water_sample"]
    stats = registry_stats([_asset(AssetEvent("2020-01-10", "commissioned")), _asset()], DAY)
    assert (stats["n_overdue_inspection"], stats["n_overdue_routine"],
            stats["n_overdue_sample"]) == (1, 1, 1)


def test_a_due_date_past_the_year_9999_is_never_due():
    """It used to raise in Python and count from a date the browser could
    not read; both now say it falls due after the last day there is."""
    asset = _asset(AssetEvent("9999-08-01", "commissioned"))
    item = _routine(asset_state(asset, date(9999, 12, 31)))
    assert item.state == "scheduled" and item.due_on is None
    assert item.detail == phrase("rounds.beyond", title="routine inspection")


# --------------------------------------------------------------- the record

@pytest.mark.parametrize(("raw", "kept"), [
    (AGGRESSIVE, AGGRESSIVE),
    ({"material": " uPVC ", "inspection_months": 6.0, "chemistry": "none"},
     {"material": "uPVC", "inspection_months": 6, "chemistry": "none"}),
    ({"material": 5, "inspection_months": 3, "chemistry": "acid"},
     {"material": "", "inspection_months": 3, "chemistry": ""}),
    ({"inspection_months": 6.5}, None),
    ({"inspection_months": True}, None),
    ({"inspection_months": 0}, None),
    ({"inspection_months": 1201}, None),
    ({"inspection_months": "6"}, None),
    # an int past a double, which the browser reads as Infinity
    ({"inspection_months": 10 ** 400}, None),
    ({"material": "steel"}, None),
    ("6 months", None),
    (None, None),
])
def test_a_record_keeps_a_rising_main_only_with_an_interval_it_can_use(raw, kept):
    assert clean_rising_main(raw) == kept


def test_the_record_carries_the_rising_main_of_the_design():
    site = SiteMetadata(community="X", district="Bo", easting=790500.0, northing=875300.0,
                        utm_zone=28)

    class Design:
        rising_main = dict(AGGRESSIVE, statement="...", upgraded=True)

    drafted = asset_from_project({"site": site, "borehole_design": Design()})
    assert drafted.rising_main == AGGRESSIVE
    assert asset_from_project({"site": site}).rising_main is None
    stored = asset_from_dict(drafted.as_dict())
    assert asset_with_rising_main(stored, PLAIN).rising_main == PLAIN
    # a design with no rising main leaves the record as it is
    assert asset_with_rising_main(stored, None) is stored


def test_a_record_with_no_mechanic_and_no_rising_main_is_written_as_before():
    plain = _asset(AssetEvent("2020-01-10", "commissioned")).as_dict()
    assert "area_mechanic" not in plain and "rising_main" not in plain
    named = _asset(area_mechanic="A. Bangura", rising_main=PLAIN)
    back = asset_from_dict(named.as_dict())
    assert (back.area_mechanic, back.rising_main) == ("A. Bangura", PLAIN)
    odd = asset_from_dict({**named.as_dict(), "area_mechanic": 7, "rising_main": "no"})
    assert (odd.area_mechanic, odd.rising_main) == ("", None)


# ------------------------------------------------------------ the mechanics

def test_area_mechanics_are_read_one_chiefdom_a_line():
    parsed = R.parse_area_mechanics(
        "Koya: A. Bangura\r\n\nno colon here\nKOYA: M. Kolleh\n : Fatmata\nKori:\n"
        "Bumpe Ngao:  I. Sesay: senior ")
    assert parsed["entries"] == [{"chiefdom": "Koya", "mechanic": "A. Bangura"},
                                 {"chiefdom": "Bumpe Ngao", "mechanic": "I. Sesay: senior"}]
    assert parsed["skipped"] == [
        phrase("rounds.mechanics_no_colon", line=3),
        phrase("rounds.mechanics_repeated", line=4, chiefdom="KOYA"),
        phrase("rounds.mechanics_blank", line=5),
        phrase("rounds.mechanics_blank", line=6),
    ]
    assert R.parse_area_mechanics(None) == {"entries": [], "skipped": []}


def test_the_mechanic_is_the_record_s_then_the_chiefdom_s_then_none():
    table = R.parse_area_mechanics("koya: M. Kolleh")["entries"]
    assert R.mechanic_for(_asset(chiefdom="Koya", area_mechanic=" A. Bangura "), table) == (
        "A. Bangura", "asset")
    assert R.mechanic_for(_asset(chiefdom=" KOYA "), table) == ("M. Kolleh", "chiefdom")
    assert R.mechanic_for(_asset(chiefdom="Kori"), table) == ("", "none")
    assert R.mechanic_for(_asset(chiefdom=""), table) == ("", "none")
    # only A to Z are folded, so both engines match the same names
    assert R.mechanic_for(_asset(chiefdom="KÖYA"), R.parse_area_mechanics("köya: X")["entries"]) == (
        "", "none")


# ---------------------------------------------------------------- the rounds

def test_a_borehole_with_no_record_of_a_visit_is_due_now_and_not_working():
    result = R.service_rounds([_asset()], DAY)
    stop = result["rounds"][0]["stops"][0]
    assert stop["condition"] == "Not known" and stop["function"] == "unknown"
    assert [i["state"] for i in stop["due"]] == ["unknown"] * 3
    assert stop["first_due"] == "2024-06-01"
    assert phrase_table("rounds.state")["unknown"] in R.round_rows(result["rounds"][0])[0][5]


def test_what_is_not_due_and_what_is_decommissioned_are_not_visited():
    fresh = _asset(AssetEvent("2024-05-20", "commissioned"), easting=694000.0)
    gone = _asset(AssetEvent("2020-01-10", "commissioned"),
                  AssetEvent("2022-08-01", "decommissioned"), easting=695000.0)
    due = _asset(AssetEvent("2020-01-10", "commissioned"), easting=696000.0)
    result = R.service_rounds([fresh, gone, due], DAY)
    assert [s["asset_id"] for r in result["rounds"] for s in r["stops"]] == [due.asset_id]
    assert (result["n_boreholes"], result["n_stops"], result["n_nothing_due"],
            result["n_decommissioned"]) == (3, 1, 1, 1)
    summary = R.rounds_text(result)["summary"]
    assert summary == [phrase("rounds.headline", date="2024-06-01", total=3, due=1, rounds=1),
                       phrase("rounds.nothing_due", count=1),
                       phrase("rounds.decommissioned", count=1)]


def test_the_round_is_a_nearest_neighbour_walk_from_the_stop_due_longest():
    """Four boreholes on a line at 0, 1, 3 and 6 km east. The one at 3 km
    has been due longest, so the walk starts there and goes each time to the
    nearest stop left: 1 km (2 km away), 0 km (1 km), then 6 km (6 km)."""
    long_ago = AssetEvent("2019-01-10", "commissioned")
    lately = AssetEvent("2022-01-10", "commissioned")
    at = {km: _asset(long_ago if km == 3 else lately, easting=694000.0 + 1000.0 * km)
          for km in (0, 1, 3, 6)}
    result = R.service_rounds(list(at.values()), DAY)
    stops = result["rounds"][0]["stops"]
    assert [s["asset_id"] for s in stops] == [at[k].asset_id for k in (3, 1, 0, 6)]
    assert [s["leg_m"] for s in stops] == [None, 2000, 1000, 6000]
    assert [s["order"] for s in stops] == [1, 2, 3, 4]
    assert result["rounds"][0]["distance_m"] == 9000
    assert R.round_text(result["rounds"][0]) == [
        phrase("rounds.round_stops", mechanic=phrase("rounds.unassigned"), stops=4),
        phrase("rounds.round_distance", km=9.0), phrase("rounds.round_unassigned")]
    assert [row[6] for row in R.round_rows(result["rounds"][0])] == ["", "2.0", "1.0", "6.0"]


def test_a_tie_goes_to_the_lower_identifier():
    """Two boreholes the same distance either side of the start go lower
    identifier first; of two due on the same day, the walk starts at the
    lower identifier."""
    lately = AssetEvent("2020-01-10", "commissioned")
    middle = _asset(AssetEvent("2019-01-10", "commissioned"), easting=695000.0)
    west = _asset(lately, easting=694000.0)
    east = _asset(lately, easting=696000.0)
    assert west.asset_id < east.asset_id
    stops = R.service_rounds([east, west, middle], DAY)["rounds"][0]["stops"]
    assert [s["asset_id"] for s in stops] == [middle.asset_id, west.asset_id, east.asset_id]
    assert [s["leg_m"] for s in stops] == [None, 1000, 2000]
    pair = R.service_rounds([east, west], DAY)["rounds"][0]["stops"]
    assert [s["asset_id"] for s in pair] == [west.asset_id, east.asset_id]


def test_stops_with_no_position_come_last_by_identifier():
    events = [AssetEvent("2020-01-10", "commissioned")]
    placed = _asset(*events)
    lost = [Asset(asset_id=_id(e), events=list(events)) for e in (700000.0, 690000.0)]
    result = R.service_rounds(lost + [placed], DAY)
    stops = result["rounds"][0]["stops"]
    assert [s["asset_id"] for s in stops] == [placed.asset_id] + sorted(a.asset_id for a in lost)
    assert [s["placed"] for s in stops] == [True, False, False]
    assert stops[1]["position"] == "" and stops[1]["leg_m"] is None
    assert phrase("rounds.unplaced", count=2) in R.round_text(result["rounds"][0])
    # a position that is not two finite numbers is no position
    for easting in ("694912", 10 ** 400, float("nan"), True):
        odd = Asset(asset_id=_id(701000.0), easting=easting, northing=938150.0,
                    events=list(events))
        assert not R.service_rounds([odd], DAY)["rounds"][0]["stops"][0]["placed"], easting


def test_rounds_are_by_mechanic_with_the_unassigned_last():
    events = [AssetEvent("2020-01-10", "commissioned")]
    table = R.parse_area_mechanics("Koya: M. Kolleh\nKori: A. Bangura")["entries"]
    result = R.service_rounds([
        _asset(*events, easting=694000.0, chiefdom="Koya"),
        _asset(*events, easting=695000.0, chiefdom="Kori"),
        _asset(*events, easting=696000.0, chiefdom="Bo Town"),
        _asset(*events, easting=697000.0, chiefdom="Koya", area_mechanic="Zainab"),
    ], DAY, mechanics=table)
    assert [r["mechanic"] for r in result["rounds"]] == ["A. Bangura", "M. Kolleh", "Zainab", ""]
    assert [s["mechanic_from"] for r in result["rounds"] for s in r["stops"]] == [
        "chiefdom", "chiefdom", "asset", "none"]


def test_records_of_one_borehole_are_one_stop_with_their_histories_merged():
    one = _asset(AssetEvent("2020-01-10", "commissioned"))
    phone = Asset(asset_id=one.asset_id, area_mechanic="A. Bangura", rising_main=AGGRESSIVE,
                  events=[AssetEvent("2024-05-30", "routine_inspection"),
                          AssetEvent("2020-01-10", "commissioned")])
    result = R.service_rounds([one, phone, {"asset_id": "not one"}], DAY)
    assert (result["n_boreholes"], result["n_merged"], result["n_unreadable"]) == (1, 1, 1)
    stop = result["rounds"][0]["stops"][0]
    assert stop["mechanic"] == "A. Bangura" and stop["community"] == "Dr. Timbo's"
    assert [i["key"] for i in stop["due"]] == ["inspection", "water_sample"]
    words = R.rounds_text(result)["summary"]
    assert phrase("rounds.merged", count=1) in words and phrase("rounds.unreadable", count=1) in words


def test_the_day_is_an_input_never_the_clock():
    asset = _asset(AssetEvent("2020-01-10", "commissioned"))
    assert R.service_rounds([asset], DAY) == R.service_rounds([asset], "2024-06-01")
    for bad in (None, "", "20240601", "1 June 2024", "2024-02-30"):
        with pytest.raises(ValueError, match="the day it is drawn up"):
            R.service_rounds([asset], bad)


def test_the_order_the_records_arrive_in_does_not_change_the_rounds():
    rng = random.Random(7)
    # distinct positions, so no two records name one borehole: those are
    # read first record first, which is the one place order counts
    assets = [_asset(AssetEvent(f"20{rng.randint(15, 23)}-0{rng.randint(1, 9)}-1{rng.randint(0, 9)}",
                                "commissioned"),
                     easting=690000.0 + east, northing=930000.0 + rng.randint(0, 9000),
                     chiefdom=rng.choice(["Koya", "Kori"]))
              for east in rng.sample(range(0, 9000, 100), 9)]
    table = R.parse_area_mechanics("Koya: M. Kolleh")["entries"]
    first = R.service_rounds(assets, DAY, mechanics=table)
    for _ in range(4):
        rng.shuffle(assets)
        assert R.service_rounds(assets, DAY, mechanics=table) == first


def test_the_intervals_are_stated_as_choices():
    words = R.rounds_text(R.service_rounds([], DAY))
    assert words["summary"] == [phrase("rounds.empty")]
    assert "not a standard" in words["intervals"] and "every 6 months" in words["intervals"]
    assert "not an optimal route" in words["method"]
    plan = replace(Schedule(), water_sample_months=24)
    assert "sample every 24" in R.rounds_text(R.service_rounds([], DAY, plan))["intervals"]


# -------------------------------------------------------------- the sheet

def test_a_stop_on_the_sheet_says_what_is_due_why_and_where_to_write():
    result = R.service_rounds([_asset(AssetEvent("2020-01-10", "commissioned"),
                                      chiefdom="Koya", area_mechanic="A. Bangura",
                                      rising_main=AGGRESSIVE)], DAY)
    stop = result["rounds"][0]["stops"][0]
    sheet = R.stop_sheet(stop)
    assert sheet["heading"] == "Stop 1: Dr. Timbo's (Western Area Rural)"
    details = dict(sheet["details"])
    words = phrase_table("rounds.sheet_details")
    assert details[words["plate"]] == f"BOREHOLE {stop['asset_id']}"
    assert details[words["position"]] == stop["position"] != ""
    assert details[words["leg"]] == phrase("rounds.sheet_first")
    assert details[words["mechanic"]] == "A. Bangura (named on the borehole's record)"
    sanitary, routine, sample = sheet["visits"]
    assert [row[1] for row in sanitary["rows"]] == form_questions()
    assert all(row[-1] == "" for visit in sheet["visits"] for row in visit["rows"])
    assert routine["text"].startswith("Routine inspection, overdue. Routine inspection has never")
    assert routine["text"].endswith("rising main in aggressive water.")
    assert [row[0] for row in sample["rows"]] == list(phrase_table("rounds.readings_sample").values())
    assert [row[0] for row in sheet["visit"]["rows"]] == list(
        phrase_table("rounds.readings_visit").values())


def test_the_round_sheet_carries_every_stop_and_rebuilds_byte_for_byte(tmp_path):
    from docx import Document

    from groundwater.reporting.rounds import build_round_sheet

    events = [AssetEvent("2020-01-10", "commissioned")]
    result = R.service_rounds([
        _asset(*events, easting=694000.0, area_mechanic="A. Bangura"),
        _asset(*events, easting=695000.0),
    ], DAY)
    first = build_round_sheet(result, tmp_path / "a.docx")
    again = build_round_sheet(result, tmp_path / "b.docx")
    assert first.read_bytes() == again.read_bytes()
    document = Document(str(first))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = [cell.text for t in document.tables for row in t.rows for cell in row.cells]
    assert "SERVICE ROUND" in text and "Area mechanic: A. Bangura" in text
    assert phrase("rounds.sheet_mechanic", mechanic=phrase("rounds.unassigned")) in text
    assert "Drawn up on 2024-06-01" in text and R.rounds_text(result)["method"] in text
    for one in result["rounds"]:
        for stop in one["stops"]:
            assert stop["asset_id"] in cells and stop["plate"] in cells
            assert R.stop_sheet(stop)["heading"] in text
    assert form_questions()[0] in cells
    empty = build_round_sheet(R.service_rounds([], DAY), tmp_path / "c.docx")
    assert phrase("rounds.empty") in "\n".join(p.text for p in Document(str(empty)).paragraphs)


# ------------------------------------------- the browser's engine, held to it

_DRIVER = r"""
import { readFileSync } from 'node:fs';
import vm from 'node:vm';
const sandbox = vm.createContext({ console, structuredClone, TextEncoder, TextDecoder });
for (const f of ['gwt-data.js', 'gwt-geo.js', 'gwt-core.js', 'support.js']) {
  vm.runInContext(readFileSync(process.argv[2] + '/' + f, 'utf8'), sandbox, { filename: f });
}
const C = sandbox.GWT.core;
const cases = JSON.parse(readFileSync(0, 'utf8'));
process.stdout.write(JSON.stringify(cases.map((r) => {
  try {
    const plan = Object.assign(C.serviceSchedule(), r.schedule || {});
    const parsed = C.parseAreaMechanics(r.mechanics);
    const result = C.serviceRounds(r.assets, r.as_of, plan, parsed.entries);
    const states = r.assets.map((raw) => {
      const asset = C.assetFromDict(raw);
      if (!asset) return null;
      const st = C.assetState(asset, r.as_of, plan);
      delete st.overdue;
      delete st.is_working;
      return st;
    });
    return { ok: {
      parsed, result, words: C.roundsText(result),
      rounds: result.rounds.map((one) => ({ text: C.roundText(one), rows: C.roundRows(one),
        sheets: one.stops.map((stop) => C.stopSheet(stop)) })),
      states, header: C.roundHeader(),
    } };
  } catch (e) {
    return { threw: String(e && e.stack || e) };
  }
})));
"""

_MECHANICS = ["", "A. Bangura", " M. Kolleh ", "Fatmata", 7, "Ibrahim \U0001f600", "zainab"]
_CHIEFDOMS = ["Koya", "KOYA", " koya ", "Kori", "", None, "Bumpe Ngao", "Kpaka"]
# with days that fall due inside the grace of the days the rounds are
# drawn up on (2023-05-20 and 2023-11-20 for 2024-06-01)
_DAYS = ["2019-03-01", "2020-01-10", "2022-08-01", "2023-05-20", "2023-11-20", "2023-12-01",
         "2024-01-31", "2024-05-01", "2024-06-01", "2024-07-15", "2099-01-01",
         "not written down"]
_KINDS = ["commissioned", "inspection", "routine_inspection", "water_sample", "repair",
          "failure", "restored", "decommissioned", "note", "something else"]


def _case(seed: int) -> dict:
    """A register, a table of mechanics, a day and a schedule from one seed."""
    rng = random.Random(seed)
    ids: list[str] = []
    assets = []
    for _ in range(rng.randint(0, 9)):
        # a small patch of zone 28 so legs tie and stops coincide, now and
        # then zone 29, and now and then a position that is not one
        easting = rng.choice([694000.0, 694500.0, 695000.0, 696000.0, 694912.0,
                              rng.uniform(690000.0, 700000.0)])
        northing = rng.choice([938150.0, 938650.0, rng.uniform(930000.0, 940000.0)])
        zone: object = 28
        if rng.random() < 0.15:
            easting, northing, zone = rng.uniform(250000.0, 300000.0), 925600.0, 29
        if ids and rng.random() < 0.2:
            asset_id = rng.choice(ids)           # a second record of one borehole
        else:
            asset_id = _id(easting, northing, int(zone),
                           rng.choice(["Western Area Rural", "Bo", "Kailahun"]))
            ids.append(asset_id)
        record: dict = {"asset_id": asset_id,
                        "community": rng.choice(["Dr. Timbo's", "Rokel", "", None, 5]),
                        "district": rng.choice(["Western Area Rural", "Bo", ""]),
                        "chiefdom": rng.choice(_CHIEFDOMS),
                        "easting": easting, "northing": northing, "utm_zone": zone}
        roll = rng.random()
        if roll < 0.1:
            record.update(easting=None)
        elif roll < 0.15:
            record.update(easting=str(easting))
        elif roll < 0.2:
            record.update(utm_zone=rng.choice([None, "28", 0, 28.5, 61]))
        if rng.random() < 0.5:
            record["area_mechanic"] = rng.choice(_MECHANICS)
        if rng.random() < 0.6:
            record["rising_main"] = rng.choice([
                AGGRESSIVE, PLAIN, {"material": "uPVC", "inspection_months": 3, "chemistry": "none"},
                {"inspection_months": 9}, {"inspection_months": 6.5}, {"material": "x"},
                {"material": " stainless steel ", "inspection_months": 6.0, "chemistry": "incomplete"}])
        record["events"] = [
            {"when": rng.choice(_DAYS), "kind": rng.choice(_KINDS),
             "by": rng.choice(["", "A. Bangura"]), "note": rng.choice(["", "seals"])}
            for _ in range(rng.choice([0, 1, 1, 2, 3, 5]))]
        if rng.random() < 0.05:
            record = {"asset_id": rng.choice(["", "SL-WAR-XXXXXXX-9", None])}
        assets.append(record)
    lines = []
    for _ in range(rng.randint(0, 5)):
        chiefdom = rng.choice(["Koya", "KOYA", "Kori", "Bumpe Ngao", ""])
        mechanic = rng.choice(["M. Kolleh", "A. Bangura", "", "I. Sesay: senior"])
        lines.append(rng.choice([f"{chiefdom}: {mechanic}", "no colon", "",
                                 "  Kpaka :Zainab  \r"]))
    schedule = rng.choice([{}, {}, {"water_sample_months": 3, "grace_days": 0},
                           {"sanitary_inspection_months": 12, "routine_inspection_months": 24}])
    return {"assets": assets, "mechanics": "\n".join(lines),
            "as_of": rng.choice(["2024-06-01", "2024-02-29", "2021-01-01"]),
            "schedule": schedule}


def _python(case: dict) -> dict:
    plan = Schedule(**{**Schedule().__dict__, **case["schedule"]})
    parsed = R.parse_area_mechanics(case["mechanics"])
    result = R.service_rounds(case["assets"], case["as_of"], plan, parsed["entries"])
    states = []
    for raw in case["assets"]:
        asset = asset_from_dict(raw)
        states.append(None if asset is None else asset_state(asset, date.fromisoformat(
            case["as_of"]), plan).as_dict())
    return json.loads(json.dumps({
        "parsed": parsed, "result": result, "words": R.rounds_text(result),
        "rounds": [{"text": R.round_text(one), "rows": R.round_rows(one),
                    "sheets": [R.stop_sheet(stop) for stop in one["stops"]]}
                   for one in result["rounds"]],
        "states": states, "header": R.round_header(),
    }))


SEEDS = range(60)


@pytest.fixture(scope="module")
def browser_answers(tmp_path_factory):
    node = shutil.which("node")
    if node is None:
        pytest.skip("node is not installed; run `node --version` to check")
    driver = tmp_path_factory.mktemp("rounds") / "rounds.mjs"
    driver.write_text(_DRIVER, encoding="utf-8")
    done = subprocess.run([node, str(driver), str(JS)],
                          input=json.dumps([_case(seed) for seed in SEEDS]),
                          capture_output=True, text=True, encoding="utf-8", check=True)
    return json.loads(done.stdout)


@pytest.mark.parametrize("seed", SEEDS)
def test_the_browser_draws_the_same_rounds(browser_answers, seed):
    answer = browser_answers[seed]
    assert "threw" not in answer, answer.get("threw")
    assert answer["ok"] == _python(_case(seed))


def test_the_seeded_registers_reach_every_rule():
    told = [_python(_case(seed)) for seed in SEEDS]
    stops = [s for t in told for r in t["result"]["rounds"] for s in r["stops"]]
    assert {s["mechanic_from"] for s in stops} == {"asset", "chiefdom", "none"}
    assert {i["state"] for s in stops for i in s["due"]} == {"overdue", "due", "unknown"}
    assert {s["function"] for s in stops} >= {"functional", "non_functional", "unknown"}
    assert any(not s["placed"] for s in stops) and any(s["leg_m"] == 0 for s in stops)
    assert any(len(r["stops"]) > 3 for t in told for r in t["result"]["rounds"])
    assert any(t["result"]["n_merged"] for t in told)
    assert any(t["result"]["n_unreadable"] for t in told)
    assert any(t["result"]["n_decommissioned"] for t in told)
    assert any(t["parsed"]["skipped"] for t in told)
    bases = " ".join(i["basis"] for s in stops for i in s["due"]
                     if i["key"] == "routine_inspection")
    assert "design sets" in bases and "record sets" in bases
    assert "names no rising main" in bases
