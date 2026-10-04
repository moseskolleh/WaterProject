"""The printed field kit (PLAN.md step 2.5).

The kit's content is held to the browser's by the parity suite
(tests/webapp/parity.mjs, against reference.json) and the browser's button
by tests/webapp/fieldkit.mjs, which commits the document it wrote to
``tests/webapp/fixtures/fieldkit_browser.docx``. Here: the schedules are one
list, read everywhere; the sheet code is the documented format and reads
back; the QR images in both engines' documents decode, with OpenCV, to
exactly the code; and a filled sheet is typed into the standard template
and read back by the Python reader.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import numpy as np
import pytest

from groundwater.config import Config, PumpingConfig
from groundwater.field_kit import (
    PAYLOAD_FORMAT,
    PayloadFields,
    airlift_yield,
    cautious_storage,
    field_kit_content,
    field_kit_payload,
    field_schedules,
    parse_field_kit_payload,
    project_identifier,
    reading_minutes,
    ves_survey_plan,
)
from groundwater.hydraulics.analysis import casing_storage_min
from groundwater.ingestion.templates import (
    PUMPING_COLUMNS,
    PUMPING_HEADER,
    RECOVERY_COLUMNS,
    write_pumping_template,
)
from groundwater.ingestion.pumping import read_pumping_workbook
from groundwater.models import SiteMetadata
from groundwater.supervision.field_checks import disinfection_dose

REPO = Path(__file__).resolve().parents[1]
JS = REPO / "docs" / "js"
BROWSER_KIT = REPO / "tests" / "webapp" / "fixtures" / "fieldkit_browser.docx"

# the site and boreholes tests/webapp/fieldkit.mjs builds the browser's kit for
SITE = SiteMetadata(project="Rokel 2026", community="Kuntolo",
                    client="Living Water International", district="Port Loko",
                    supervisor="WiNGiN")
BOREHOLES = ["KTL-01", "KTL|02 %x"]
CODES = ["GWT-FK/1|pumping|Rokel 2026|KTL-01",
         "GWT-FK/1|pumping|Rokel 2026|KTL%7C02 %25x"]


# ------------------------------------------------------- one list, read everywhere

def test_the_schedule_is_plan_md_s():
    """PLAN.md step 2.1's list, then every 30 minutes."""
    assert reading_minutes(120) == [0.5, 1, 1.5, 2, 3, 4, 5, 6, 8, 10, 12, 15, 20, 25,
                                    30, 40, 50, 60, 75, 90, 120]
    assert reading_minutes(240)[-4:] == [120, 150, 180, 210, 240][-4:]
    assert reading_minutes(0.4) == []


def test_no_browser_script_types_out_a_schedule():
    """The co-pilots and the blank template read data/field.yaml; none of
    the scripts carries a copy of a list or of a field rule."""
    field = field_schedules()
    source = "\n".join(p.read_text(encoding="utf-8") for p in sorted(JS.glob("*.js"))
                       if p.name not in ("gwt-data.js", "gwt-geo.js", "gwt-samples.js"))
    for values in (field["pumping"]["schedule_min"], field["ves"]["ab2_series_m"],
                   field["ves"]["mn_series_m"]):
        # eight of a list's entries in a row is a copy of it (fewer are not:
        # the map's tick steps run 3, 4, 5, 6, 8, 10 too)
        for i in range(len(values) - 7):
            run = r",\s*".join(re.escape(f"{v:g}") for v in values[i:i + 8])
            assert not re.search(r"(?<![\d.])" + run + r"(?![\d.])", source), values[i:i + 8]
    for name in ("LOGAN", "T_LOW", "LATE_EVERY_MIN", "MIN_AB_PER_MN", "MAX_AB_PER_MN"):
        assert not re.search(rf"\b{name}\s*=\s*\d", source), name
    # the time column of the browser's blank pumping template, too
    assert "[1, 2, 3, 4, 5, 6, 8, 10, 15, 20" not in source


def test_the_browser_bundle_carries_the_file_as_python_reads_it():
    text = (JS / "gwt-data.js").read_text(encoding="utf-8")
    match = re.search(r'"field":\s*(\{.*?"disinfection_card":\s*\{[^}]*\}\s*\})', text)
    assert match, "gwt-data.js carries no field schedules; run python web/build_webapp_data.py"
    assert json.loads(match.group(1)) == field_schedules()


def test_the_storage_time_is_the_engine_s_at_the_cautious_floor():
    """casing_storage_min at the co-pilot's low transmissivity, through
    Logan's factor: the figure the pumping co-pilot gives before pumping."""
    storage = cautious_storage(PumpingConfig())
    assert storage["t_low"] == 1 and storage["t_high"] == 10
    assert storage["low_t_min"] == casing_storage_min(1 / (1.22 * 24), PumpingConfig())
    # a 5 inch casing and a 1.25 inch riser: about five hours at 1 m2/day
    assert storage["low_t_min"] == pytest.approx(306.8, abs=0.1)
    assert storage["high_t_min"] == pytest.approx(30.7, abs=0.1)


def test_the_sheet_reads_until_the_pump_may_stop():
    content = field_kit_content(SITE, ["BH-1"])
    assert content["constant_min"] == pytest.approx(306.8, abs=0.1)
    # the first scheduled minute past it
    assert content["planned_min"] == 330
    sheet = content["sheets"][0]
    times = [float(row[0]) for block in sheet["blocks"] for row in block["rows"]]
    assert times == reading_minutes(330)
    assert [float(r[0]) for r in sheet["recovery"]["rows"]] == reading_minutes(330)
    assert [b["caption"] for b in sheet["blocks"]] == [
        "Constant discharge 0-60 min", "Constant discharge 61-120 min",
        "Constant discharge 121-180 min", "Constant discharge 181-330 min"]
    assert any("do not stop the pump before 307 minutes" in n for n in sheet["notes"])


def test_the_ves_card_is_the_co_pilot_s_proposal_carried_to_the_end_of_the_line():
    content = field_kit_content(SITE, ["BH-1"])
    card = next(c for c in content["cards"] if c["key"] == "ves")
    rows = card["tables"][0]["rows"]
    full = ves_survey_plan(250.0)
    assert [(float(r[1]), float(r[2])) for r in rows] == [(s["ab2"], s["mn"])
                                                          for s in full["steps"]]
    # a shorter target is the first rows of it
    plan = ves_survey_plan(40.0)
    assert plan["max_ab2"] == 80 and plan["steps"] == full["steps"][:len(plan["steps"])]
    # every MN change repeats its AB/2, and AB is never under five MN
    for before, after in zip(plan["steps"], plan["steps"][1:], strict=False):
        if after["mn"] != before["mn"]:
            assert after["ab2"] == before["ab2"]
    assert all(2 * s["ab2"] >= 5 * s["mn"] for s in full["steps"])


def test_the_dose_card_is_the_supervision_calculator_s():
    content = field_kit_content(SITE, ["BH-1"])
    card = next(c for c in content["cards"] if c["key"] == "disinfection")
    row = card["tables"][0]["rows"][3]          # 40 m of water
    dose = disinfection_dose(40, 125)
    assert row[0] == "40"
    assert row[3:5] == [f"{dose.solution_02pct_l:.1f}", f"{dose.hth_grams:.0f}"]
    assert "20 mg of chlorine per litre" in card["lines"][0]
    assert "at least 4 hours" in card["lines"][0]
    assert "3.1 g of 65 percent HTH" in card["lines"][1]
    assert any("Adekile" in r for r in content["references"])


# ------------------------------------------------------------- the sheet code

def test_the_code_is_the_documented_format():
    assert field_kit_payload("Rokel 2026", "KTL-01") == "GWT-FK/1|pumping|Rokel 2026|KTL-01"
    assert field_kit_payload(" A|B%C\n ", "BH\t 3") == "GWT-FK/1|pumping|A%7CB%25C|BH 3"
    assert PAYLOAD_FORMAT == "GWT-FK/1"


def test_a_field_is_trimmed_of_the_ascii_spaces_only():
    """str.strip() takes U+0085 and U+001C to U+001F off the ends, which
    String.trim() keeps, and String.trim() takes a byte-order mark, which
    str.strip() keeps, so the two engines made different codes for one
    name. Both now trim the six ASCII spaces and nothing else; parity holds
    the browser to it on these names and more."""
    assert field_kit_payload("﻿Rokel", "BH-1\x85") == "GWT-FK/1|pumping|﻿Rokel|BH-1\x85"
    assert field_kit_payload("\x1cP\x1f", " B　") == "GWT-FK/1|pumping|\x1cP\x1f| B　"
    assert parse_field_kit_payload("﻿GWT-FK/1|pumping|P|B") is None
    assert parse_field_kit_payload(" \t\nGWT-FK/1|pumping|P|B\x85\r\n").borehole == "B\x85"


@pytest.mark.parametrize("project,borehole", [
    ("Rokel 2026", "KTL-01"), ("A|B%C", "%41"), ("", "RK-1"), ("Kɔnɔ", "Ø-1|%7C"),
])
def test_a_code_reads_back_to_what_went_into_it(project, borehole):
    parsed = parse_field_kit_payload(field_kit_payload(project, borehole))
    assert parsed == PayloadFields("GWT-FK/1", "pumping", project, borehole)


@pytest.mark.parametrize("text", [
    "GWT-FK/2|pumping|Rokel 2026|KTL-01",    # a version this reader does not know
    "GWT-FK/1|pumping|Rokel 2026",           # a field short
    "GWT-FK/1||Rokel 2026|KTL-01",           # no sheet
    "BOREHOLE SL-WAR-8FEEVKQ-T",             # the headworks plate's code
])
def test_text_that_is_not_a_code_is_not_read_as_one(text):
    assert parse_field_kit_payload(text) is None


def test_the_project_is_its_reference_else_its_name():
    assert project_identifier(SiteMetadata(project="Name", project_ref=" REF/1 ")) == "REF/1"
    assert project_identifier(SiteMetadata(project="Name")) == "Name"
    sheet = field_kit_content(SiteMetadata(community="Rokel"), ["RK-1"])["sheets"][0]
    assert sheet["payload"] == "GWT-FK/1|pumping||RK-1"
    assert "cannot match a photographed sheet to a project" in sheet["warning"]


def test_a_code_too_long_for_its_symbol_is_refused(tmp_path):
    from groundwater.reporting.field_kit import build_field_kit

    with pytest.raises(ValueError):
        build_field_kit(SiteMetadata(project="P" * 200), ["BH-1"], tmp_path / "k.docx",
                        tmp_path)
    with pytest.raises(ValueError, match="at least one borehole"):
        build_field_kit(SITE, [" ", ""], tmp_path / "k.docx", tmp_path)


# ------------------------------------------------- the documents, and their QR

def _decoded_symbols(docx_path: Path) -> list[str]:
    cv2 = pytest.importorskip("cv2", reason="no decoder installed")
    with zipfile.ZipFile(docx_path) as z:
        names = sorted((n for n in z.namelist() if n.startswith("word/media/")),
                       key=lambda n: int(re.sub(r"\D", "", n)))
        images = [z.read(n) for n in names]
    out = []
    for png in images:
        image = cv2.imdecode(np.frombuffer(png, dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
        text, *_ = cv2.QRCodeDetector().detectAndDecode(image)
        out.append(text)
    return out


@pytest.fixture(scope="module")
def python_kit(tmp_path_factory):
    from groundwater.reporting.field_kit import build_field_kit

    folder = tmp_path_factory.mktemp("kit")
    return build_field_kit(SITE, BOREHOLES, folder / "field_kit.docx", folder)


def test_the_python_kit_s_symbols_decode_to_their_codes(python_kit):
    assert _decoded_symbols(python_kit) == CODES


def test_the_browser_kit_s_symbols_decode_to_their_codes():
    """The document the browser's button wrote (tests/webapp/fieldkit.mjs)."""
    assert _decoded_symbols(BROWSER_KIT) == CODES


def test_the_python_kit_prints_its_content(python_kit):
    import docx

    document = docx.Document(str(python_kit))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = [c.text for t in document.tables for row in t.rows for c in row.cells]
    content = field_kit_content(SITE, BOREHOLES)
    for sheet in content["sheets"]:
        assert sheet["title"] in text and sheet["code"] in text
        for note in sheet["notes"]:
            assert note in text
    for card in content["cards"]:
        assert card["title"] in text
        for line in card["lines"]:
            assert line in text
    # the header labels and the column headings are the template's
    labels = [label for _, label, _ in PUMPING_HEADER]
    assert all(label in cells for label in labels)
    assert all(h in cells for h in PUMPING_COLUMNS + RECOVERY_COLUMNS)
    assert cells.count("Time (min)") == 2 * 5        # four blocks and the recovery


# ----------------------------------------- transcribed into the standard template

def test_a_filled_sheet_types_into_the_template_and_reads_back(tmp_path):
    """Rule 6: the paper sheet's header and Time column go into
    template_pumping_test.xlsx cell for cell, and the reader finds them."""
    from openpyxl import load_workbook

    sheet = field_kit_content(SITE, ["KTL-01"])["sheets"][0]
    path = write_pumping_template(tmp_path / "template_pumping_test.xlsx")
    wb = load_workbook(path)
    ws = wb.active
    printed = dict(sheet["header"])
    # what the crew writes in on the day
    printed.update({"Date": "2026-10-03", "Static water level (m)": "9.44",
                    "Pump setting (m)": "67", "Depth of Borehole (m)": "70"})
    for _, label, value_cell in PUMPING_HEADER:
        value = printed[label]
        ws[value_cell] = float(value) if re.fullmatch(r"[\d.]+", value) else value
    ws["C9"] = 2.93
    # each block's Time column, as printed, into the template's group of
    # the same heading, with a level beside each; the recovery likewise
    times = []
    for block, column in zip(sheet["blocks"], (1, 4, 7, 10), strict=True):
        for k, row in enumerate(block["rows"]):
            ws.cell(row=12 + k, column=column, value=float(row[0]))
            ws.cell(row=12 + k, column=column + 1, value=10 + 0.1 * len(times))
            times.append(float(row[0]))
    recovery = [float(r[0]) for r in sheet["recovery"]["rows"]]
    for k, t in enumerate(recovery):
        ws.cell(row=12 + k, column=13, value=t)
        ws.cell(row=12 + k, column=14, value=40 - 0.1 * k)
    wb.save(path)

    test = read_pumping_workbook(path)
    assert test.borehole_ref == "KTL-01"
    assert test.site.community == "Kuntolo" and test.site.district == "Port Loko"
    assert test.site.supervisor == "WiNGiN"
    assert test.static_water_level_m == 9.44 and test.pump_setting_m == 67
    assert test.test_type.startswith("constant")
    # one constant series, the whole schedule to 330 minutes, in order
    assert len(test.steps) == 1
    assert test.steps[0].discharge_m3_per_h == pytest.approx(2.93)
    np.testing.assert_array_equal(test.steps[0].time_min, reading_minutes(330))
    np.testing.assert_array_equal(test.recovery_time_min, recovery)
    assert not [f for f in test.flags if f.level == "error"]


def test_the_kit_follows_the_configuration():
    config = Config()
    config.pumping.casing_diameter_in = 6.0
    config.pumping.min_constant_test_min = 120.0
    content = field_kit_content(SITE, ["BH-1"], config)
    assert content["storage"]["casing_in"] == 6.0
    assert content["constant_min"] == pytest.approx(
        casing_storage_min(1 / (1.22 * 24), config.pumping))


# ------------------------------------------------------------ airlift yield

def test_the_airlift_yield_of_a_timed_container_is_volume_over_mean_time():
    result = airlift_yield("bucket", volume_l=20, timings_s=[25, 24.6, 25.4])
    assert result["q_l_per_s"] == pytest.approx(0.8)
    assert result["basis"] == "Timed container: 20 L filled in a mean of 25.0 s over 3 timings."
    assert not result["flags"]


def test_the_v_notch_reads_as_the_published_90_degree_formula():
    """Kindsvater-Shen at Ce = 0.58 against Thomson's 1.38 h^2.5 m3/s, the
    form most field tables print for a 90-degree notch: within 1 percent
    over the head range the coefficient holds for, and a head outside it
    is said to be."""
    for head_mm in (50, 100, 200, 380):
        q = airlift_yield("vnotch", head_mm=head_mm)
        assert q["q_l_per_s"] == pytest.approx(1.38 * (head_mm / 1000) ** 2.5 * 1000, rel=0.01)
        assert not q["flags"]
    low = airlift_yield("vnotch", head_mm=34)
    assert [f["code"] for f in low["flags"]] == ["vnotch_head_outside"]


def test_an_airlift_reading_that_gives_no_yield_is_refused():
    for method, options in (("bucket", {"volume_l": 20, "timings_s": []}),
                            ("bucket", {"volume_l": 20, "timings_s": [10, 0]}),
                            ("vnotch", {"head_mm": 0}), ("none", {"reason": "  "}),
                            ("pump", {})):
        with pytest.raises(ValueError):
            airlift_yield(method, **options)
    assert airlift_yield("none", reason=" compressor  down ")["basis"] == (
        "Not measured: compressor down.")
