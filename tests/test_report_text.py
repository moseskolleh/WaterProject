"""What the reports say is what the inputs support.

The geology paragraph used to be chosen by the substring "western" in the
district name and said "Freetown Basic Complex" of a site the report's own
maps placed on the Bullom Group; the field-work section asserted a
reconnaissance, a geomorphological survey and pegs for every survey; the
table of contents told the reader to right-click; total coliforms were a
WHO health failure and "faecal contamination" with E. coli at zero.
"""

from __future__ import annotations

from docx import Document

from groundwater.models import SiteMetadata, WaterQualityResult, WaterQualitySample
from groundwater.quality import assess_sample
from groundwater.quality.diagrams import facies_of
from groundwater.reporting.docx_utils import ReportBuilder
from groundwater.reporting.geophysical import _geology_for
from groundwater.utils import plural, plural_noun, utm_text


def _text(path) -> str:
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_the_geology_paragraph_comes_from_the_map_under_the_site():
    # Rokel A (1): Western Area by district, Bullom Group by the map
    rokel = SiteMetadata(community="Rokel", district="Western Area",
                         easting=708958, northing=926355, utm_zone=28)
    text = _geology_for(rokel, "")
    assert "Bullom Group" in text
    assert "Freetown" not in text
    assert "BGS Africa Groundwater Atlas classes the aquifer" in text
    # without a fix the district's region decides, and the peninsula is
    # the Freetown Complex whatever the district is called
    assert "Freetown Basic Complex" in _geology_for(SiteMetadata(district="Western Area Rural"), "")
    assert "Freetown" not in _geology_for(SiteMetadata(district="Port Loko"), "")
    assert _geology_for(rokel, "my own words") == "my own words"


def test_the_geology_paragraph_names_a_polygon_as_the_map_key_does():
    """A sheet that gets the district wrong does not rename the rock.

    The key on the geology map is scoped by where each polygon is; the
    paragraph was scoped by the sheet's district. A site on the Freetown
    Complex written down as Port Loko got "Paleozoic Igneous (Pi)" and
    nothing else in the paragraph, the age the crosswalk itself calls
    wrong, beside a key naming the Freetown Layered Complex.
    """
    from groundwater.geo import geographic_to_utm
    from groundwater.mapping import geology_unit_at
    from groundwater.mapping.lithology import lithology_for
    from groundwater.mapping.regional import _unit_district

    utm = geographic_to_utm(8.40, -13.18)
    site = SiteMetadata(community="X", district="Port Loko", easting=utm.easting,
                        northing=utm.northing, utm_zone=utm.zone)
    unit = geology_unit_at(*site.latlon)
    assert unit is not None and unit.glg == "Pi"
    key = lithology_for(unit.glg, _unit_district(unit) or site.district)
    text = _geology_for(site, "")
    assert key is not None and key.formation_name == "Freetown Layered Complex"
    assert "Freetown Layered Complex (Jf)" in text


def test_the_field_work_section_reports_only_what_was_recorded(sample_data, tmp_path):
    from groundwater.ingestion import read_ves_workbook
    from groundwater.reporting import build_geophysical_report
    from groundwater.reporting.geophysical import GeophysicalReportInputs
    from groundwater.ves import interpret_model, invert_sounding

    soundings = read_ves_workbook(sample_data / "rokel" / "rokel_ves.xlsx")[:1]
    inversions = [invert_sounding(s) for s in soundings]
    interps = [interpret_model(s, r.model) for s, r in zip(soundings, inversions, strict=True)]
    bare = build_geophysical_report(
        GeophysicalReportInputs(soundings=soundings, inversions=inversions,
                                interpretations=interps, figures_dir=tmp_path),
        tmp_path / "bare.docx",
    )
    text = _text(bare)
    assert "No reconnaissance record" in text
    assert "marked with pegs" not in text
    assert "Geomorphological survey of the area" not in text
    assert "point(s)" not in text and "zone(s)" not in text
    assert "half-space" in text and "0/0" not in text
    assert "708958 m E (UTM zone 28N)" in text and "708,958" not in text
    noted = build_geophysical_report(
        GeophysicalReportInputs(soundings=soundings, inversions=inversions,
                                interpretations=interps, figures_dir=tmp_path,
                                reconnaissance_date="7 December 2015",
                                reconnaissance_notes="Valley floor, stream 200 m east."),
        tmp_path / "noted.docx",
    )
    text = _text(noted)
    assert "conducted on 7 December 2015" in text
    assert "Valley floor, stream 200 m east." in text
    assert "No reconnaissance record" not in text


def test_the_table_of_contents_reads_without_word(tmp_path):
    rb = ReportBuilder(title="t")
    rb.table_of_contents()
    rb.heading("1. Introduction", 1)
    rb.heading("1.1 Setting", 2)
    rb.heading("Table of Contents", 1, numbered=False)
    path = rb.save(tmp_path / "toc.docx")
    doc = Document(str(path))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Right-click" not in body
    assert "1. Introduction" in body and "    1.1 Setting" in body
    settings = doc.settings.element
    assert any(el.tag.endswith("updateFields") for el in settings)


def test_captions_use_the_caption_style_and_empty_cells_are_empty(tmp_path):
    rb = ReportBuilder(title="t")
    rb.table([["a", None]], header=["x", "y"], caption="One row.")
    rb.table([], caption="Nothing.")           # used to crash on max() of nothing
    path = rb.save(tmp_path / "t.docx")
    doc = Document(str(path))
    captions = [p for p in doc.paragraphs if p.text.startswith("Table ")]
    assert captions and all(p.style.name == "Caption" for p in captions)
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert "None" not in cells and "(no entries)" in cells


def test_plurals_and_grid_coordinates():
    assert plural(1, "point") == "1 point" and plural(2, "point") == "2 points"
    assert plural_noun(1, "exceedance") == "exceedance"
    assert plural_noun(3, "exceedance") == "exceedances"
    site = SiteMetadata(easting=778000.0, northing=946000.0, utm_zone=28)
    assert utm_text(site, "easting") == "778000 m E (UTM zone 28N)"
    assert utm_text(site, "northing") == "946000 m N"
    assert utm_text(SiteMetadata(), "easting") == ""


def _sample(*results):
    return WaterQualitySample(site=SiteMetadata(community="X"), sample_id="s",
                              results=list(results))


def test_the_corrosivity_verdict_says_what_the_ph_is():
    low = assess_sample(_sample(
        WaterQualityResult("pH", 5.9), WaterQualityResult("Calcium", 4.0),
        WaterQualityResult("Alkalinity", 10.0), WaterQualityResult("TDS", 60.0),
    )).corrosivity
    assert low.is_aggressive
    assert "The pH of 5.9 is below the 6.5 to 8.5 acceptability range" in low.verdict
    assert "within the acceptability range" not in low.verdict
    mid = assess_sample(_sample(
        WaterQualityResult("pH", 7.0), WaterQualityResult("Calcium", 4.0),
        WaterQualityResult("Alkalinity", 10.0), WaterQualityResult("TDS", 60.0),
    )).corrosivity
    assert "The pH of 7.0 is within the 6.5 to 8.5 acceptability range" in mid.verdict


def test_the_facies_section_says_what_the_water_is(sample_data):
    from groundwater.ingestion import read_quality_workbook

    sample = read_quality_workbook(sample_data / "dr_timbo" / "dr_timbo_water_quality.xlsx")
    facies = facies_of(sample)
    assert facies["facies"] == "mixed-cation-HCO3"
    assert facies["sentence"].startswith("The water is a mixed-cation-HCO3 type")
    assert "milliequivalent percent" in facies["sentence"]
    assert facies_of(_sample(WaterQualityResult("pH", 7.0))) is None
    assessment = assess_sample(sample)
    # manganese fails the WHO guideline, total coliforms the national limit,
    # and the summary names both without calling the second faecal
    assert assessment.verdict_state == "health_fail"
    assert "Total coliforms" in {r.parameter for r in assessment.national_exceedances}
    remark = next(r.remark for r in assessment.all_exceedances if r.parameter == "Total coliforms")
    assert "not of faecal contamination in itself" in remark
    assert "WHO sets no health based guideline" in remark


def test_a_provisional_report_qualifies_its_own_summary(sample_data, tmp_path):
    """The stamp on the cover was contradicted by an unqualified summary."""
    from groundwater.ingestion import read_pumping_workbook
    from groundwater.hydraulics import analyse_pumping_test
    from groundwater.readiness import assess_readiness
    from groundwater.reporting.pumping import PumpingReportInputs, build_pumping_report

    analysis = analyse_pumping_test(
        read_pumping_workbook(sample_data / "dr_timbo" / "dr_timbo_constant_test.xlsx"))
    readiness = assess_readiness({"pump_analysis": analysis}, "pumping")
    assert not readiness.is_certifiable
    path = build_pumping_report(
        PumpingReportInputs(analysis=analysis, figures_dir=tmp_path, readiness=readiness),
        tmp_path / "p.docx",
    )
    text = _text(path)
    summary = text[text.index("Executive Summary"):text.index("1. Test Details")]
    assert "This report is provisional and not a certification" in summary
    assert "Yield established" in summary and "Site position" in summary
    # a certifiable report gets no such sentence
    clean = build_pumping_report(
        PumpingReportInputs(analysis=analysis, figures_dir=tmp_path),
        tmp_path / "c.docx",
    )
    assert "This report is provisional" not in _text(clean)


def test_the_handover_and_completion_reports_say_what_they_hold(sample_data, tmp_path):
    """No marker-less map captioned as the water point, no blank cover
    lines, no duplicated rows, handpump care for a submersible pump, and
    the national limits called provisional where they are printed."""
    from groundwater.design import design_borehole
    from groundwater.hydraulics import analyse_pumping_test
    from groundwater.ingestion import (
        read_drilling_workbook,
        read_pumping_workbook,
        read_quality_workbook,
    )
    from groundwater.reporting.completion import CompletionReportInputs, build_completion_report
    from groundwater.reporting.handover import HandoverReportInputs, build_handover_report, om_guidance

    d = sample_data / "dr_timbo"
    log = read_drilling_workbook(d / "dr_timbo_drilling_log.xlsx")
    analysis = analyse_pumping_test(read_pumping_workbook(d / "dr_timbo_constant_test.xlsx"))
    quality = assess_sample(read_quality_workbook(d / "dr_timbo_water_quality.xlsx"))
    design = design_borehole(log=log, static_water_level_m=9.44)
    handover = build_handover_report(
        HandoverReportInputs(site=log.site, log=log, design=design, pumping=analysis,
                             quality=quality, figures_dir=tmp_path,
                             pump_type="Submersible pump"),
        tmp_path / "h.docx",
    )
    text = _text(handover)
    assert "Location of the water point" not in text
    assert "No GPS position is recorded" in text
    assert "Project: " not in text.split("Executive Summary")[0]
    assert text.count("Total depth") == 1
    assert "strokes per day" not in text and "pump rods" not in text
    assert "meter reading" in text
    assert "provisional" in text.lower()
    assert om_guidance("India Mark II")[1][1][1].startswith("Record the approximate hours")
    completion = build_completion_report(
        CompletionReportInputs(
            log=log, design=design, pumping=analysis, quality=quality, figures_dir=tmp_path,
            development_record=[("17:00", "17:17", "", "Muddy"), ("17:17", "18:00", "2.5", "Clear")],
        ),
        tmp_path / "c.docx",
    )
    text = _text(completion)
    assert "Status as recorded by the driller" in text
    assert "The record covers 1 h 00 min (17:00 to 18:00)." in text
    assert "provisional" in text.lower()


def _tables(path) -> list[list[list[str]]]:
    return [[[c.text for c in row.cells] for row in t.rows] for t in Document(str(path)).tables]


def test_a_tie_is_carried_through_the_whole_report(tmp_path):
    """Two points 2.8 weighted points apart: the suitability section called
    them indistinguishable while the summary recommended one, the conclusions
    selected it "according to the results", and the preference table opened
    on the 2nd under "in order of preference"."""
    import numpy as np

    from groundwater.models import LayeredModel, VESSounding
    from groundwater.reporting import build_geophysical_report
    from groundwater.reporting.geophysical import GeophysicalReportInputs
    from groundwater.ves import interpret_model, invert_sounding
    from groundwater.ves.forward import forward_schlumberger

    ab2 = np.array([1, 1.5, 2, 3, 4, 5, 7, 10, 15, 20, 30, 40, 50, 60, 70, 80.0])
    soundings, inversions, interps = [], [], []
    for sid, thickness in (("VES 1", 18.0), ("VES 2", 20.0)):
        truth = LayeredModel([800, 60, 4000], [4, thickness])
        s = VESSounding(SiteMetadata(community="Testville", district="Bo"), sid, ab2,
                        np.full(len(ab2), 1.0), forward_schlumberger(truth, ab2))
        r = invert_sounding(s)
        soundings.append(s)
        inversions.append(r)
        interps.append(interpret_model(s, r.model))
    path = build_geophysical_report(
        GeophysicalReportInputs(soundings=soundings, inversions=inversions,
                                interpretations=interps, figures_dir=tmp_path),
        tmp_path / "tie.docx",
    )
    text = _text(path)
    assert "VES 2 is ahead by 2.8 points, within the 3-point margin" in text
    assert "by name only" not in text
    assert "is recommended as the preferred drilling location" not in text
    assert "Points VES 2 and VES 1 cannot be told apart on geophysical grounds" in text
    assert "Drilling points the survey cannot separate: VES 2 and VES 1." in text
    assert "is selected as the preferred point" not in text
    assert "Points VES 2 and VES 1 cannot be separated by the results" in text
    assert "Drilling should be carried out at point VES 2 or point VES 1" in text
    preference = next(t for t in _tables(path) if t[0][-1] == "Ranking")
    assert [(row[1], row[-1]) for row in preference[1:]] == [("VES 2", "=1st"),
                                                             ("VES 1", "=1st")]
    assert "The two points marked =1st cannot be told apart" in text


def test_the_rokel_report_labels_its_half_space_and_carries_the_sheet_warnings(
        sample_data, tmp_path):
    """Every model table labelled layer 1, the surface, as the half-space; the
    two overlap warnings the Rokel sheets raise never reached Annex A; and
    the narrative called a boundary known to x/ 3.7 resolved."""
    from groundwater.ingestion import read_ves_workbook
    from groundwater.reporting import build_geophysical_report
    from groundwater.reporting.geophysical import GeophysicalReportInputs
    from groundwater.ves import interpret_model, invert_sounding

    soundings = read_ves_workbook(sample_data / "rokel" / "rokel_ves.xlsx")
    inversions = [invert_sounding(s) for s in soundings]
    interps = [interpret_model(s, r.model) for s, r in zip(soundings, inversions, strict=True)]
    path = build_geophysical_report(
        GeophysicalReportInputs(soundings=soundings, inversions=inversions,
                                interpretations=interps, figures_dir=tmp_path,
                                include_qa_annex=True),
        tmp_path / "rokel.docx",
    )
    text = _text(path)
    models = [t for t in _tables(path) if t[0] == ["N", "rho (ohm-m)", "h (m)", "z (m)"]]
    assert len(models) == 2
    for table in models:
        assert table[1][3] == "0"                       # the surface
        assert table[-1][2] == "half-space"             # the half-space
        assert sum("half-space" in cell for row in table for cell in row) == 1
    assert "[WARNING] segment_overlap_discrepancy (A (1)): " in text
    assert "AB/2 40 m: 156.1 and 78.7 ohm-m (ratio 1.98)" in text
    assert "[WARNING] segment_overlap_discrepancy (B (2)): " in text
    assert "The boundary at 1.02 m (x/ 3.7) is poorly resolved" in text
    assert "The data at A (1) are fitted with a 3 layer model" in text
    assert "The data at A (1) resolves" not in text


def test_models_tried_says_why_a_simpler_model_that_reached_the_target_lost():
    """"The 3-layer model is the simplest that reaches the 10 percent target"
    of a 2-layer model at 4.2 percent: the 2-layer model reached it too, and
    was passed over because the 3-layer one more than halved its misfit."""
    from types import SimpleNamespace

    from groundwater.models import LayeredModel
    from groundwater.reporting.geophysical import models_tried_text

    def trials(tried, err):
        return SimpleNamespace(trials=tried, fit_error_percent=err,
                               model=LayeredModel([300, 100, 150], [3, 30]))

    assert models_tried_text(trials([(2, 4.2), (3, 0.01)], 0.01)) == (
        "Models tried: 2 layers, 4.2%; 3 layers, 0.0%. Of these the 2-layer model "
        "also reaches the 10 percent target, but the 3-layer model more than halves "
        "its misfit and is preferred.")
    assert ("the 4-layer model more than halves its misfit, and the 3-layer model is "
            "the simplest it does not better that far") in models_tried_text(
               trials([(2, 8.0), (3, 3.5), (4, 2.0)], 3.5))
    assert models_tried_text(trials([(2, 15.4), (3, 8.0)], 8.0)).endswith(
        "Of these the 3-layer model is the simplest that reaches the 10 percent target.")


def test_the_limitations_describe_a_wenner_survey_in_its_own_terms():
    """"A Schlumberger sounding ... with AB/2 expanded to 60 m" of a Wenner
    survey whose spacing a was 60 m, so AB/2 was 90 m."""
    import numpy as np

    from groundwater.models import LayeredModel, VESSounding
    from groundwater.reporting.geophysical import _limitations
    from groundwater.ves import interpret_model

    a = np.array([1, 2, 5, 10, 20, 40, 60.0])
    wenner = VESSounding(SiteMetadata(), "W 1", a, np.full(len(a), np.nan),
                         np.full(len(a), 100.0), array_type="wenner")
    interp = interpret_model(wenner, LayeredModel([600, 40, 3000], [3, 25]))
    text = " ".join(_limitations([], [interp], soundings=[wenner]))
    assert ("A Wenner sounding resolves the ground to roughly half its largest "
            "electrode spacing a") in text
    assert ("with a expanded to 60 m (AB/2 of 90 m) the depth of investigation here "
            "is about 30 m") in text
    assert "Schlumberger" not in text


def test_a_depth_cut_back_to_the_depth_of_investigation_reads_as_a_minimum():
    """A zone resolved to 38 m in a sounding that sees 40 m: the margin below
    it was cut to 40 m and the recommendation said "about 40 m"."""
    import numpy as np

    from groundwater.models import LayeredModel, VESSounding
    from groundwater.reporting.geophysical import _recommendations
    from groundwater.ves import interpret_model

    ab2 = np.array([1, 2, 5, 10, 20, 40, 80.0])
    sounding = VESSounding(SiteMetadata(), "S1", ab2, np.full(len(ab2), np.nan),
                           np.full(len(ab2), 100.0))
    interp = interpret_model(sounding, LayeredModel([1000, 100, 5000], [5, 33]))
    interp.rank = 1
    depth = next(item for item in _recommendations([interp])
                 if item.startswith("The drilling depth"))
    assert depth.startswith("The drilling depth should be at least 40 m at point S1")
    assert ("Where the depth is a minimum, the margin drilled below the deepest water "
            "zone runs past what the survey resolves") in depth
