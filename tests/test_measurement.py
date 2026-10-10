"""The value of one more measurement (PLAN.md step 3.5).

The worked examples are computed by hand in the comments beside them. The
one most of them use: the point's chance p = 0.6, the alternative's q = 0.5,
a completed borehole US$ 5,000 and a dry attempt US$ 2,000.

    going to the alternative  K = 5000 + 2000 x (1 - 0.5) / 0.5 = 7,000
    drilling here first       0.6 x 5000 + 0.4 x (2000 + 7000)   = 6,600

so drilling here is the cheaper choice now. A measurement with two readings,
factors 2 and 0.5, is spread among dry boreholes as 1/3 and 2/3 (the only
spread whose factors average 1: f + g = 1 and 2f + 0.5g = 1), and among
working ones as 2/3 and 1/3.

    reading "2":   chance 0.6 x 2/3 + 0.4 x 1/3 = 8/15, chance after 0.75
    reading "0.5": chance 0.6 x 1/3 + 0.4 x 2/3 = 7/15, chance after 3/7

After "0.5", drilling here costs 3/7 x 5000 + 4/7 x 9000 = 7,285.71, more
than the 7,000 of going to the alternative, so the choice changes and saves
285.71. The value is 7/15 x 285.71 = 133.33. Knowing for certain would save
the dry attempt whenever there would be one: 0.4 x 2000 = 800.
"""

from __future__ import annotations

import copy
import math
import random
from types import SimpleNamespace

import pytest

from groundwater.siting import measurement as measurement_module
from groundwater.siting import odds as odds_module
from groundwater.siting.measurement import (
    KINDS,
    even_spread,
    measurement_cost_usd,
    measurement_decision_rows,
    measurement_reading_rows,
    measurement_summary,
    measurement_text,
    measurement_value,
    measurement_values,
    preposterior,
)
from groundwater.siting.odds import odds_tables
from groundwater.text import phrase

TWO_READINGS = [[(2.0, 2 / 3, 1 / 3), (0.5, 1 / 3, 2 / 3)]]


def _point(p=0.6, prior=0.5, sid="P1", weight=1.0, sampled=False):
    """What measurement_value reads of a point's odds."""
    return SimpleNamespace(probability=p, prior=prior, sounding_id=sid,
                           range_sampled=sampled,
                           evidence=[SimpleNamespace(key="fit", weight=weight)])


@pytest.fixture
def two_ratios(monkeypatch):
    """The real tables with the resistivity's "no zone" at 0.5 and every
    band at 2: a profiling line then reads, in effect, the two readings of
    the worked example. The most even spread puts 2/3 on the one band at 0.5
    among dry boreholes and 1/15 on each of the five at 2 (a + 5b = 1 and
    0.5a + 10b = 1)."""
    tables = copy.deepcopy(odds_tables())
    resistivity = tables["evidence"]["resistivity"]
    resistivity["none"]["lr"] = 0.5
    for band in resistivity["bands"]:
        band["lr"] = 2.0
    monkeypatch.setattr(measurement_module, "odds_tables", lambda: tables)
    monkeypatch.setattr(odds_module, "odds_tables", lambda: tables)
    return tables


# ------------------------------------------------------------- by hand

def test_the_worked_example_by_hand():
    v = preposterior(0.6, 0.5, 5000.0, 2000.0, TWO_READINGS)
    assert v["alternative_usd"] == pytest.approx(7000.0)
    assert v["drill_here_usd"] == pytest.approx(6600.0)
    assert v["decision_now"] == "here"
    assert v["now_usd"] == pytest.approx(6600.0)
    assert v["keep_chance"] == pytest.approx(8 / 15)
    assert v["keep_probability"] == pytest.approx(0.75)
    # 0.75 x 5000 + 0.25 x 9000
    assert v["keep_usd"] == pytest.approx(6000.0)
    assert v["change_chance"] == pytest.approx(7 / 15)
    assert v["change_probability"] == pytest.approx(3 / 7)
    assert v["change_usd"] == pytest.approx(7000.0)
    assert v["value_usd"] == pytest.approx(400 / 3)
    # 8/15 x 6000 + 7/15 x 7000 = 6466.67 = 6600 - 133.33
    assert v["after_usd"] == pytest.approx(6600.0 - 400 / 3)
    assert v["perfect_usd"] == pytest.approx(800.0)
    # the choice changes where the factor takes the odds, 1.5, under 1
    assert v["threshold"] == pytest.approx(1 / 1.5)


def test_the_worked_example_from_the_other_side():
    """p = 0.4 against q = 0.5: going to the alternative is cheaper now.

        drilling here first   0.4 x 5000 + 0.6 x 9000 = 7,400 > 7,000
        reading "2":   chance 0.4 x 2/3 + 0.6 x 1/3 = 7/15, after 4/7
        reading "0.5": chance 0.4 x 1/3 + 0.6 x 2/3 = 8/15, after 1/4

    After "2" drilling here costs 4/7 x 5000 + 3/7 x 9000 = 6,714.29, saving
    285.71 on 7,000, so the value is 7/15 x 285.71 = 133.33 again. Knowing
    for certain saves p (K - 5000) = 0.4 x 2000 = 800."""
    v = preposterior(0.4, 0.5, 5000.0, 2000.0, TWO_READINGS)
    assert v["decision_now"] == "move"
    assert v["now_usd"] == pytest.approx(7000.0)
    assert v["change_chance"] == pytest.approx(7 / 15)
    assert v["change_probability"] == pytest.approx(4 / 7)
    assert v["change_usd"] == pytest.approx(5000 * 4 / 7 + 9000 * 3 / 7)
    assert v["value_usd"] == pytest.approx(400 / 3)
    assert v["perfect_usd"] == pytest.approx(800.0)


def test_the_value_has_a_closed_form_in_the_dry_attempt():
    """Drilling here first less going to the alternative is
    dry x (q - p) / q, so the value is dry / q times the expected shortfall
    of the chance after below q: the completed borehole's cost cancels."""
    for completed in (3000.0, 5000.0, 20000.0):
        v = preposterior(0.6, 0.5, completed, 2000.0, TWO_READINGS)
        # 2000 / 0.5 x 7/15 x (0.5 - 3/7)
        assert v["value_usd"] == pytest.approx(2000 / 0.5 * 7 / 15 * (0.5 - 3 / 7))


def test_the_spread_of_two_readings_is_the_only_one_their_ratios_allow():
    dry, wet = even_spread([2.0, 0.5])
    assert dry == pytest.approx([1 / 3, 2 / 3], abs=1e-15)
    assert wet == pytest.approx([2 / 3, 1 / 3], abs=1e-15)


def test_a_profiling_line_reads_the_worked_example(two_ratios):
    v = measurement_value(_point(), None, 5000.0, 2000.0, "profiling")
    # the alternative is an unsurveyed site at the point's prior, 0.5
    assert v.alternative_id is None and v.alternative_probability == 0.5
    assert [r.if_dry for r in v.readings] == pytest.approx([2 / 3] + [1 / 15] * 5)
    assert [r.if_success for r in v.readings] == pytest.approx([1 / 3] + [2 / 15] * 5)
    assert v.value_usd == pytest.approx(400 / 3)
    assert v.perfect_usd == pytest.approx(800.0)
    assert v.cost_usd == 150.0
    assert measurement_summary(v) == (
        "A profiling line through P1 is worth up to US$ 133 to this decision and "
        "costs about US$ 150. It costs more than it is worth to this decision.")


def test_the_fit_weight_tempers_the_readings(two_ratios):
    """Weighted 0.5, the factors are 2^0.5 and 0.5^0.5, and the spread is
    the one those allow: f + g = 1 and 1.41421 f + 0.70711 g = 1 give
    f = (1 - 0.70711) / (1.41421 - 0.70711) = 0.41421 on the five bands."""
    v = measurement_value(_point(weight=0.5), None, 5000.0, 2000.0, "profiling")
    assert v.readings[0].factor == pytest.approx(math.sqrt(0.5))
    assert sum(r.if_dry for r in v.readings[1:]) == pytest.approx(math.sqrt(2) - 1)
    assert v.weight == 0.5
    # weaker evidence is worth less to the same decision
    assert v.value_usd < 400 / 3


def test_a_tie_keeps_the_first_ranked_point():
    """At p = q = 0.5 the two choices cost the same, 7,000 each, and the tie
    keeps the first-ranked point. The reading "0.5" takes it to 1/3, so the
    choice changes on it and saves 2000 / 0.5 x (0.5 - 1/3) = 666.67, half
    the time: 333.33."""
    v = preposterior(0.5, 0.5, 5000.0, 2000.0, TWO_READINGS)
    assert v["drill_here_usd"] == v["alternative_usd"] == 7000.0
    assert v["decision_now"] == "here"
    assert v["change_chance"] == pytest.approx(0.5)
    assert v["change_probability"] == pytest.approx(1 / 3)
    assert v["value_usd"] == pytest.approx(1000 / 3)
    # and two points at the same odds: the first-ranked one is drilled first
    values = measurement_values([_point(0.5, sid="A"), _point(0.5, sid="B")], [0, 1],
                                5000.0, 2000.0)
    assert {v.decision_now for v in values} == {"here"}
    assert "Without it, drilling at A first is the cheaper choice" in \
        measurement_text(values[0])[1]


def test_the_even_spread_is_a_choice_not_a_bound():
    """Other spreads the ratios allow give other values, larger and smaller,
    so the page calls the even spread a choice and not a bound. With the
    profiling line's ratios (0.3 to 1.4), all the dry boreholes on the two
    extremes (4/11 at 0.3, 7/11 at 1.4, which average 1) read the most; all
    on 0.9 and 1.4 (4/5 and 1/5) read less than the even spread."""
    ratios = [r for _names, r in measurement_module._classes("profiling", False)[0][1]]
    assert (min(ratios), max(ratios)) == (0.3, 1.4)

    def value(dry):
        assert math.fsum(dry) == pytest.approx(1.0)
        assert math.fsum(f * r for f, r in zip(dry, ratios, strict=True)) == \
            pytest.approx(1.0)
        wet = [f * r for f, r in zip(dry, ratios, strict=True)]
        return preposterior(0.6, 0.55, 8000.0, 3500.0,
                            [list(zip(ratios, wet, dry, strict=True))])["value_usd"]

    dry, _wet = even_spread(ratios)
    even = value(dry)
    extremes = [4 / 11 if r == 0.3 else 7 / 11 if r == 1.4 else 0.0 for r in ratios]
    middle = [0.8 if r == 0.9 else 0.2 if r == 1.4 else 0.0 for r in ratios]
    assert value(extremes) > even > value(middle)
    assert "the even spread is a choice and not a bound" in \
        " ".join(phrase("measurement.basis", completed="1", dry="1", weight=1.0).split())


# ---------------------------------------------------------- properties

def _random_classes(rng, sampled):
    """The real tables' classes at a random fit weight."""
    weight = rng.choice([1.0, 0.8, 0.6, 0.3])
    kind = rng.choice(KINDS)
    classes = []
    for _evidence, bands in measurement_module._classes(kind, sampled):
        factors = [ratio ** weight for _names, ratio in bands]
        spread = even_spread(factors)
        if spread is not None:
            dry, wet = spread
            classes.append(list(zip(factors, wet, dry, strict=True)))
    return classes


def test_the_value_is_never_negative_nor_above_perfect_information():
    rng = random.Random(35)
    for _ in range(3000):
        p = rng.choice([rng.random(), 1e-6, 1 - 1e-6, rng.uniform(0.3, 0.7)])
        q = rng.choice([rng.random(), p, 1e-6, 1 - 1e-6])
        p, q = min(max(p, 1e-6), 1 - 1e-6), min(max(q, 1e-6), 1 - 1e-6)
        completed = rng.uniform(1000.0, 30000.0)
        dry = rng.uniform(100.0, 20000.0)
        classes = _random_classes(rng, rng.random() < 0.5)
        v = preposterior(p, q, completed, dry, classes)
        assert v["value_usd"] >= 0.0
        assert v["value_usd"] <= v["perfect_usd"] * (1 + 1e-9) + 1e-9
        # by the definition, worked out apart: every reading, its chance and
        # the chance after it, the best choice on each, and the best now
        readings = [(1.0, 1.0)]
        for bands in classes:
            readings = [(s0 * s, f0 * f) for s0, f0 in readings for _r, s, f in bands]
        chances = [p * s + (1 - p) * f for s, f in readings]
        afters = [p * s / c for (s, _f), c in zip(readings, chances, strict=True)]
        assert math.fsum(chances) == pytest.approx(1.0)
        # the chance after a reading averages back to the chance now
        assert math.fsum(c * a for c, a in zip(chances, afters, strict=True)) == \
            pytest.approx(p, rel=1e-9, abs=1e-12)
        move = completed + dry * (1 - q) / q

        def here(x, move=move, completed=completed, dry=dry):
            return x * completed + (1 - x) * (dry + move)

        best_after = math.fsum(c * min(here(a), move)
                               for c, a in zip(chances, afters, strict=True))
        assert v["value_usd"] == pytest.approx(min(here(p), move) - best_after,
                                               rel=1e-6, abs=1e-6 * move)
        # the expected cost with the measurement adds up from its groups
        assert v["keep_chance"] * v["keep_usd"] + v["change_chance"] * v["change_usd"] == \
            pytest.approx(v["after_usd"], rel=1e-9)


def test_the_value_is_zero_where_no_reading_changes_the_choice():
    """At 99.9 percent here against 10 percent there, the worst reading of
    the real tables (0.36 x 0.3 = 0.108) takes odds of 999 to about 108:
    still far above the alternative's 0.11."""
    point = _point(p=0.999, prior=0.1, sampled=True)
    for kind in KINDS:
        v = measurement_value(point, None, 5000.0, 2000.0, kind)
        assert v.change_chance == 0.0
        assert v.value_usd == 0.0
        assert v.after_usd == v.now_usd
        assert measurement_summary(v) == phrase(
            "measurement.zero", what=phrase(f"measurement.what_{kind}", sid="P1"),
            sid="P1", cost=f"{measurement_cost_usd(kind):,.0f}")
        assert len(measurement_text(v)) == 3
        # no row for readings that change it
        assert len(measurement_decision_rows(v)) == 5
    # and a measurement that reads nothing is worth nothing
    v = preposterior(0.6, 0.5, 5000.0, 2000.0, [])
    assert v["value_usd"] == 0.0 and v["change_chance"] == 0.0


def test_the_ratios_are_spread_so_they_average_one():
    for _evidence, bands in measurement_module._classes("sounding", True):
        dry, wet = even_spread([r for _n, r in bands])
        assert sum(dry) == pytest.approx(1.0) and sum(wet) == pytest.approx(1.0)
        for (_n, r), f, s in zip(bands, dry, wet, strict=True):
            assert s == pytest.approx(r * f)
            assert f > 0.0
    # all on one side of 1, or all 1, cannot be spread so: no information
    assert even_spread([1.0, 0.9, 0.8]) is None
    assert even_spread([1.0, 1.0]) is None
    assert even_spread([1.2, 1.5]) is None


def test_a_sounding_is_worth_at_least_a_profiling_line():
    """A sounding reads everything a profiling line does and the depth too,
    so it cannot be worth less (Blackwell); with no range sampled the depth
    is not read and the two are worth the same."""
    rng = random.Random(7)
    for _ in range(200):
        p, q = rng.uniform(0.05, 0.95), rng.uniform(0.05, 0.95)
        point = _point(p=p, prior=q, sampled=True)
        s = measurement_value(point, None, 6000.0, 2500.0, "sounding")
        line = measurement_value(point, None, 6000.0, 2500.0, "profiling")
        assert s.value_usd >= line.value_usd - 1e-9
        bare = _point(p=p, prior=q, sampled=False)
        assert measurement_value(bare, None, 6000.0, 2500.0, "sounding").value_usd == \
            pytest.approx(measurement_value(bare, None, 6000.0, 2500.0, "profiling").value_usd)


# ----------------------------------------------------------- edge cases

def test_no_survey_gives_nothing_to_price():
    assert measurement_values([], [], 5000.0, 2000.0) == []


def test_odds_at_the_extremes_stay_finite():
    for p in (1e-6, 1 - 1e-6):
        for q in (1e-6, 0.5, 1 - 1e-6):
            for kind in KINDS:
                v = measurement_value(_point(p=p, prior=q, sampled=True), None,
                                      5000.0, 2000.0, kind)
                for name in ("now_usd", "after_usd", "value_usd", "perfect_usd",
                             "alternative_usd", "threshold"):
                    assert math.isfinite(getattr(v, name)), (p, q, name)
                assert 0.0 <= v.value_usd <= v.perfect_usd * (1 + 1e-9)
                assert "nan" not in " ".join(measurement_text(v))


def test_impossible_inputs_are_refused():
    with pytest.raises(ValueError):
        preposterior(1.0, 0.5, 5000.0, 2000.0, TWO_READINGS)
    with pytest.raises(ValueError):
        preposterior(0.5, 0.0, 5000.0, 2000.0, TWO_READINGS)
    with pytest.raises(ValueError):
        preposterior(0.5, 0.5, 5000.0, 0.0, TWO_READINGS)
    with pytest.raises(ValueError):
        measurement_cost_usd("borehole")


def test_with_no_range_sampled_a_sounding_says_its_depth_is_left_out():
    v = measurement_value(_point(sampled=False), None, 5000.0, 2000.0, "sounding")
    assert {r.evidence for r in v.readings} == {"resistivity"}
    assert measurement_text(v)[-1] == phrase("measurement.not_sampled", sid="P1")
    sampled = measurement_value(_point(sampled=True), None, 5000.0, 2000.0, "sounding")
    assert {r.evidence for r in sampled.readings} == {"depth", "resistivity"}
    assert phrase("measurement.not_sampled", sid="P1") not in measurement_text(sampled)
    # nine depth readings (four bands, resolved or partly, and unresolved)
    # and six of the resistivity, each named
    rows = measurement_reading_rows(sampled)
    assert len(rows) == 15
    assert rows[4][1] == "15 to 35 m, favourable; basement resolved"
    assert rows[8][1] == "basement not resolved, no depth quoted"
    assert rows[12][1] == "50 to 300 ohm-m: productive weathered zone"


def test_the_alternative_is_the_other_point_with_the_best_odds():
    a, b, c = _point(0.6, sid="A"), _point(0.3, sid="B"), _point(0.45, sid="C")
    values = measurement_values([a, b, c], [0, 1, 2], 5000.0, 2000.0)
    assert [v.kind for v in values] == list(KINDS)
    assert {v.alternative_id for v in values} == {"C"}
    assert values[0].alternative_probability == 0.45
    # ranked otherwise, the first-ranked point is the one priced
    values = measurement_values([a, b, c], [2, 0, 1], 5000.0, 2000.0)
    assert values[0].sounding_id == "C" and values[0].alternative_id == "A"
    assert values[0].decision_now == "move"
    assert "Without it, going to point A is the cheaper choice" in \
        measurement_text(values[0])[1]


def test_the_costs_are_provisional_and_say_they_have_no_source():
    from groundwater._resources import bundled_text

    assert measurement_cost_usd("sounding") == 100.0
    assert measurement_cost_usd("profiling") == 150.0
    # the comment over the two figures, its lines joined
    text = " ".join(line.strip("# ") for line in bundled_text("field.yaml").splitlines())
    assert "NO SOURCE IS RECORDED for either figure" in text
    assert "provisional" in phrase("measurement.basis", completed="1", dry="1", weight=1.0)


# --------------------------------------------------------------- reports

def test_the_cost_report_prints_what_one_more_measurement_is_worth(tmp_path):
    from docx import Document

    from groundwater.costing import CostingInputs, estimate_borehole_cost, sample_cost
    from groundwater.reporting.costing import CostReportInputs, build_cost_report

    inputs = CostingInputs(total_depth_m=40.0, mobilisation_distance_km=100.0)
    spread = sample_cost(inputs, samples=500, success_probability=0.6, odds_source="P1")
    values = measurement_values([_point(), _point(0.4, sid="P2")], [0, 1],
                                spread.mean, spread.dry_mean)
    out = build_cost_report(CostReportInputs(
        estimate=estimate_borehole_cost(inputs), figures_dir=tmp_path,
        distribution=spread, measurements=values), tmp_path / "cost.docx")
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "4.2 The value of one more measurement" in text
    for v in values:
        assert " ".join(measurement_text(v)) in text
    cells = [[c.text for c in row.cells] for t in doc.tables for row in t.rows]
    assert measurement_decision_rows(values[0])[0] in cells
    # and with no survey, it says there is nothing to improve
    out = build_cost_report(CostReportInputs(
        estimate=estimate_borehole_cost(inputs), figures_dir=tmp_path,
        distribution=spread, measurements=[]), tmp_path / "bare.docx")
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert phrase("measurement.no_survey") in text


@pytest.mark.slow
def test_the_geophysical_report_prices_it_under_the_odds(tmp_path):
    from pathlib import Path

    from docx import Document

    from groundwater.costing import CostingInputs, sample_cost
    from groundwater.ingestion.ves import read_ves_workbook
    from groundwater.reporting import build_geophysical_report
    from groundwater.reporting.geophysical import GeophysicalReportInputs
    from groundwater.siting import assess_siting, survey_odds
    from groundwater.ves import interpret_model
    from groundwater.ves.inversion import invert_sounding

    soundings = read_ves_workbook(Path(__file__).resolve().parents[1] / "examples"
                                  / "data" / "rokel" / "rokel_ves.xlsx")
    inversions = [invert_sounding(s) for s in soundings]
    interps = [interpret_model(s, r.model) for s, r in zip(soundings, inversions, strict=True)]

    def build(distribution, name):
        path = build_geophysical_report(
            GeophysicalReportInputs(soundings=soundings, inversions=inversions,
                                    interpretations=interps, figures_dir=tmp_path,
                                    cost_distribution=distribution),
            tmp_path / name)
        return "\n".join(p.text for p in Document(str(path)).paragraphs)

    text = build(None, "bare.docx")
    assert text.index("The value of one more measurement") > \
        text.index("Chance of a working borehole")
    assert phrase("measurement.no_cost") in text
    spread = sample_cost(CostingInputs(total_depth_m=40.0), samples=500)
    text = build(spread, "priced.docx")
    site = soundings[0].site
    odds = survey_odds(interps, None, site.utm_zone, site.latlon)
    values = measurement_values(odds, [s.index for s in assess_siting(interps)],
                                spread.mean, spread.dry_mean)
    assert len(values) == 2
    for v in values:
        assert " ".join(measurement_text(v)) in text
