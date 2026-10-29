"""End to end tests of the Streamlit app using the bundled samples.

These drive the same flow a demo visitor uses: open a page, pick a
bundled sample, run the analysis, build the report. AppTest executes the
real app script. Only the page on screen runs (PLAN.md step 1.1), so each
test opens its page first, with ``goto``, as the sidebar navigation does.
"""

import json
from pathlib import Path

import pytest

streamlit = pytest.importorskip("streamlit")

from streamlit.testing.v1 import AppTest  # noqa: E402

from conftest import goto  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
APP = str(REPO / "app" / "streamlit_app.py")

# Each AppTest executes the whole app script, several times over: this module
# is a quarter of the suite's running time, so pull request CI runs it on one
# Python version rather than four.
pytestmark = pytest.mark.slow


@pytest.fixture(scope="module")
def app(sample_data):
    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    assert not at.exception, at.exception
    return at


def test_app_renders(app):
    assert app.title[0].value == "Groundwater Investigation Toolkit"
    # bundled samples offered on every data page
    for page, key in (("Geophysics (VES)", "sample_ves"),
                      ("Pumping test", "sample_pump"),
                      ("Water quality", "sample_wq"),
                      ("Borehole design", "sample_log")):
        goto(app, page)
        assert app.selectbox(key=key) is not None


def test_ves_flow_with_sample(app):
    goto(app, "Geophysics (VES)")
    app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
    app.run()
    assert not app.exception
    # the Port Loko copy-over error in the sample is surfaced as a warning
    warnings = " ".join(str(w.value) for w in app.warning)
    assert "Port Loko" in warnings

    app.button(key="run_ves").click()
    app.run()
    assert not app.exception
    assert "ves_results" in app.session_state
    soundings, results, interps = app.session_state["ves_results"]
    assert len(results) == 2 and results[0].fit_error_percent < 21.5

    # the range of models, beside each best fit, asked for with its button
    app.button(key="ves_range").click()
    app.run()
    assert not app.exception
    kept, ranges = app.session_state["ves_ranges"]
    assert kept is results and len(ranges) == 2
    shown = " ".join(str(i.value) for i in app.info)
    assert "of the models that fit" in shown and "Metropolis-Hastings" in shown

    # the chance of a working borehole beside the score, reading the range
    from groundwater.siting import odds_basis_text, odds_point_text, survey_odds

    site = soundings[0].site
    expected = survey_odds(interps, ranges, site.utm_zone, site.latlon)
    written = " ".join(str(m.value) for m in app.markdown)
    for o in expected:
        assert " ".join(odds_point_text(o)) in written
    assert "has not been sampled" not in written
    captions = " ".join(str(c.value) for c in app.caption)
    assert " ".join(odds_basis_text(expected[0])) in captions

    app.button(key="build_geo_report").click()
    app.run()
    assert not app.exception


def test_the_programme_offers_the_survey_odds_without_using_them(app):
    """The typed success rate stays until the odds are chosen with a click."""
    if "ves_results" not in app.session_state:
        goto(app, "Geophysics (VES)")
        app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
        app.run()
        app.button(key="run_ves").click()
        app.run()
    goto(app, "Costing & BoQ")
    assert not app.exception
    typed = app.number_input(key="cost_prog_success").value
    offered = " ".join(str(i.value) for i in app.info)
    assert "the first-ranked point, at about" in offered
    rate = programme_rate_of(app)
    assert app.button(key="use_survey_odds").label == f"Use {rate:g} percent"
    assert app.number_input(key="cost_prog_success").value == typed
    app.button(key="use_survey_odds").click()
    app.run()
    assert not app.exception
    assert app.number_input(key="cost_prog_success").value == rate


def test_the_cost_distribution_draws_from_the_survey(app):
    """With the survey inverted and its range sampled, the planning figure
    draws the depth from the first-ranked point's range and dry holes at its
    odds, and says so on the page (PLAN.md step 3.4)."""
    from groundwater.costing.distribution import cost_range_text
    from groundwater.siting import assess_siting, survey_odds

    if "ves_results" not in app.session_state:
        goto(app, "Geophysics (VES)")
        app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
        app.run()
        app.button(key="run_ves").click()
        app.run()
    soundings, results, interps = app.session_state["ves_results"]
    kept = app.session_state.get("ves_ranges")
    if kept is None or kept[0] is not results:
        goto(app, "Geophysics (VES)")
        app.button(key="ves_range").click()
        app.run()
    ranges = app.session_state["ves_ranges"][1]
    goto(app, "Costing & BoQ")
    app.button(key="run_cost").click()
    app.run()
    assert not app.exception
    spread = app.session_state["cost_spread"][1]
    i = assess_siting(interps)[0].index
    site = soundings[0].site
    odds = survey_odds([interps[i]], [ranges[i]], site.utm_zone, site.latlon)[0]
    assert spread.odds_source == interps[i].sounding_id
    assert spread.success_probability == odds.probability
    assert spread.depth_source == ranges[i].sounding_id
    assert spread.depth_p10_m <= spread.depth_p90_m <= ranges[i].investigation_depth_m
    assert spread.expected_per_working == (
        spread.mean + spread.dry_mean * ((1.0 - odds.probability) / odds.probability))
    shown = " ".join(str(m.value) for m in app.markdown)
    assert all(t in shown for t in cost_range_text(spread)[1:])
    # what one more measurement is worth, beside it (PLAN.md step 3.5):
    # priced from this distribution, at the first-ranked point, against the
    # other point
    from groundwater.siting import measurement_values
    from groundwater.siting.measurement import measurement_summary

    every = survey_odds(interps, ranges, site.utm_zone, site.latlon)
    values = measurement_values(every, [s.index for s in assess_siting(interps)],
                                spread.mean, spread.dry_mean)
    assert [v.sounding_id for v in values] == [interps[i].sounding_id] * 2
    assert all(v.alternative_id == interps[1 - i].sounding_id for v in values)
    assert all(f"**{measurement_summary(v)}**" in shown for v in values)
    # and the Geophysics page prices it too, now there is a cost
    goto(app, "Geophysics (VES)")
    assert not app.exception
    shown = " ".join(str(m.value) for m in app.markdown)
    assert all(f"**{measurement_summary(v)}**" in shown for v in values)


def programme_rate_of(app):
    """The rate the page offers, worked out as the page works it out."""
    from groundwater.siting import assess_siting, survey_odds
    from groundwater.siting.odds import programme_rate

    soundings, results, interps = app.session_state["ves_results"]
    kept = app.session_state.get("ves_ranges", None)
    ranges = kept[1] if kept is not None and kept[0] is results else [None] * len(results)
    i = assess_siting(interps)[0].index
    site = soundings[0].site
    return programme_rate(survey_odds([interps[i]], [ranges[i]], site.utm_zone,
                                      site.latlon)[0])


def test_two_points_with_one_id_each_show_their_own_odds(app):
    """A sheet copied without renumbering gives two points one id. The pages
    looked the odds up by id, so both rows showed the last point's odds and
    the first point's never appeared, and the Costing page offered the first
    point carrying the leader's id rather than the leader."""
    import copy

    from groundwater.siting import assess_siting, odds_point_text, odds_short, survey_odds
    from groundwater.siting.odds import programme_offer
    from groundwater.text import phrase

    if "ves_results" not in app.session_state:
        goto(app, "Geophysics (VES)")
        app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
        app.run()
        app.button(key="run_ves").click()
        app.run()
    held = app.session_state["ves_results"]
    kept = app.session_state.get("ves_ranges", None)
    soundings, results, interps = held
    ranges = kept[1] if kept is not None and kept[0] is results else [None] * len(results)
    # the second point copied over with the first's id, and listed first
    renamed = copy.copy(interps[1])
    renamed.sounding_id = interps[0].sounding_id
    points = [renamed, interps[0]]
    results2 = [results[1], results[0]]
    ranges2 = [ranges[1], ranges[0]]
    site = soundings[1].site
    expected = survey_odds(points, ranges2, site.utm_zone, site.latlon)
    assert odds_short(expected[0]) != odds_short(expected[1])
    # the leader, scored a point at a time: the one listed second
    lead = max((0, 1), key=lambda k: assess_siting([points[k]])[0].weighted)
    assert lead == 1
    try:
        app.session_state["ves_results"] = ([soundings[1], soundings[0]], results2, points)
        app.session_state["ves_ranges"] = (results2, ranges2)
        goto(app, "Geophysics (VES)")
        table = next(d.value for d in app.dataframe
                     if phrase("odds.col_short") in d.value.columns)
        assert table[phrase("odds.col_short")].tolist() == [
            odds_short(expected[lead]), odds_short(expected[1 - lead])]
        written = " ".join(str(m.value) for m in app.markdown)
        for o in expected:
            assert " ".join(odds_point_text(o)) in written
        goto(app, "Costing & BoQ")
        offered = " ".join(str(i.value) for i in app.info)
        assert programme_offer(expected[lead]) in offered
    finally:
        app.session_state["ves_results"] = held
        if kept is not None:
            app.session_state["ves_ranges"] = kept


def test_pumping_flow_with_sample(app):
    goto(app, "Pumping test")
    app.selectbox(key="sample_pump").select("dr_timbo/dr_timbo_constant_test.xlsx")
    app.run()
    assert not app.exception
    analysis = app.session_state["pump_analysis"]
    assert analysis.transmissivity_m2_per_day is not None
    app.button(key="build_pump_report").click()
    app.run()
    assert not app.exception


def test_pending_pumping_sample(app):
    goto(app, "Pumping test")
    app.selectbox(key="sample_pump").select("kuntolo/kuntolo_step_test.xlsx")
    app.run()
    assert not app.exception
    analysis = app.session_state["pump_analysis"]
    assert analysis.transmissivity_m2_per_day is None  # discharge pending


def test_quality_flow_with_sample(app):
    goto(app, "Water quality")
    app.selectbox(key="sample_wq").select("dr_timbo/dr_timbo_water_quality.xlsx")
    app.run()
    assert not app.exception
    errors = " ".join(str(e.value) for e in app.error)
    assert "Manganese" in errors  # health exceedance verdict
    app.button(key="build_wq_report").click()
    app.run()
    assert not app.exception


def test_a_report_built_is_issued_and_checked_on_the_verify_page(app):
    """PLAN.md step 8.1. The water quality report, built on its page, keeps
    its issue record in the session (which the project file saves) and is
    issued from the sample workbook it was read from, by that file's hash.
    On the Verify a document page it is verified against the session's
    record; with one digit changed it is not, and the part is named; and its
    code, typed in lower case, is found."""
    import hashlib
    import io
    import re
    import zipfile

    import yaml

    from groundwater.project_io import serialize_project
    from groundwater.verification import document_record, verify_document

    goto(app, "Water quality")
    app.selectbox(key="sample_wq").select("dr_timbo/dr_timbo_water_quality.xlsx")
    app.run()
    app.button(key="build_wq_report").click()
    app.run()
    assert not app.exception, app.exception
    path = Path(app.session_state["artifacts"]["Download water quality report (.docx)"])
    data = path.read_bytes()
    record = document_record(data)
    assert record is not None and record["kind"] == "quality"
    assert record in app.session_state["issued_documents"]
    assert any(c.value == f"Document code {record['code']}" for c in app.caption)
    sample = (REPO / "examples" / "data" / "dr_timbo" / "dr_timbo_water_quality.xlsx").read_bytes()
    evidence = verify_document(data)["facts"]["evidence"]
    assert {"what": "water quality sheet", "name": "dr_timbo_water_quality.xlsx",
            "sha256": hashlib.sha256(sample).hexdigest()} in evidence
    saved = yaml.safe_load(serialize_project(
        {"issued_documents": app.session_state["issued_documents"]}, "0.3.0"))
    assert record in saved["issued"]

    goto(app, "Verify a document")
    app.file_uploader(key="verify_document").set_value(
        ("report.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    app.run()
    assert not app.exception, app.exception
    assert any("Verified" in s.value for s in app.success)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    text = entries["word/document.xml"]
    found = re.search(rb"<w:t(?: [^>]*)?>[^<]*?([0-9])", text)
    at = found.start(1)
    entries["word/document.xml"] = text[:at] + (b"8" if text[at:at + 1] == b"9" else b"9") + text[at + 1:]
    changed = io.BytesIO()
    with zipfile.ZipFile(changed, "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in entries.items():
            z.writestr(name, body)
    app.file_uploader(key="verify_document").set_value(
        ("changed.docx", changed.getvalue(),
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    app.run()
    assert any("Changed since it was issued" in e.value for e in app.error)
    assert any("word/document.xml (the text of the document)" in m.value for m in app.markdown)
    app.text_input(key="verify_code").input(record["code"].lower()).run()
    assert any(f"Document {record['code']} was issued" in s.value for s in app.success)


def test_the_evidence_bundle_is_built_on_the_handover_page_and_checked(app):
    """PLAN.md step 8.3. The water quality report, built on its page, is
    listed on the Handover page's Evidence bundle section; a file that is not
    a sealed document, added there, is refused in the catalogue's words;
    without it the bundle carries the report, the sheet it was read from by
    that file's hash and the session's issue records, and the Verify a
    document page verifies it against them. One byte changed in the sheet
    fails the check and names the file."""
    import hashlib

    from groundwater.bundle import _stored_zip, verify_evidence_bundle
    from groundwater.text import phrase_table
    from groundwater.verification import read_package

    goto(app, "Water quality")
    app.selectbox(key="sample_wq").select("dr_timbo/dr_timbo_water_quality.xlsx")
    app.run()
    app.button(key="build_wq_report").click()
    app.run()
    assert not app.exception, app.exception
    report = Path(app.session_state["artifacts"]["Download water quality report (.docx)"])

    goto(app, "Handover")
    assert any(report.name in str(df.value) for df in app.dataframe)
    app.file_uploader(key="bundle_documents").set_value(("notes.docx", b"not a document", "application/octet-stream"))
    app.run()
    app.button(key="build_bundle").click()
    app.run()
    assert not app.exception, app.exception
    assert any(e.value.startswith("notes.docx cannot go into an evidence bundle") for e in app.error)
    app.file_uploader(key="bundle_documents").set_value([])
    app.run()
    app.button(key="build_bundle").click()
    app.run()
    assert not app.exception, app.exception
    path = Path(app.session_state["artifacts"]["Download evidence bundle (.zip)"])
    data = path.read_bytes()
    records = app.session_state["issued_documents"]
    checked = verify_evidence_bundle(data, records)
    assert checked["verdict"] == "verified", checked["lines"]
    facts = checked["facts"]
    assert f"reports/{report.name}" in [d["path"] for d in facts["documents"]]
    sample = (REPO / "examples" / "data" / "dr_timbo" / "dr_timbo_water_quality.xlsx").read_bytes()
    sheet = "raw/dr_timbo_water_quality.xlsx"
    assert {"path": sheet, "what": "water quality sheet", "bytes": len(sample),
            "sha256": hashlib.sha256(sample).hexdigest()} in facts["files"]
    assert all(r in facts["issued"] for r in records)

    goto(app, "Verify a document")
    app.file_uploader(key="verify_bundle").set_value((path.name, data, "application/zip"))
    app.run()
    assert not app.exception, app.exception
    assert any(phrase_table("bundle.verdicts")["verified"] in s.value for s in app.success)
    parts = read_package(data)
    parts[sheet] = bytes([parts[sheet][0] ^ 1]) + parts[sheet][1:]
    app.file_uploader(key="verify_bundle").set_value(("changed.zip", _stored_zip(sorted(parts.items())),
                                                      "application/zip"))
    app.run()
    assert any(phrase_table("bundle.verdicts")["changed"] in e.value for e in app.error)
    assert any(f"Changed since the bundle was made: {sheet}." in m.value for m in app.markdown)


def test_design_flow_with_sample(app):
    goto(app, "Borehole design")
    app.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    app.run()
    assert not app.exception


def test_the_design_page_designs_against_the_pumping_test(app):
    """The design this page hands on is the one the completion report carries.

    It used to ask for a static water level with nothing in the box and pass
    no pump intake at all, so the intake checks never ran here and the two
    pages disagreed about the same borehole.
    """
    goto(app, "Pumping test")
    app.selectbox(key="sample_pump").select("dr_timbo/dr_timbo_constant_test.xlsx")
    app.run()
    goto(app, "Borehole design")
    app.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    app.run()
    assert not app.exception

    analysis = app.session_state["pump_analysis"]
    design = app.session_state["borehole_design"]
    assert design.static_water_level_m == pytest.approx(
        analysis.test.static_water_level_m, abs=0.01
    )
    # the intake the test supports, not a level nobody typed
    assert design.pump_intake_m is not None
    assert design.pump_intake_m > analysis.test.static_water_level_m


def test_costing_flow(app):
    goto(app, "Costing & BoQ")
    app.button(key="run_cost").click()
    app.run()
    assert not app.exception
    estimate = app.session_state["cost_estimate"]
    assert estimate.direct_cost_usd > 0
    assert estimate.items
    # the planning figure is sampled with the estimate and shown beside it,
    # saying which figure is the contract document
    kept, spread, chart = app.session_state["cost_spread"][:3]
    assert kept is estimate and chart.exists()
    assert spread.boq_usd == pytest.approx(estimate.price_with_vat_usd)
    from groundwater.text import phrase
    assert phrase("cost_range.which_is_which") in [str(i.value) for i in app.info]
    # with no survey interpreted, the value of one more measurement says
    # there are no odds to improve
    if "ves_results" not in app.session_state:
        assert phrase("measurement.no_survey") in [str(i.value) for i in app.info]
    app.button(key="build_cost_report").click()
    app.run()
    assert not app.exception


def test_no_water_no_pay_is_priced_on_dry_attempts(app):
    """The warning priced every expected failure as a whole completed
    borehole; it now carries a dry attempt for each, and says what that is."""
    from groundwater.utils import fmt_usd

    goto(app, "Costing & BoQ")
    app.number_input(key="cost_success").set_value(80.0)
    app.button(key="run_cost").click()
    app.run()
    assert not app.exception
    estimate = app.session_state["cost_estimate"]
    loaded = estimate.price_per_successful_well_usd(80.0)
    assert loaded == pytest.approx(estimate.price_with_vat_usd
                                   + estimate.dry_attempt_price_usd / 4)
    said = [str(w.value) for w in app.warning if "No water no pay" in str(w.value)]
    assert said and f"${fmt_usd(loaded)}" in said[0]
    assert f"(${fmt_usd(estimate.dry_attempt_price_usd)})" in said[0]
    assert "siting, set-up and drilling" in said[0]
    app.number_input(key="cost_success").set_value(100.0)
    app.run()


def test_supervision_flow(app):
    from groundwater.supervision import load_checklists

    first = load_checklists()[0]
    goto(app, "Supervision")
    # the radio is keyed chkw_ and writes through to the chk_ answer store,
    # so an answer survives a stage the run does not draw
    app.radio(key=f"chkw_{first.item_id}").set_value("Yes")
    app.run()
    assert not app.exception
    assert app.session_state[f"chk_{first.item_id}"] == "Yes"
    app.button(key="build_sup_report").click()
    app.run()
    assert not app.exception


def test_a_field_check_nobody_measured_is_not_in_the_record(app):
    """The boxes started at plausible values and all four checks went into
    the report every run, so a record built without opening the expander
    said verticality, sand content, specific capacity and the drilled
    metres had passed, against numbers nobody measured."""
    from docx import Document

    goto(app, "Supervision")
    for key in ("fx_watercol", "fx_casingid", "fx_dev", "fx_vdepth", "fx_sand1",
                "fx_sand2", "fx_sand3", "fx_q", "fx_dd", "fx_logs_m", "fx_invoiced_m"):
        assert app.number_input(key=key).value is None, key
    app.button(key="build_sup_report").click()
    app.run()
    assert not app.exception
    assert app.session_state["field_checks"] == []
    path = Path(app.session_state["workdir"]) / "Supervision_Checklist_Report.docx"
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Field Acceptance Checks" not in text
    # what is entered is checked, and only that
    app.number_input(key="fx_dev").set_value(80.0)
    app.number_input(key="fx_vdepth").set_value(60.0)
    app.number_input(key="fx_casingid").set_value(103.0)
    app.number_input(key="fx_sand2").set_value(0.1)
    app.run()
    assert not app.exception
    assert [c.name for c in app.session_state["field_checks"]] == [
        "Verticality (plumb test)", "Sand content"]
    assert "1 sample" in app.session_state["field_checks"][1].measured
    for key in ("fx_dev", "fx_vdepth", "fx_casingid", "fx_sand2"):
        app.number_input(key=key).set_value(None)
    app.run()
    assert not app.exception


def test_degrees_typed_into_the_utm_boxes_are_refused_until_converted(app):
    """The gate, the identifier and the sidebar say the same thing of them,
    and converting them is the operator's own button."""
    from groundwater.geo import degrees_in_utm_boxes

    saved = {key: app.session_state.get(key, None)
             for key in ("meta_easting", "meta_northing", "meta_zone")}
    try:
        goto(app, "Asset registry")
        app.number_input(key="meta_easting").set_value(13.2294)
        app.number_input(key="meta_northing").set_value(8.4831)
        app.run()
        assert not app.exception
        sentence = degrees_in_utm_boxes(13.2294, 8.4831)
        warnings = [str(w.value) for w in app.warning]
        assert sentence in warnings                     # the sidebar
        assert sentence + (" The sidebar's \"Convert these degrees to UTM\" "
                           "button does it.") in warnings
        app.button(key="convert_degree_boxes").click()
        app.run()
        assert not app.exception
        assert app.session_state["meta_zone"] == "28N"
        assert abs(app.session_state["meta_easting"] - 694912) <= 1
        assert abs(app.session_state["meta_northing"] - 938150) <= 1
        # converted on a sign the parser supplied, and said
        assert "13.2294 was read as 13.2294 W" in app.session_state["latlon_assumed"]
        assert all(sentence not in str(w.value) for w in app.warning)
    finally:
        for key in ("meta_easting", "meta_northing"):
            app.number_input(key=key).set_value(saved[key] or 0.0)
        app.selectbox(key="meta_zone").select(saved["meta_zone"] or "29N")
        app.text_input(key="latlon_paste").set_value("")
        app.session_state["latlon_assumed"] = ""
        app.run()


def test_supervision_photo_evidence_holds_the_record_until_attached(app):
    """The gate lists each item missing its photograph, in the catalogue's
    words, and a photograph kept with its provenance lifts it."""
    import base64

    from groundwater.photos import photo_provenance
    from groundwater.supervision import load_checklists
    from groundwater.text import phrase

    wanted = [i for i in load_checklists() if i.photo_required]
    app.session_state["sup_evidence"] = {}
    goto(app, "Supervision")
    app.selectbox(key="sup_stage").select("design")
    app.run()
    assert not app.exception
    shown = " ".join(str(m.value) for m in app.markdown)
    assert "**Photo evidence**" in shown
    assert phrase("evidence.photo_missing", stage="Design and installation",
                  item=wanted[0].text) in shown
    # what the page's uploader keeps: the file and the record made of it
    data = (Path(__file__).resolve().parent / "webapp" / "fixtures"
            / "photo_exif.jpg").read_bytes()
    app.session_state["sup_evidence"] = {item.item_id: {
        "name": "photo.jpg", "mime": "image/jpeg",
        "b64": base64.b64encode(data).decode("ascii"),
        "provenance": photo_provenance(data, "2026-10-03T09:00:00Z",
                                       position_note="unsupported"),
    } for item in wanted}
    app.run()
    assert not app.exception
    shown = " ".join(str(m.value) for m in app.markdown)
    assert "**Photo evidence**" not in shown
    captions = " ".join(str(c.value) for c in app.caption)
    assert "SHA-256 " in captions and "camera clock" in captions
    app.session_state["sup_evidence"] = {}


def test_programme_flow(app):
    goto(app, "Costing & BoQ")
    app.button(key="run_programme").click()
    app.run()
    assert not app.exception
    programme, gantt_path, spread, spread_path = app.session_state["programme_estimate"][:4]
    assert programme.n_attempted >= programme.n_successful
    assert gantt_path.exists()
    assert spread.estimate_usd == pytest.approx(programme.price_with_vat_usd)
    assert spread.n_boreholes == programme.n_successful and spread_path.exists()


def test_the_programme_figure_follows_the_range_as_the_borehole_figure_does(app):
    """A range sampled, or dropped, after the programme was estimated moves
    the depth the borehole's planning figure is drawn at; the programme's is
    drawn again at the same depth rather than left at the old one."""
    if "ves_results" not in app.session_state:
        goto(app, "Geophysics (VES)")
        app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
        app.run()
        app.button(key="run_ves").click()
        app.run()
    results = app.session_state["ves_results"][1]
    kept = app.session_state.get("ves_ranges")
    if kept is None or kept[0] is not results:
        goto(app, "Geophysics (VES)")
        app.button(key="ves_range").click()
        app.run()
    ranges = app.session_state["ves_ranges"]
    goto(app, "Costing & BoQ")
    app.button(key="run_programme").click()
    app.run()
    assert not app.exception
    drawn = app.session_state["programme_estimate"][2]
    assert drawn.depth_source is not None
    try:
        del app.session_state["ves_ranges"]
        app.run()
        assert not app.exception
        held = app.session_state["programme_estimate"][2]
        assert held.depth_source is None and held is not drawn
    finally:
        app.session_state["ves_ranges"] = ranges
    app.run()
    assert app.session_state["programme_estimate"][2].depth_source == drawn.depth_source


def test_handover_flow(app):
    goto(app, "Handover")
    app.button(key="build_handover").click()
    app.run()
    assert not app.exception


def test_maps_flow(app):
    """With no position and no area recorded, the national maps still draw."""
    goto(app, "Site maps")
    app.button(key="run_maps").click()
    app.run()
    assert not app.exception
    paths = app.session_state["map_paths"]
    assert len(paths) >= 3 and all(p.exists() for p in paths)
    # a study area needs somewhere to centre on, and the sidebar has
    # recorded neither a position nor a district, so it is skipped rather
    # than drawn over the whole country
    assert not any("study_area" in p.name for p in paths)


def test_a_recorded_district_is_enough_for_a_study_area_map(app):
    """No GPS fix is not no map: the district is still an area worth drawing."""
    app.session_state["meta_district"] = "Port Loko"
    app.session_state["meta_community"] = "Kuntoloh"
    goto(app, "Site maps")
    app.button(key="run_maps").click()
    app.run()
    assert not app.exception
    paths = app.session_state["map_paths"]
    study = [p for p in paths if "study_area" in p.name]
    assert study and study[0].exists()


def test_subsurface_maps_flow(app):
    """The maps built from the survey's own soundings, not from a dataset."""
    goto(app, "Geophysics (VES)")
    app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
    app.run()
    assert not app.exception
    goto(app, "Site maps")
    app.button(key="run_subsurface").click()
    app.run()
    assert not app.exception
    made = [Path(p) for p in app.session_state.get("subsurface_paths", [])]
    assert made and all(p.exists() for p in made)
    # Rokel has two soundings, so the interpolated surfaces cannot be
    # built - they need three - and the page says so rather than failing.
    # What two soundings do support is the survey point map and the
    # pseudo-section of their readings.
    assert any("site_location" in p.name for p in made)
    assert Path(app.session_state["pseudosection_path"]).exists()
    traverse = app.session_state["traverse"]
    assert {"length_m", "bearing_deg", "max_offset_m"} <= set(traverse)
    # and the two Rokel points really are 20 km apart, which is what the
    # consistency checker flags. A section joining two soundings with
    # nothing measured between them is refused, and the page says why
    # rather than drawing a horizon across two chiefdoms.
    assert traverse["length_m"] > 20_000
    assert "section_path" not in app.session_state
    captions = " ".join(str(c.value) for c in app.caption)
    assert "No geoelectric section" in captions
    assert "no measurement between" in captions


def test_the_topography_tab_draws_a_map_from_geotiff_tiles(sample_data):
    """A GeoTIFF was refused with a gdal_translate command to run first, and
    one file was all the tab took. Tiles uploaded together are now opened as
    they arrive, a file that cannot be read is refused beside its name, and
    the map is drawn from one window read across the tiles, with the read
    report and the Copernicus notice under it."""
    from groundwater.text import phrase

    dem = Path(__file__).resolve().parent / "webapp" / "fixtures" / "dem"
    names = ["Copernicus_DSM_COG_10_N08_00_W012_00_DEM.tif",
             "Copernicus_DSM_COG_10_N08_00_W011_00_DEM.tif"]
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Site maps"
    # a GPS fix on the two tiles' shared edge, 8.995 N, 11.0 W
    at.session_state["meta_easting"] = 280126.0
    at.session_state["meta_northing"] = 994900.0
    at.session_state["meta_zone"] = "29N"
    at.run()
    assert not at.exception, at.exception
    captions = [str(c.value) for c in at.caption]
    assert phrase("terrain.python_only") in captions and phrase("terrain.not_in_reports") in captions
    at.slider(key="map_radius").set_value(5)

    # a projected copy is refused by name, and the map cannot be asked for
    at.file_uploader(key="dem_upload").set_value(
        [("refuse_epsg_3857.tif", (dem / "refuse_epsg_3857.tif").read_bytes(), "image/tiff")])
    at.run()
    assert not at.exception, at.exception
    assert any("refuse_epsg_3857.tif" in str(e.value) for e in at.error)
    assert at.button(key="run_topo").disabled
    assert at.button(key="run_terrain").disabled

    at.file_uploader(key="dem_upload").set_value(
        [(n, (dem / n).read_bytes(), "image/tiff") for n in names])
    at.run()
    assert not at.exception and not at.error, [e.value for e in at.error]
    at.button(key="run_topo").click()
    at.run()
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]
    assert Path(at.session_state["topo_path"]).stat().st_size > 30_000
    report = at.session_state["topo_report"]
    assert report["notes"][0].startswith(f"Read 2 files ({names[0]} and {names[1]})")
    assert report["notice"] == phrase("terrain.copernicus_liability")
    captions = [str(c.value) for c in at.caption]
    assert all(note in captions for note in report["notes"]) and report["notice"] in captions
    # the tiles are kept on disk, under the names they were uploaded with
    kept = at.session_state["dem_upload_kept"]["files"]
    assert [name for name, _ in kept] == names
    assert all(Path(path).is_file() for _, path in kept)


def test_a_topographic_map_is_not_shown_beside_files_it_was_not_drawn_from(sample_data):
    """The topographic map was cleared only when a draw failed, so the map
    of earlier files stayed on the page, with their credit, read report and
    Copernicus notice, under a refusal of the files uploaded since, after
    another source was declared, and after the uploader was cleared. It is
    shown only beside the files and the declared source it was drawn from,
    and it still outlives a rerun and a visit to another page, where the
    uploader hands its files back under the same ids."""
    from groundwater.text import phrase

    dem = Path(__file__).resolve().parent / "webapp" / "fixtures" / "dem"
    pair = [(n, (dem / n).read_bytes(), "image/tiff")
            for n in ("Copernicus_DSM_COG_10_N08_00_W012_00_DEM.tif",
                      "Copernicus_DSM_COG_10_N08_00_W011_00_DEM.tif")]
    notice = phrase("terrain.copernicus_liability")
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Site maps"
    # a GPS fix on the two tiles' shared edge, 8.995 N, 11.0 W
    at.session_state["meta_easting"] = 280126.0
    at.session_state["meta_northing"] = 994900.0
    at.session_state["meta_zone"] = "29N"
    at.run()
    at.slider(key="map_radius").set_value(5)

    def draw():
        at.file_uploader(key="dem_upload").set_value(pair)
        at.run()
        at.button(key="run_topo").click()
        at.run()
        assert not at.exception, at.exception
        assert not at.error, [e.value for e in at.error]

    def shown() -> bool:
        """Whether the map is on the page, its credit and notice with it."""
        assert not at.exception, at.exception
        captions = [str(c.value) for c in at.caption]
        credited = any(c.startswith("Elevation from") for c in captions)
        assert credited == (notice in captions) == ("topo_path" in at.session_state)
        return credited

    draw()
    assert shown()
    at.run()
    assert shown(), "nothing changed"
    goto(at, "Overview")
    goto(at, "Site maps")
    assert any("uploaded earlier" in str(c.value) for c in at.caption)
    assert shown(), "the same files, handed back by the uploader"

    at.file_uploader(key="dem_upload").set_value(
        [("refuse_epsg_3857.tif", (dem / "refuse_epsg_3857.tif").read_bytes(), "image/tiff")])
    at.run()
    assert any("refuse_epsg_3857.tif" in str(e.value) for e in at.error)
    assert not shown(), "other files, refused as they arrived"

    draw()
    at.selectbox(key="terrain_source").select("other")
    at.run()
    assert not shown(), "the same files, declared as from another source"

    at.selectbox(key="terrain_source").select("detect")
    at.run()
    draw()
    at.file_uploader(key="dem_upload").set_value([])
    at.run()
    assert not shown(), "no files at all"


def _terrain_page(files, fix=(280126.0, 994900.0)) -> AppTest:
    """The Site maps page with ``files`` in the elevation model's uploader
    (as ``set_value`` takes them), the terrain radius at 1 km and the
    drainage threshold at 0.05 km², and the GPS fix ``fix`` in zone 29N, by
    default on the two Copernicus fixture tiles' shared edge, 8.995 N,
    11.0 W; None records no location."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Site maps"
    if fix is not None:
        at.session_state["meta_easting"], at.session_state["meta_northing"] = fix
        at.session_state["meta_zone"] = "29N"
    at.run()
    assert not at.exception, at.exception
    at.file_uploader(key="dem_upload").set_value(files)
    at.run()
    at.number_input(key="terrain_radius_km").set_value(1.0)
    at.number_input(key="terrain_area_km2").set_value(0.05)
    at.run()
    return at


def _tiles(*names: str) -> list:
    """Fixture tiles from tests/webapp/fixtures/dem as the uploader takes
    them; by default the two Copernicus tiles either side of 11 W, each 40
    cells, about 1.2 km, wide."""
    dem = Path(__file__).resolve().parent / "webapp" / "fixtures" / "dem"
    names = names or ("Copernicus_DSM_COG_10_N08_00_W012_00_DEM.tif",
                      "Copernicus_DSM_COG_10_N08_00_W011_00_DEM.tif")
    return [(n, (dem / n).read_bytes(), "image/tiff") for n in names]


def _analyse(at: AppTest) -> dict:
    """Press "Analyse the terrain", check the run was neither refused nor
    failed, and give the result the page keeps."""
    at.button(key="run_terrain").click()
    at.run()
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]
    return at.session_state["terrain_result"]


def _analysis_kept(at: AppTest) -> list[str]:
    """Where the session holds an array, the analysis, its core, the read or
    an elevation grid, searched through every value's dicts, lists, tuples
    and sets: the page keeps the files and the sentences of an analysis,
    never the analysis itself."""
    import numpy as np

    from groundwater.mapping.elevation import ElevationGrid, ElevationRead
    from groundwater.mapping.terrain_analysis import TerrainAnalysis, TerrainCore

    found, seen = [], set()

    def walk(value, where: str) -> None:
        if id(value) in seen:
            return
        seen.add(id(value))
        if isinstance(value, (np.ndarray, TerrainAnalysis, TerrainCore, ElevationRead, ElevationGrid)):
            found.append(where)
        elif isinstance(value, dict):
            for key, item in value.items():
                walk(item, f"{where}[{key!r}]")
        elif isinstance(value, (list, tuple, set)):
            for n, item in enumerate(value):
                walk(item, f"{where}[{n}]")

    for key, value in at.session_state.items():
        walk(value, key)
    return found


def test_the_topography_tab_analyses_a_geotiff_and_offers_its_downloads(sample_data):
    """The Topography tab drew a map of the tiles and nothing more. "Analyse
    the terrain" reads the window with its margin, analyses it, finds its
    lineament candidates, lays the survey plan across them, and draws the
    four terrain figures, the candidates' map (and their rose, where there
    is a candidate) and the plan, each with the engine's captions under it
    and its PNG and SVG offered, the plan's soundings in a table, the
    drainage lines and the candidates as GeoJSON, the plan as GPX and a
    GeoLibre project, the read report, and the Copernicus liability
    sentence under the downloads. The session keeps the files and the
    sentences, not the analysis, and the radius and the threshold are
    carried through a visit to another page. A radius over the cap is
    refused, and the figures of the square before go with it."""
    import json
    import pickle
    import re
    import xml.etree.ElementTree as ET

    from streamlit.runtime.forward_msg_queue import ForwardMsgQueue

    from groundwater.text import phrase

    dem = Path(__file__).resolve().parent / "webapp" / "fixtures" / "dem"
    pair = [(n, (dem / n).read_bytes(), "image/tiff")
            for n in ("Copernicus_DSM_COG_10_N08_00_W012_00_DEM.tif",
                      "Copernicus_DSM_COG_10_N08_00_W011_00_DEM.tif")]
    notice = phrase("terrain.copernicus_liability")
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Site maps"
    # a GPS fix on the two tiles' shared edge, 8.995 N, 11.0 W
    at.session_state["meta_easting"] = 280126.0
    at.session_state["meta_northing"] = 994900.0
    at.session_state["meta_zone"] = "29N"
    at.run()
    assert not at.exception, at.exception
    at.file_uploader(key="dem_upload").set_value(pair)
    at.run()
    # the tiles are 40 cells, about 1.2 km, either side of 11 W
    at.number_input(key="terrain_radius_km").set_value(1.0)
    at.number_input(key="terrain_area_km2").set_value(0.05)
    at.run()
    sent = []
    ForwardMsgQueue.on_before_enqueue_msg(sent.append)
    try:
        at.button(key="run_terrain").click()
        at.run()
    finally:
        ForwardMsgQueue.on_before_enqueue_msg(None)
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]
    # the run's status opens as it starts, so the bar in it is seen while
    # it runs, and not only by whoever opens it
    opened = next(m for m in sent if m.delta.add_block.expandable.label == phrase("terrain.stage_read"))
    assert opened.delta.add_block.expandable.expanded
    assert any(m.delta.new_element.WhichOneof("type") == "progress"
               and m.metadata.delta_path[:-1] == opened.metadata.delta_path for m in sent)

    result = at.session_state["terrain_result"]
    assert set(result) == {"figures", "geojson", "lineament_geojson", "folder", "summary", "captions", "notes",
                           "notice", "inputs", "plan_gpx", "plan_geolibre", "plan_rows", "desk"}
    assert not _analysis_kept(at), "the analysis's layers are dropped"
    # what the desk study says of the terrain is kept beside the figures
    assert result["desk"]["candidates"] == 1 and len(result["desk"]["soundings"]) == 3
    assert len(pickle.dumps(result)) < 200_000, "the result holds paths and sentences alone"
    candidates = json.loads(Path(result["lineament_geojson"]).read_text(encoding="utf-8"))
    # the tiles hold one valley within 500 m of the fix: one line across it
    # and three soundings, a GPX waypoint each, and the line a route
    assert len(candidates["features"]) == 1
    keys = ["relief", "wetness", "curvature", "directions", "lineaments", "rose", "plan"]
    assert [key for key, _, _ in result["figures"]] == keys
    gpx = ET.fromstring(Path(result["plan_gpx"]).read_bytes())
    ns = {"g": "http://www.topografix.com/GPX/1/1"}
    assert [w.findtext("g:name", namespaces=ns) for w in gpx.findall("g:wpt", ns)] == [
        "Site", "VES 1", "VES 2", "VES 3", "Line 1 start", "Line 1 end"]
    assert len(gpx.findall("g:rte/g:rtept", ns)) == 41
    project = json.loads(Path(result["plan_geolibre"]).read_text(encoding="utf-8"))
    assert [layer["name"] for layer in project["layers"]] == [
        "Lineament candidates", "Service radius", "Profiling lines", "Profile stations", "Proposed soundings", "Site"]
    assert [row["Sounding"] for row in result["plan_rows"]] == ["VES 1", "VES 2", "VES 3"]
    table = next(str(m.value) for m in at.markdown if str(m.value).startswith("| Sounding |"))
    assert [line.split(" | ")[0].lstrip("| ") for line in table.splitlines()[2:]] == ["VES 1", "VES 2", "VES 3"]
    assert "not confirmed by geophysics" in result["captions"]["plan_note"]
    assert "No water points were loaded" in result["captions"]["plan"]
    # the run's files are in a folder of its own, in the session's
    assert Path(result["folder"]).parent == at.session_state["workdir"]
    names = []
    for key, png, svg in result["figures"]:
        assert Path(png) == Path(result["folder"]) / f"terrain_{key}.png"
        assert Path(svg) == Path(png).with_suffix(".svg")
        assert Path(png).stat().st_size > 20_000 and Path(svg).stat().st_size > 20_000
        names += [Path(png).name, Path(svg).name]
    assert Path(result["geojson"]) == Path(result["folder"]) / "drainage_lines.geojson"
    assert Path(result["lineament_geojson"]) == Path(result["folder"]) / "lineament_candidates.geojson"
    # every panel marks the site, labelled "Site" for want of a community
    # name; matplotlib writes each line of text in an SVG as a comment; a
    # rose is not a map, and has nowhere to mark it
    panels = {"relief": 1, "wetness": 1, "curvature": 2, "directions": 4, "lineaments": 1, "rose": 0, "plan": 1}
    for key, _, svg in result["figures"]:
        drawn = re.findall(r"<!-- (.*?) -->", Path(svg).read_text(encoding="utf-8"), re.S)
        assert " ".join(drawn).split().count("Site") == panels[key], key
    lines = json.loads(Path(result["geojson"]).read_text(encoding="utf-8"))
    assert lines["features"], "a 0.05 km2 threshold traces lines on the tiles"
    # the default 0.5 km2 traces a line on these tiles too, so the threshold
    # asked for is checked where the analysis and its sentences give it
    assert result["summary"]["threshold_units"] == 5_000_000, "the analysis ran at the 0.05 km2 asked for"
    assert "more than 0.05 km² drains" in result["captions"]["relief"]
    assert "more than 0.05 km² drains" in lines["description"]
    assert lines["description"] == result["captions"]["geojson_description"]
    assert lines["notice"] == result["notice"] == notice
    assert candidates["notice"] == notice and candidates["attribution"] == lines["attribution"]
    assert "not confirmed by geophysics" in candidates["description"]
    assert "not confirmed by geophysics" in result["captions"]["lineaments"]
    assert result["notes"][0].startswith("Read 2 files (")

    captions = [str(c.value) for c in at.caption]
    for key in keys:
        assert result["captions"][key] in captions, key
    assert all(note in captions for note in result["notes"])
    assert notice in captions
    offered = [element.proto.label for element in at.get("download_button")]
    assert [f"Download {name}" for name in names] + ["Download the drainage lines (GeoJSON)",
                                                      "Download the lineament candidates (GeoJSON)",
                                                      "Download the survey plan (GPX)",
                                                      "Download the survey plan (GeoLibre project)"] == [
        label for label in offered
        if "terrain_" in label or "drainage" in label or "lineament" in label or "survey plan" in label]
    # the run's progress is cleared once the figures are drawn, rather than
    # left ticked under the title of the last of them
    assert not at.status, [status.label for status in at.status]

    goto(at, "Overview")
    goto(at, "Site maps")
    assert not at.exception, at.exception
    assert at.number_input(key="terrain_radius_km").value == 1.0
    assert at.number_input(key="terrain_area_km2").value == 0.05
    assert at.session_state["terrain_result"] == result, "the same files, handed back by the uploader"
    assert result["captions"]["relief"] in [str(c.value) for c in at.caption]
    assert not _analysis_kept(at), "nor is it kept by a visit to another page"

    at.number_input(key="terrain_radius_km").set_value(50.0)
    at.run()
    at.button(key="run_terrain").click()
    at.run()
    assert not at.exception, at.exception
    assert [str(e.value)[:len("A 50 km radius")] for e in at.error] == ["A 50 km radius"]
    assert "terrain_result" not in at.session_state, "the figures were of the 1 km square"
    assert result["captions"]["relief"] not in [str(c.value) for c in at.caption]
    assert not Path(result["folder"]).exists()
    assert result["geojson"] not in at.session_state["artifacts"].values()
    # the refused run's status stays, at the stage that refused it, with its bar
    status, = at.status
    assert status.state == "error" and status.label == phrase("terrain.stage_read")
    assert [child.type for child in status.children.values()] == ["progress"]


def _desk_study_text(at: AppTest) -> tuple[str, int]:
    """Press "Build the desk study (.docx)" and give its text, paragraphs
    and table cells in order, and how many pictures it holds."""
    import docx

    at.button(key="build_desk_study").click()
    at.run()
    assert not at.exception, at.exception
    doc = docx.Document(at.session_state["artifacts"]["Download the desk study (.docx)"])
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells), len(doc.inline_shapes)


def test_the_desk_study_carries_the_terrain_press(sample_data):
    """PLAN.md step 4.4: after a terrain press, the desk study on Site maps
    carries what the press found, its figures, and the plan's tables, with
    the pre-field stamp on the page after the cover; the press keeps what
    the desk study says of the terrain, small, beside its figures."""
    from groundwater.text import phrase

    at = _terrain_page(_tiles())
    result = _analyse(at)
    desk = result["desk"]
    assert desk["site"] is not None and desk["label"]
    assert len(json.dumps(desk)) < 20000
    text, pictures = _desk_study_text(at)
    assert phrase("desk_study.stamp_title") in text
    assert phrase("desk_study.terrain_heading") in text
    assert phrase("desk_study.terrain_none") not in text
    drawn = {key for key, _, _ in result["figures"]}
    # the four area maps, and each terrain figure the content names that the press drew
    named = {"relief", "lineaments"} | ({"rose"} if desk["candidates"] else set()) | (
        {"plan"} if desk["lines"] else set())
    assert pictures == 4 + len(named & drawn)
    if desk["soundings"]:
        assert desk["soundings"][0][0] in text
    assert "Earlier projects within 5 km" in text and phrase("desk_study.projects_not_loaded") in text


def test_a_lookup_short_of_a_rings_reach_is_said_under_the_plan(sample_data):
    """A lookup on the Water points page at 300 m never loads a borehole 300
    to 550 m from the site, whose 50 m ring a sounding inside the service
    radius may stand in, and the plan said nothing of it. The press now
    tells the plan how far round the site the lookup went, and the caption
    says it fell short; it no longer says no water points were loaded."""
    from groundwater.models import SiteMetadata
    from groundwater.text import phrase

    site = SiteMetadata(easting=280126.0, northing=994900.0, utm_zone=29)
    at = _terrain_page(_tiles())
    at.session_state["wp_result"] = {"decision": {"search_radius_m": 300.0}, "rows": [], "points": [],
                                     "centre": site.latlon}
    caption = _analyse(at)["captions"]["plan"]
    assert phrase("survey_plan.water_points_short", searched=300.0, reach=550.0) in caption
    assert phrase("survey_plan.no_water_points") not in caption


def test_the_desk_study_without_an_elevation_model(sample_data):
    """With no terrain analysed, the desk study still gives the ground, the
    odds and the area maps, and says the terrain and plan wait for one."""
    from groundwater.text import phrase

    at = _terrain_page([])
    text, pictures = _desk_study_text(at)
    assert phrase("desk_study.terrain_none") in text and phrase("desk_study.plan_no_terrain") in text
    assert phrase("desk_study.water_not_loaded") in text
    assert "USGS Geologic Map of Africa" in text and pictures == 4


def test_the_topography_tab_analyses_an_ascii_grid_in_a_window_of_its_own(sample_data):
    """An ESRI ASCII grid is read whole, not as a window of tiles, so the
    analysis cuts the window out of it around the site, the margin as far
    as the grid reaches, and credits the figures and the drainage lines
    with the grid's own line; there is no Copernicus notice to give. The
    read says of the window cut what the read of tiles says: a square asked
    wider than the grid was reported as partly a gap, under figures cut to
    the grid, which draw none, and it is now given as the square shown.
    Without a location, the window is centred on the grid, and the read
    says so."""
    import json
    import re

    from groundwater.text import phrase

    # a valley along 11 W falling to the south, at 3 arc-seconds
    n, step = 120, 1 / 1200
    rows = []
    for r in range(n):                       # the file runs north first
        lat = 8.95 + (n - 1 - r + 0.5) * step
        rows.append(" ".join(
            f"{50 + 2000 * abs(-11.05 + (c + 0.5) * step + 11.0) + 1000 * (lat - 8.95):.1f}"
            for c in range(n)))
    text = (f"ncols {n}\nnrows {n}\nxllcorner -11.05\nyllcorner 8.95\ncellsize {step!r}\n"
            "NODATA_value -9999\n" + "\n".join(rows) + "\n")
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Site maps"
    at.session_state["meta_easting"] = 280126.0
    at.session_state["meta_northing"] = 994900.0
    at.session_state["meta_zone"] = "29N"
    at.run()
    at.file_uploader(key="dem_upload").set_value([("valley.asc", text.encode(), "text/plain")])
    at.run()
    at.number_input(key="terrain_radius_km").set_value(1.0)
    at.number_input(key="terrain_area_km2").set_value(0.05)
    at.run()
    at.button(key="run_terrain").click()
    at.run()
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]

    def shown(result: dict) -> tuple[str, str]:
        """The radius of the square shown and its place, as the GeoJSON's
        description measures them on the window."""
        said = re.search(r"within ([0-9.]+) km each way of (.+?), analysed",
                         result["captions"]["geojson_description"])
        assert said, result["captions"]["geojson_description"]
        return said.group(1), said.group(2)

    result = at.session_state["terrain_result"]
    assert result["notice"] is None
    assert not _analysis_kept(at), "the analysis's layers are dropped"
    lines = json.loads(Path(result["geojson"]).read_text(encoding="utf-8"))
    assert lines["features"], "the valley drains more than 0.05 km2"
    assert result["summary"]["threshold_units"] == 5_000_000, "the analysis ran at the 0.05 km2 asked for"
    assert "more than 0.05 km² drains" in result["captions"]["relief"]
    assert "more than 0.05 km² drains" in lines["description"]
    assert lines["attribution"].startswith("Elevation: ") and "valley.asc" in lines["attribution"]
    assert "notice" not in lines and lines["description"] == result["captions"]["geojson_description"]
    # the grid ends 5 km south of the site, so the square shown has about 4
    # km of the 5 km margin on that side, and the caption says what it had
    said = re.search(r"over a square ([0-9.]+) km wider on every side", result["captions"]["relief"])
    assert said and 3.8 <= float(said.group(1)) <= 4.1, result["captions"]["relief"]
    assert phrase("terrain.caption_edge", margin=float(said.group(1))) in result["captions"]["relief"]
    assert all(Path(png).is_file() and Path(svg).is_file() for _, png, svg in result["figures"])
    # a GPS fix inside the grid: the read gives the file, the square shown
    # and the margin, as the read of tiles does
    radius, where = shown(result)
    read, margin = result["notes"]
    assert read.startswith("Read 1 file (valley.asc): ") and read.endswith(f" around {where}, {radius} km each way.")
    assert margin == phrase("terrain.read_margin", margin=float(said.group(1)))
    captions = [str(c.value) for c in at.caption]
    assert all(note in captions for note in result["notes"])

    # the grid is about 11 km across: a 10 km radius is cut to it, so the
    # read gives the square the figures show and no gap they do not draw
    at.number_input(key="terrain_radius_km").set_value(10.0)
    at.run()
    cut = _analyse(at)
    radius, where = shown(cut)
    assert float(radius) < 10.0
    assert [note[-len(f" around {where}, {radius} km each way."):] for note in cut["notes"]] == [
        f" around {where}, {radius} km each way."], cut["notes"]
    assert phrase("terrain.caption_edge_none") in cut["captions"]["relief"], "the grid has no margin beyond it"

    # with no location recorded, the window is centred on the middle of the
    # grid, and the read says so
    bare = _terrain_page([("valley.asc", text.encode(), "text/plain")], fix=None)
    centred = _analyse(bare)
    assert phrase("terrain.centre_no_site") in centred["notes"]
    assert phrase("terrain.centre_no_site") in [str(c.value) for c in bare.caption]


def test_a_terrain_run_stopped_part_way_leaves_the_figures_shown_whole(sample_data):
    """Every run wrote its figures over the same files, so a second analysis
    stopped part-way, as a click anywhere on the page stops a run at its
    next progress update, left the first result on the page over figures the
    second had already redrawn: an image of the new threshold under the old
    one's caption, beside the old one's other figures and drainage lines.
    Each run now writes into a folder of its own. One stopped part-way
    leaves the result shown untouched and nothing of its own behind, and a
    result's folder is deleted once a finished run has taken its place."""
    import hashlib
    import sys

    from streamlit.runtime.scriptrunner_utils.script_requests import RerunData
    from streamlit.runtime.scriptrunner_utils.script_run_context import get_script_run_ctx

    def digests(result: dict) -> dict:
        files = [path for _, png, svg in result["figures"] for path in (png, svg)] + [result["geojson"],
                                                                                    result["lineament_geojson"]]
        return {path: hashlib.sha256(Path(path).read_bytes()).hexdigest() for path in files}

    at = _terrain_page(_tiles())
    first = _analyse(at)
    drawn = digests(first)
    runs = Path(at.session_state["workdir"])

    # a rerun asked for once the first figure is drawn, as a click on the
    # page asks for one: Streamlit stops the run at its next progress update
    site_maps = sys.modules["views.site_maps"]
    figures = site_maps._TERRAIN_FIGURES

    def then_rerun(draw):
        def drawing(*args, **kwargs):
            png = draw(*args, **kwargs)
            ctx = get_script_run_ctx()
            ctx.script_requests.request_rerun(RerunData(page_script_hash=ctx.page_script_hash))
            return png
        return drawing

    at.number_input(key="terrain_area_km2").set_value(0.5)
    at.run()
    site_maps._TERRAIN_FIGURES = ((figures[0][0], then_rerun(figures[0][1])),) + figures[1:]
    try:
        at.button(key="run_terrain").click()
        at.run()
    finally:
        site_maps._TERRAIN_FIGURES = figures
    assert not at.exception, at.exception
    assert at.session_state["terrain_result"] == first
    assert digests(first) == drawn, "the figures shown were not drawn over"
    assert sorted(runs.glob("terrain_*")) == [Path(first["folder"])], "nothing of the stopped run is left"

    second = _analyse(at)
    assert "more than 0.5 km² drains" in second["captions"]["relief"]
    assert all(Path(path).is_file() for path in digests(second))
    assert sorted(runs.glob("terrain_*")) == [Path(second["folder"])], "the first run's folder is deleted"


def test_the_terrain_figures_are_not_shown_beside_files_they_were_not_worked_out_from(sample_data):
    """The terrain figures, with their read report and Copernicus notice,
    describe the files and the declared source they were worked out from,
    and nothing held that the page drops them when either changes. They go,
    with their files, when another source is declared, a file is removed or
    added, or the source is put in other words; a visit to another page,
    which hands the same files back, keeps them (the analysis test holds
    that)."""
    pair = _tiles()
    at = _terrain_page(pair)

    def dropped(result: dict, why: str) -> None:
        assert not at.exception, at.exception
        assert "terrain_result" not in at.session_state, why
        captions = [str(c.value) for c in at.caption]
        assert not set(result["notes"]) & set(captions), why
        assert result["captions"]["relief"] not in captions and result["notice"] not in captions, why
        assert not Path(result["folder"]).exists(), why

    result = _analyse(at)
    at.selectbox(key="terrain_source").select("srtm")
    at.run()
    dropped(result, "another source declared")

    at.selectbox(key="terrain_source").select("detect")
    at.run()
    result = _analyse(at)
    # the site is on the eastern tile's western edge
    at.file_uploader(key="dem_upload").set_value(pair[1:])
    at.run()
    dropped(result, "a tile removed")

    result = _analyse(at)
    at.file_uploader(key="dem_upload").set_value(pair)
    at.run()
    dropped(result, "a tile added")

    at.selectbox(key="terrain_source").select("other")
    at.run()
    result = _analyse(at)
    at.text_input(key="terrain_source_text").set_value("viewfinderpanoramas")
    at.run()
    dropped(result, "the source put in other words")


def test_the_terrain_analysis_takes_a_missing_tile_as_sea_as_the_read_does(sample_data):
    """Copernicus publishes no tiles of open sea, so the read takes a missing
    tile whose whole edge with the tiles supplied is at or below 0 m as sea
    with no height, and the page hands that on to the analysis. Nothing held
    that it does: without it the sea is a gap, the lines draining to it are
    dashed, and the relief caption gives 38% of the square as sea rather
    than 63%."""
    import dataclasses

    from groundwater.config import Config
    from groundwater.geo import geographic_to_utm
    from groundwater.mapping import (
        analyse_terrain,
        lineament_candidates,
        lineament_captions,
        read_elevation,
        survey_plan,
        survey_plan_captions,
        terrain_captions,
    )
    from groundwater.models import SiteMetadata

    # three tiles around a missing one, the site at their corner, 8 N 12 W
    names = ("holes_N08W012.tif", "holes_sea_N08W013.tif", "holes_sea_N07W012.tif")
    fix = geographic_to_utm(8.0, -12.0, 29)
    at = _terrain_page(_tiles(*names), fix=(fix.easting, fix.northing))
    at.selectbox(key="terrain_source").select("copernicus_glo30")
    at.number_input(key="terrain_radius_km").set_value(15.0)
    at.number_input(key="terrain_area_km2").set_value(1.0)
    at.run()
    result = _analyse(at)
    assert result["summary"]["assumed_sea"] > 0

    dem = Path(__file__).resolve().parent / "webapp" / "fixtures" / "dem"
    cfg = dataclasses.replace(Config().terrain, drainage_area_km2=1.0)
    site = SiteMetadata(easting=fix.easting, northing=fix.northing, utm_zone=29)
    read = read_elevation([(n, dem / n) for n in names], site, 15.0, cfg, declared=("copernicus_glo30", ""))
    analysis = analyse_terrain(read.grid, cfg, assumed_sea=read.assumed_sea, shown=read.shown, precision="float32")
    where = read.centre.label
    candidates = lineament_candidates(analysis)
    lat, lon = site.latlon
    plan = survey_plan(candidates, analysis.core, analysis.parts, (lon, lat), (), cfg)
    assert result["captions"] == {**terrain_captions(analysis, read, where), **lineament_captions(candidates, where),
                                  **survey_plan_captions(plan, where)}


def test_a_project_load_takes_the_site_maps_off_the_deliverables(sample_data):
    """Loading a project dropped the outgoing site's maps from the Site maps
    page, but the Overview's Deliverables still offered them, the terrain
    figures and drainage lines among them, as though they were the loaded
    site's. They leave the list with the page, and the terrain figures are
    deleted."""
    import groundwater
    from groundwater.project_io import serialize_project

    at = _terrain_page(_tiles())
    at.slider(key="map_radius").set_value(5)
    at.run()
    at.button(key="run_topo").click()
    at.run()
    assert not at.exception and not at.error, (at.exception, [e.value for e in at.error])
    result = _analyse(at)
    built = {at.session_state["topo_path"], result["geojson"]} | {
        path for _, png, svg in result["figures"] for path in (png, svg)}
    assert built <= set(at.session_state["artifacts"].values())

    saved = serialize_project({"meta_community": "Kuntolo"}, groundwater.__version__)
    at.file_uploader(key="project_upload").set_value(("kuntolo_project.yaml", saved, "application/x-yaml"))
    at.run()
    at.button(key="project_load").click()
    at.run()
    assert not at.exception, at.exception
    assert any("Project loaded" in str(s.value) for s in at.success)
    assert not built & set(at.session_state["artifacts"].values())
    assert not Path(result["folder"]).exists()
    goto(at, "Overview")
    assert not at.exception, at.exception
    assert "📦 Deliverables" not in [str(h.value) for h in at.subheader], "nothing else was built"


def test_a_loaded_projects_ledger_is_kept_and_the_next_load_replaces_it(app):
    """PLAN.md step 5.1: the ledger of predictions and outcomes rides in the
    project file. The app keeps a loaded file's ledger for its next save,
    and a project loaded after it brings its own, empty or not, never the
    outgoing one's."""
    from types import SimpleNamespace

    import yaml

    import groundwater
    from groundwater import ledger as L
    from groundwater.project_io import serialize_project

    interp = SimpleNamespace(
        sounding_id="VES 1", depth_to_basement_m=28.4, basement_not_resolved=False,
        investigation_depth_m=60.0, water_zones=[(8.0, 20.0)], aquifer_thickness_m=12.0,
        max_drilling_depth_m=35.0, drilling_depth_capped=False, fit_error_percent=3.2,
        site_easting=708958.0, site_northing=926355.0)
    site = SimpleNamespace(community="Rokel", district="Port Loko", utm_zone=28)
    ledger, _ = L.record_predictions(L.empty_ledger(), L.survey_predictions([interp], None, None, site),
                                     "2026-10-01", False, groundwater.__version__)
    at = app
    for name, held in (("rokel", ledger), ("kuntolo", None)):
        saved = serialize_project({"meta_community": name, "prediction_ledger": held}, groundwater.__version__)
        at.file_uploader(key="project_upload").set_value((f"{name}_project.yaml", saved, "application/x-yaml"))
        at.run()
        at.button(key="project_load").click()
        at.run()
        assert not at.exception, at.exception
        assert at.session_state["prediction_ledger"] == L.clean_ledger(held), name
        # the key serialize_project reads, so the next save writes it again
        resaved = serialize_project({"prediction_ledger": at.session_state["prediction_ledger"]},
                                    groundwater.__version__)
        assert yaml.safe_load(resaved)["ledger"] == L.clean_ledger(held), name


def test_project_state_tracked(app):
    """The state the project file persists is maintained by the app."""
    assert "rates_overrides" in app.session_state
    assert app.session_state["rates_overrides"], "working rates not tracked"


def test_guided_wizard_flow(app):
    """Walk the wizard: site details -> siting -> costing -> done."""
    goto(app, "Guided start")
    # step 0 gate: needs community and district from the sidebar
    app.text_input(key="meta_community").set_value("Kuntolo")
    app.selectbox(key="meta_district").select("Bombali")
    app.run()
    assert not app.exception
    app.button(key="wiz_next").click()
    app.run()
    assert app.session_state["wiz_step"] == 1

    # step 1: run the siting analysis on the bundled sample
    app.selectbox(key="sample_wiz_ves").select("rokel/rokel_ves.xlsx")
    app.run()
    app.button(key="wiz_run_ves").click()
    app.run()
    assert not app.exception
    assert "ves_results" in app.session_state
    # the fresh result must unlock Next in the same rerun (review fix)
    next_button = app.button(key="wiz_next")
    assert not getattr(next_button, "disabled", next_button.proto.disabled)
    app.button(key="wiz_next").click()
    app.run()
    assert app.session_state["wiz_step"] == 2

    # step 2: costing prefilled from the siting result
    app.button(key="wiz_cost_run").click()
    app.run()
    assert not app.exception
    assert app.session_state["cost_estimate"].direct_cost_usd > 0
    app.button(key="wiz_next").click()
    app.run()
    assert app.session_state["wiz_step"] == 3


def test_wizard_unlocks_after_first_ves_run(sample_data):
    """Regression: the first siting run must enable Next immediately.

    Uses a fresh app instance (no prior VES results) to reproduce the
    reported state: reaching step 1 cold and running the analysis once.
    """
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Guided start"
    at.run()
    at.text_input(key="meta_community").set_value("Kuntolo")
    at.selectbox(key="meta_district").select("Bombali")
    at.run()
    at.button(key="wiz_next").click()
    at.run()
    at.selectbox(key="sample_wiz_ves").select("rokel/rokel_ves.xlsx")
    at.run()
    at.button(key="wiz_run_ves").click()
    at.run()
    assert not at.exception
    next_button = at.button(key="wiz_next")
    assert not getattr(next_button, "disabled", next_button.proto.disabled)


def test_wizard_grace_survives_until_costing_step(sample_data):
    """Regression: a project loaded on a step other than costing must
    keep its restored wizard costing values when the costing step is
    finally visited, even though the load rerun has long passed."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Guided start"
    # state as the project loader would leave it: saved at the final
    # step with a siting-derived signature that the fresh session
    # cannot reproduce, plus the wizard grace marker
    at.session_state["wiz_step"] = 3
    at.session_state["wiz_cost_depth"] = 85.0
    at.session_state["wiz_cost_over"] = 12.0
    at.session_state["wiz_prefill_sig"] = "72.0:24.0"
    at.session_state["_wiz_load_grace"] = True
    at.run()  # the load rerun: costing block does not execute
    assert not at.exception
    at.session_state["wiz_step"] = 2  # user navigates back to costing
    at.run()
    assert not at.exception
    assert at.session_state["wiz_cost_depth"] == 85.0
    assert at.session_state["wiz_cost_over"] == 12.0
    # grace is consumed: a genuine source change afterwards resets
    assert "_wiz_load_grace" not in at.session_state


def test_wizard_grace_cleared_by_siting_change(sample_data):
    """Regression: a siting change made after loading a project must
    invalidate the load grace, so costing follows the new depth
    instead of the stale loaded values."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Guided start"
    at.session_state["wiz_step"] = 1
    at.session_state["wiz_cost_depth"] = 85.0
    at.session_state["wiz_cost_over"] = 12.0
    at.session_state["wiz_prefill_sig"] = "72.0:24.0"
    at.session_state["_wiz_load_grace"] = True
    at.run()
    assert not at.exception
    # the user changes the planned depth on the siting step
    at.number_input(key="wiz_manual_depth").set_value(95.0)
    at.run()
    assert "_wiz_load_grace" not in at.session_state
    at.session_state["wiz_step"] = 2
    at.run()
    assert not at.exception
    # costing adopted the new siting depth, not the stale loaded value
    assert at.session_state["wiz_cost_depth"] == 95.0


def _fail_inversion_for(monkeypatch, sounding_id):
    """Make the inversion raise for one sounding and run for the rest.

    The app imports invert_sounding from groundwater.ves on every run,
    so patching the package attribute reaches the next run of the
    script.
    """
    import groundwater.ves as ves_pkg

    real = ves_pkg.invert_sounding

    def invert(sounding, *args, **kwargs):
        if sounding.sounding_id == sounding_id:
            raise ValueError("synthetic inversion failure")
        return real(sounding, *args, **kwargs)

    monkeypatch.setattr(ves_pkg, "invert_sounding", invert)


def test_a_failed_inversion_keeps_the_wizard_grace(sample_data, monkeypatch):
    """Regression (PR #17): a siting run that fails part way must neither
    crash the tab nor consume a loaded project's wizard grace, since it
    produced no siting result to change the costing prefill."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Guided start"
    at.session_state["wiz_step"] = 1
    at.run()
    assert not at.exception
    _fail_inversion_for(monkeypatch, "B (2)")
    at.session_state["_wiz_load_grace"] = True
    at.selectbox(key="sample_wiz_ves").select("rokel/rokel_ves.xlsx")
    at.run()
    bars = [e.proto.value for e in at.get("progress")]
    at.button(key="wiz_run_ves").click()
    at.run()
    assert not at.exception, at.exception
    errors = " ".join(str(e.value) for e in at.error)
    assert "Inversion failed for sounding B (2)" in errors
    # the run's part-full progress bar is cleared, not left beside the error
    assert [e.proto.value for e in at.get("progress")] == bars
    assert "ves_results" not in at.session_state
    assert at.session_state["_wiz_load_grace"] is True


def test_a_failed_inversion_keeps_the_previous_result(sample_data, monkeypatch):
    """A failed rerun leaves the last successful siting result in place:
    it is still the best answer anyone has, and the reports, the Overview
    and the wizard's costing prefill all read it."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Geophysics (VES)"
    at.run()
    at.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
    at.run()
    at.button(key="run_ves").click()
    at.run()
    assert not at.exception
    before = at.session_state["ves_results"]
    at.session_state["_wiz_load_grace"] = True
    _fail_inversion_for(monkeypatch, "B (2)")
    at.button(key="run_ves").click()
    at.run()
    assert not at.exception, at.exception
    errors = " ".join(str(e.value) for e in at.error)
    assert "Inversion failed for sounding B (2)" in errors
    assert "previous successful run" in errors
    after = at.session_state["ves_results"]
    assert [r.fit_error_percent for r in after[1]] == [
        r.fit_error_percent for r in before[1]]
    assert [i.rank for i in after[2]] == [i.rank for i in before[2]]
    assert at.session_state["_wiz_load_grace"] is True


def test_templates_tab(app):
    goto(app, "Templates")
    app.button(key="gen_templates").click()
    app.run()
    assert not app.exception


def test_the_field_kit_is_built_from_the_templates_page(app):
    """PLAN.md step 2.5: the Field kit button writes the sheets and cards
    for the boreholes named, and says so when none is."""
    goto(app, "Templates")
    app.text_area(key="templates_fieldkit_boreholes").set_value("KTL-01\nKTL-02")
    app.button(key="templates_fieldkit_build").click()
    app.run()
    assert not app.exception, app.exception
    kit = app.session_state["artifacts"]["Download the field kit (.docx)"]
    import docx

    text = "\n".join(p.text for p in docx.Document(kit).paragraphs)
    assert "borehole KTL-01" in text and "borehole KTL-02" in text
    app.text_area(key="templates_fieldkit_boreholes").set_value("  ")
    app.button(key="templates_fieldkit_build").click()
    app.run()
    assert any("at least one borehole" in str(e.value) for e in app.error)


def test_the_field_kit_box_offers_a_loaded_test_and_keeps_what_was_typed(sample_data):
    """The box offers the borehole of a test loaded on the same page, and
    what is typed into it is still there on the other page and on coming
    back, as in the browser app. Streamlit keeps a keyed widget's first
    value and forgets it when the page is left, so neither held before."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    goto(at, "Pumping test")
    assert at.text_area(key="pump_fieldkit_boreholes").value == ""
    at.selectbox(key="sample_pump").select("dr_timbo/dr_timbo_constant_test.xlsx")
    at.run()
    ref = at.session_state["pump_analysis"].test.borehole_ref
    assert ref and at.text_area(key="pump_fieldkit_boreholes").value == ref
    at.text_area(key="pump_fieldkit_boreholes").set_value("TYPED-1\nTYPED-2")
    at.run()
    goto(at, "Templates")
    assert at.text_area(key="templates_fieldkit_boreholes").value == "TYPED-1\nTYPED-2"
    goto(at, "Pumping test")
    assert at.text_area(key="pump_fieldkit_boreholes").value == "TYPED-1\nTYPED-2"
    assert not at.exception, at.exception


def test_the_survey_s_predictions_meet_the_borehole_in_the_ledger(app):
    """PLAN.md step 5.1: the Geophysics page records the survey's
    predictions, a drilling log's outcome goes to the point chosen as drilled,
    a forecast is kept as it was made, and building the report records too."""
    from groundwater.ingestion import read_drilling_workbook

    goto(app, "Geophysics (VES)")
    app.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
    app.run()
    app.button(key="run_ves").click()
    app.run()
    assert not app.exception
    assert "No prediction is recorded yet." in [str(i.value) for i in app.info]
    # the Rokel sheets name no field supervisor, so the project's is the team
    app.session_state["meta_supervisor"] = "  R.  Conteh "
    app.button(key="record_predictions").click()
    app.run()
    assert not app.exception, app.exception
    ledger = app.session_state["prediction_ledger"]
    assert [e["point"]["sounding_id"] for e in ledger["entries"]] == ["A (1)", "B (2)"]
    assert [e["prediction"]["survey_team"] for e in ledger["entries"]] == ["R. Conteh", "R. Conteh"]
    assert all(e["before_outcome"] for e in ledger["entries"])
    assert all(e["prediction"]["odds"] is not None for e in ledger["entries"])
    # no range was sampled, and the prediction says so
    assert not any(e["prediction"]["range_sampled"] for e in ledger["entries"])

    # Dr Timbo's log carries no position: with two points it cannot be paired
    app.session_state["drilling_log"] = read_drilling_workbook(
        REPO / "examples" / "data" / "dr_timbo" / "dr_timbo_drilling_log.xlsx")
    app.run()
    assert not app.exception
    assert any("could not be matched to a point" in str(w.value) for w in app.warning)
    drilled = ledger["entries"][1]["id"]
    app.selectbox(key="ledger_drilled_pick").select(drilled)
    app.run()
    assert not app.exception, app.exception
    assert app.session_state["prediction_ledger"]["drilled"] == drilled
    assert any("B (2)" in str(m.value) and "chosen on this page" in str(m.value) for m in app.markdown)
    shown = " ".join(" ".join(t.value.astype(str).to_numpy().ravel()) for t in app.table)
    assert "Water struck at 12, 30 m" in shown and "Drilled to 70 m" in shown

    # once drilled, the forecasts are kept as they were made, by the button
    # and by the report alike
    app.button(key="record_predictions").click()
    app.run()
    assert any("2 point(s) kept the forecast" in str(i.value) for i in app.info)
    app.button(key="build_geo_report").click()
    app.run()
    assert not app.exception
    after = app.session_state["prediction_ledger"]
    assert [e["recorded_on"] for e in after["entries"]] == [e["recorded_on"] for e in ledger["entries"]]
    assert all(e["before_outcome"] for e in after["entries"])


def test_the_portfolio_pools_the_projects_ledgers(sample_data):
    """PLAN.md step 5.1: the Portfolio page pools the ledgers of the files
    it is given, each point once, and offers them as one file."""
    from types import SimpleNamespace

    from groundwater import ledger as L
    from groundwater.project_io import serialize_project

    def saved(community, easting, status):
        site = SimpleNamespace(community=community, district="Port Loko", utm_zone=28)
        interp = SimpleNamespace(
            sounding_id="VES 1", depth_to_basement_m=28.4, basement_not_resolved=False,
            investigation_depth_m=60.0, water_zones=[(8.0, 20.0)], aquifer_thickness_m=12.0,
            max_drilling_depth_m=35.0, drilling_depth_capped=False, fit_error_percent=3.2,
            site_easting=easting, site_northing=926355.0)
        ledger, _ = L.record_predictions(L.empty_ledger(), L.survey_predictions([interp], None, None, site),
                                         "2026-10-01", False, "0.3.0")
        log = SimpleNamespace(borehole_ref="BH-1", total_depth_m=45.0, intervals=[], water_strikes_m=[],
                              status=status, completion_date="",
                              site=SimpleNamespace(easting=None, northing=None, utm_zone=None))
        ledger = L.attach_outcome(ledger, L.completion_outcome(log, None))
        return serialize_project({"project_summary": {"community": community, "district": "Port Loko",
                                                      "status": status},
                                  "prediction_ledger": ledger}, "0.3.0")

    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Portfolio"
    at.run()
    rokel = saved("Rokel", 708958.0, "Successful")
    at.file_uploader(key="portfolio_upload").set_value([
        ("rokel.yaml", rokel, "application/x-yaml"),
        ("rokel_copy.yaml", rokel, "application/x-yaml"),
        ("kuntolo.yaml", saved("Kuntolo", 709500.0, "Dry hole"), "application/x-yaml"),
    ])
    at.run()
    assert not at.exception, at.exception
    assert any("2 survey points predicted in 2 projects, 2 of them drilled" in str(c.value)
               for c in at.caption)
    table = next(d.value for d in at.dataframe if "Basement predicted" in d.value.columns)
    assert list(table["Project"]) == ["Rokel (Port Loko)", "Kuntolo (Port Loko)"]
    assert any(b.label == "Download the pooled ledger (.json)" for b in at.download_button)

    # PLAN.md step 5.2: the calibration report over the pooled ledgers, its
    # summary on the page and the report offered for these files only
    from groundwater.text import phrase

    assert "Calibration" in [str(h.value) for h in at.subheader]
    shown = [str(m.value) for m in at.markdown]
    assert phrase("calibration.summary_scored", forecast=2) in shown
    assert phrase("calibration.summary_depth_too_few", judged=0, minimum=10) in shown
    at.button(key="build_calibration").click()
    at.run()
    assert not at.exception, at.exception
    built = at.session_state["calibration_report"]
    from docx import Document

    text = "\n".join(para.text for para in Document(built["path"]).paragraphs)
    assert phrase("calibration.title") in text and phrase("calibration.hits_heading") in text
    assert any(b.label == "Download the calibration report (.docx)" for b in at.download_button)
    # another set of files wants the report built again
    at.file_uploader(key="portfolio_upload").set_value([("rokel.yaml", rokel, "application/x-yaml")])
    at.run()
    assert not at.exception, at.exception
    assert not any(b.label == "Download the calibration report (.docx)" for b in at.download_button)


def _drilled_project(glg: str = "pCm", code: str = "B-L") -> bytes:
    """A saved project whose ledger holds 30 drilled points on one geology
    unit, two in three of them successful, logged with basement from 22 m
    where the survey read 2,400 ohm-m: enough to fit a set, whose basement
    edge moves from 3,000 to 579.7 ohm-m, the geometric midpoint of
    140 and 2,400."""
    from groundwater.project_io import serialize_project

    log = [[0.0, 2.0, "laterite"], [2.0, 8.0, "clay"], [8.0, 22.0, "weathered"], [22.0, 50.0, "basement"]]
    layers = [[0.0, 2.0, 900.0], [2.0, 8.0, 12.0], [8.0, 22.0, 140.0], [22.0, None, 2400.0]]
    entries = []
    for k in range(30):
        success = k % 3 != 0
        entries.append({
            "id": f"e{k:03d}", "before_outcome": True, "recorded_on": "2026-03-01", "toolkit": "0.3.0",
            "point": {"sounding_id": f"VES {k}"},
            "prediction": {"investigation_depth_m": 60.0, "layers": layers,
                           "suitability": {"aquifer_thickness": 0.5, "resistivity_fit": 0.5, "overburden": 0.5,
                                           "basal_fracture": 0.9 if success else 0.1},
                           "odds": {"probability": 0.6, "prior": 0.55, "bgs_code": code, "glg": glg}},
            "outcome": {"total_depth_m": 50.0, "success": success, "success_yield_m3_per_h": 1.0,
                        "success_basis": "tested", "intervals": log}})
    return serialize_project({"project_summary": {"community": "Rokel", "district": "Port Loko",
                                                  "status": "Successful"},
                              "prediction_ledger": {"entries": entries}}, "0.3.0")


def test_the_portfolio_fits_a_calibration_set_for_a_named_person_to_adopt(sample_data):
    """PLAN.md step 5.3: a set fitted to the pooled ledgers is reviewed
    beside the set in force, changes nothing until a named person adopts it,
    and is then in force for the session, to download, replace or drop."""
    from groundwater import calibration_sets as CS
    from groundwater.project_io import PERSIST_PREFIXES, stale_on_load
    from groundwater.text import phrase

    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Portfolio"
    at.run()
    assert not at.exception, at.exception
    # with no files the set in force is still shown, and can be replaced
    assert phrase("calibration_sets.page_heading") in [str(h.value) for h in at.subheader]
    default = phrase("calibration_sets.page_in_force", name=CS.set_name(CS.default_set()))
    assert default in [str(m.value) for m in at.markdown]
    assert phrase("calibration_sets.page_fit_none") in [str(i.value) for i in at.info]
    assert any(b.label == phrase("calibration_sets.page_download") for b in at.download_button)

    at.file_uploader(key="portfolio_upload").set_value([("rokel.yaml", _drilled_project(), "application/x-yaml")])
    at.run()
    at.button(key="calibration_fit").click()
    at.run()
    assert not at.exception, at.exception
    shown = [str(m.value) for m in at.markdown]
    assert any(text.startswith("It uses 30 boreholes from the pooled ledgers") for text in shown)
    edges = next(d.value for d in at.dataframe if phrase("calibration_sets.edges_col_layers") in d.value.columns)
    assert list(edges[phrase("calibration_sets.edges_col_fitted")]) == ["20 ohm-m, 0 misread",
                                                                       "580 ohm-m, 0 misread"]
    # nothing is in force until a named person adopts it
    assert "calibration_set" not in at.session_state
    at.button(key="calibration_adopt").click()
    at.run()
    assert [str(e.value) for e in at.error] == [phrase("calibration_sets.adopt_unnamed")]
    at.text_input(key="calibration_adopter").set_value("  A. Kamara ")
    at.button(key="calibration_adopt").click()
    at.run()
    assert not at.exception, at.exception
    held = at.session_state["calibration_set"]
    assert (held["version"], held["adopted_by"]) == (1, "A. Kamara")
    # issued under the licence in the settings, for no organisation, since
    # none was typed and the sidebar names none
    assert (held["licence"], held["organisation"], held["contact"]) == ("CC BY 4.0", "", "")
    name = CS.set_name(held)
    assert phrase("calibration_sets.page_put_in_force", name=name) in [str(m.value) for m in at.success]
    shown = [str(m.value) for m in at.markdown]
    assert phrase("calibration_sets.page_in_force", name=name) in shown
    assert "- " + CS.set_changes(held)[0] in shown
    # the review is of a set no longer on offer
    assert not any(text.startswith("It uses 30 boreholes") for text in shown)

    # back to the default, and the adopted set's file put in force again
    at.button(key="calibration_default").click()
    at.run()
    assert at.session_state["calibration_set"] is None
    at.file_uploader(key="calibration_upload").set_value(
        ("calibration.yaml", CS.set_yaml(held).encode("utf-8"), "application/x-yaml"))
    at.run()
    at.button(key="calibration_load").click()
    at.run()
    assert not at.exception, at.exception
    assert at.session_state["calibration_set"] == held
    # a set nobody adopted is refused, and the one in force stays
    candidate = CS.fit_set([], "2026-10-24")
    at.file_uploader(key="calibration_upload").set_value(
        ("calibration.yaml", CS.set_yaml(candidate).encode("utf-8"), "application/x-yaml"))
    at.run()
    at.button(key="calibration_load").click()
    at.run()
    assert [str(e.value) for e in at.error] == [phrase("calibration_sets.not_adopted")]
    assert at.session_state["calibration_set"] == held
    # the set is the session's: no project file carries it, and opening one
    # leaves it in force
    assert not "calibration_set".startswith(PERSIST_PREFIXES)
    assert "calibration_set" not in stale_on_load({"calibration_set": held})


def _set_file(glg: str, code: str, organisation: str, on: str, start: int = 0) -> bytes:
    """A calibration.yaml another organisation fitted to 30 drilled points
    of its own on one geology unit, adopted and issued under CC BY 4.0."""
    from groundwater import calibration_sets as CS
    from groundwater import ledger as L

    log = [[0.0, 2.0, "laterite"], [2.0, 8.0, "clay"], [8.0, 22.0, "weathered"], [22.0, 50.0, "basement"]]
    layers = [[0.0, 2.0, 900.0], [2.0, 8.0, 12.0], [8.0, 22.0, 140.0], [22.0, None, 2400.0]]
    entries = [{"id": f"o{start + k}", "before_outcome": True, "recorded_on": "2026-03-01", "toolkit": "0.3.0",
                "point": {"sounding_id": f"VES {k}"},
                "prediction": {"investigation_depth_m": 60.0, "layers": layers, "suitability": None,
                               "odds": {"probability": 0.6, "prior": 0.55, "bgs_code": code, "glg": glg}},
                "outcome": {"total_depth_m": 50.0, "success": k % 3 != 0, "success_yield_m3_per_h": 1.0,
                            "success_basis": "tested", "intervals": log}} for k in range(30)]
    fitted = CS.fit_set(L.pool_ledgers([("theirs", {"entries": entries})]), on)
    return CS.set_yaml(CS.adopt_set(fitted, "A. Kamara", on, organisation=organisation)).encode("utf-8")


def test_the_portfolio_pools_sets_from_other_organisations_for_a_named_person_to_adopt(sample_data):
    """PLAN.md step 5.5: set files other organisations adopted are pooled
    into one set and reviewed beside the set in force, with no project file
    on the page; files that cannot be pooled are refused in a sentence
    naming the file; and nothing changes until a named person adopts the
    pooled set, issued in the name of their organisation."""
    from groundwater import calibration_sets as CS
    from groundwater.text import phrase

    trust = _set_file("pCm", "B-L", "Rural Water Trust", "2026-10-20")
    kono = _set_file("Pi", "I-L", "Kono Wells", "2026-09-01", start=100)
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Portfolio"
    at.session_state["org_name"] = "Salone Water Forum"
    at.run()
    assert not at.exception, at.exception
    assert f"**{phrase('calibration_sets.page_pool')}**" in [str(m.value) for m in at.markdown]
    assert phrase("calibration_sets.page_share") in [str(c.value) for c in at.caption]

    # the same set twice, under another organisation's name, is refused by file
    renamed = trust.replace(b'organisation: "Rural Water Trust"', b'organisation: "Someone Else"')
    at.file_uploader(key="calibration_pool_upload").set_value(
        [("trust.yaml", trust, "application/x-yaml"), ("copy.yaml", renamed, "application/x-yaml")])
    at.run()
    at.button(key="calibration_pool").click()
    at.run()
    assert not at.exception, at.exception
    assert [str(e.value) for e in at.error] == [phrase("calibration_sets.pool_same", label="copy.yaml",
                                                       other="trust.yaml")]
    assert "calibration_pooled" not in at.session_state

    at.file_uploader(key="calibration_pool_upload").set_value(
        [("trust.yaml", trust, "application/x-yaml"), ("kono.yaml", kono, "application/x-yaml")])
    at.run()
    at.button(key="calibration_pool").click()
    at.run()
    assert not at.exception, at.exception
    shown = [str(m.value) for m in at.markdown]
    assert any(text.startswith("It pools the 2 calibration sets below") for text in shown)
    sources = next(d.value for d in at.dataframe if phrase("calibration_sets.sources_col_licence") in d.value.columns)
    assert sorted(sources[phrase("calibration_sets.sources_col_organisation")]) == ["Kono Wells",
                                                                                    "Rural Water Trust"]
    # nothing is in force until a named person adopts it
    assert "calibration_set" not in at.session_state
    at.button(key="calibration_pool_adopt").click()
    at.run()
    assert [str(e.value) for e in at.error] == [phrase("calibration_sets.adopt_unnamed")]
    assert at.text_input(key="calibration_pool_organisation").value == "Salone Water Forum"
    at.text_input(key="calibration_pool_adopter").set_value("C. Conteh")
    at.text_input(key="calibration_pool_contact").set_value("forum@example.org")
    at.button(key="calibration_pool_adopt").click()
    at.run()
    assert not at.exception, at.exception
    held = at.session_state["calibration_set"]
    assert (held["version"], held["adopted_by"], held["organisation"], held["contact"]) == (
        1, "C. Conteh", "Salone Water Forum", "forum@example.org")
    assert sorted(s["organisation"] for s in held["sources"]) == ["Kono Wells", "Rural Water Trust"]
    name = CS.set_name(held)
    assert name.startswith("calibration set 1, pooled on ")
    assert phrase("calibration_sets.page_put_in_force", name=name) in [str(m.value) for m in at.success]
    assert "calibration_pooled" not in at.session_state
    # a set saved before step 5.5 declares no licence, and is refused by file
    old = CS.read_set(kono.decode("utf-8"))
    old = {key: value for key, value in old.items() if key not in ("licence", "organisation", "contact", "sources")}
    old.update(format_version=1, boreholes=[{"id": f"x{k}", "success": True, "forecast": True, "uses": ["prior"]}
                                            for k in range(old["boreholes"])])
    at.file_uploader(key="calibration_pool_upload").set_value(
        [("trust.yaml", trust, "application/x-yaml"),
         ("kono-old.yaml", CS.set_yaml(old).encode("utf-8"), "application/x-yaml")])
    at.run()
    at.button(key="calibration_pool").click()
    at.run()
    assert [str(e.value) for e in at.error] == [phrase("calibration_sets.pool_unlicensed", label="kono-old.yaml")]
    assert at.session_state["calibration_set"] == held


def test_a_set_put_in_force_reads_the_survey_on_show_again(sample_data):
    """The survey already inverted is read by the set put in force: each
    sounding by the edges the set fitted for the unit under it, its range of
    models dropped where those edges are not the ones it was sampled with."""
    from groundwater.calibration_sets import config_for_sounding
    from groundwater.config import Config
    from groundwater.text import phrase

    at = AppTest.from_file(APP, default_timeout=600)
    goto(at, "Geophysics (VES)")
    at.selectbox(key="sample_ves").select("rokel/rokel_ves.xlsx")
    at.run()
    at.button(key="run_ves").click()
    at.run()
    at.button(key="ves_range").click()
    at.run()
    assert not at.exception, at.exception
    soundings, results, before = at.session_state["ves_results"]
    # both soundings sit on the Quaternary sands, U-M/H on Qe
    assert before[0].depth_to_basement_m is None

    goto(at, "Portfolio")
    at.file_uploader(key="portfolio_upload").set_value(
        [("rokel.yaml", _drilled_project("Qe", "U-M/H"), "application/x-yaml")])
    at.run()
    at.button(key="calibration_fit").click()
    at.run()
    at.text_input(key="calibration_adopter").set_value("A. Kamara")
    at.button(key="calibration_adopt").click()
    at.run()
    assert not at.exception, at.exception
    held = at.session_state["calibration_set"]
    kept, again, after = at.session_state["ves_results"]
    assert kept is soundings and again is results
    config = config_for_sounding(Config(calibration=held), soundings[0])
    assert config.ves.fresh_basement_min_rho == 579.7
    # read by the set's basement edge, the 1,637 ohm-m layer is as resistive
    # as basement now; over the 47 ohm-m water zone it is a resistive cover,
    # not bedrock (it used to be read as basement at 1.0 m)
    assert before[0].layers[1].unit == "compact or dry weathered rock (regolith)"
    assert after[0].layers[1].unit == phrase("ves.unit_resistive_cover")
    assert after[0].depth_to_basement_m is None
    assert "ves_ranges" not in at.session_state

    at.button(key="calibration_default").click()
    at.run()
    assert not at.exception, at.exception
    assert at.session_state["ves_results"][2][0].depth_to_basement_m is None


def test_the_project_file_carries_the_works_record_and_brings_the_contract_back(sample_data, monkeypatch):
    """PLAN.md step 6.3: what the project records of its contractor's work
    goes into the project file as the works record, from the checklist, the
    drilled metres typed (blank until they are), the contract awarded with
    its variation orders and completion date, and the log. The contract, its
    tables and its date come back when the file is loaded, and a project
    loaded after it does not inherit them."""
    import sys

    import yaml

    if str(REPO / "app") not in sys.path:
        sys.path.insert(0, str(REPO / "app"))
    import shared
    from groundwater.project_io import serialize_project
    from groundwater.text import phrase

    saved = []

    def capture(session, version):
        data = serialize_project(session, version)
        saved.append(data)
        return data

    monkeypatch.setattr(shared, "serialize_project", capture)
    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    goto(at, "Borehole design")
    at.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    at.run()
    goto(at, "Costing & BoQ")
    at.button(key="run_cost").click()
    at.run()
    goto(at, "Supervision")
    # nothing typed is nothing reconciled, and no metres go into the record
    assert at.session_state["fx_logs_m"] is None and at.session_state["fx_invoiced_m"] is None
    reconciled = "Drilled metres reconciliation"
    assert reconciled not in [c.name for c in at.session_state["field_checks"]]
    assert yaml.safe_load(saved[-1])["works"]["logged_m"] is None
    at.number_input(key="fx_logs_m").set_value(70.0)
    at.number_input(key="fx_invoiced_m").set_value(75.0)
    at.session_state["chk_drl-daily-drilling-log"] = "No"
    at.run()
    assert not at.exception, at.exception
    checks = {c.name: c.passed for c in at.session_state["field_checks"]}
    assert checks[reconciled] is False

    goto(at, "Procurement")
    at.text_input(key="proc_ref").set_value("WSD/2018/1")
    at.text_input(key="proc_due").set_value("10/05/2018")
    at.run()
    assert phrase("scorecards.page_due_unread") in [str(w.value) for w in at.warning]
    at.text_input(key="proc_due").set_value("2018-05-10")
    at.button(key="proc_award").click()
    at.run()
    assert not at.exception, at.exception
    assert phrase("scorecards.page_due_unread") not in [str(w.value) for w in at.warning]
    lines = at.session_state["proc_contract_lines"]
    at.session_state["proc_variations"] = [
        {"Ref": "VO-1", "Date": "2018-05-12", "Code": lines[0]["code"], "Quantity change": 2.0,
         "New rate (USD)": None, "Reason": "Water deeper", "Authorised by": "A. Kamara"},
        {"Ref": "VO-2", "Date": "", "Code": lines[0]["code"], "Quantity change": 1.0,
         "New rate (USD)": None, "Reason": "", "Authorised by": ""}]
    at.run()
    assert not at.exception, at.exception

    # what Save project holds now; the next session saves its own
    data = saved[-1]
    payload = yaml.safe_load(data)
    works = payload["works"]
    assert (works["contractor"], works["drilled"], works["completion_date"]) == (
        "WiNGiN Heavy Duty Machines Co. Ltd", True, "13th May, 2018")
    assert (works["logged_m"], works["claimed_m"]) == (70.0, 75.0)
    assert works["checklist"]["drilling"]["failed"] == ["drl-daily-drilling-log"]
    assert (works["contract"], works["variations_signed"], works["variations_unsigned"]) == (True, 1, 1)
    assert works["variation_usd"] == round(2.0 * lines[0]["rate_usd"], 2)
    assert works["completion_due"] == "2018-05-10"
    assert len(payload["procurement"]["lines"]) == len(lines)
    assert [v["Ref"] for v in payload["procurement"]["variations"]] == ["VO-1", "VO-2"]
    assert payload["state"]["proc_ref"] == "WSD/2018/1" and payload["state"]["proc_due"] == "2018-05-10"
    assert "proc_award" not in payload["state"] and "proc_build" not in payload["state"]

    fresh = AppTest.from_file(APP, default_timeout=600)
    fresh.run()
    fresh.file_uploader(key="project_upload").set_value(("timbo_project.yaml", data, "application/x-yaml"))
    fresh.run()
    fresh.button(key="project_load").click()
    fresh.run()
    assert not fresh.exception, fresh.exception
    assert fresh.session_state["proc_contract_lines"] == payload["procurement"]["lines"]
    assert fresh.session_state["proc_variations"] == payload["procurement"]["variations"]
    assert fresh.session_state["proc_due"] == "2018-05-10"
    goto(fresh, "Procurement")
    assert "Contract sum" in [m.label for m in fresh.metric]
    # the next project brings its own contract, here none
    other = serialize_project({"meta_community": "Kuntolo"}, "0.3.0")
    fresh.file_uploader(key="project_upload").set_value(("kuntolo_project.yaml", other, "application/x-yaml"))
    fresh.run()
    fresh.button(key="project_load").click()
    fresh.run()
    assert not fresh.exception, fresh.exception
    assert "proc_contract_lines" not in fresh.session_state
    assert not fresh.session_state["proc_due"]


def test_the_variation_table_is_read_once_for_the_certificate_and_the_record():
    """PLAN.md step 6.3: the Procurement page's certificate and the works
    record read its variation orders through one function. A signature of
    white space is none, as the browser's form has it, and an empty number
    cell, None or NaN, is no quantity and no new rate."""
    import sys

    if str(REPO / "app") not in sys.path:
        sys.path.insert(0, str(REPO / "app"))
    import shared

    rows = [{"Ref": "VO-1", "Date": None, "Code": "DRL", "Quantity change": float("nan"), "New rate (USD)": "",
             "Reason": None, "Authorised by": " \t "},
            {"Ref": "VO-2", "Date": "", "Code": "  ", "Quantity change": 5.0, "Authorised by": "A. Kamara"},
            {"Ref": "VO-3", "Date": "2018-05-12", "Code": "CAS", "Quantity change": 2, "New rate (USD)": 12.5,
             "Reason": "Deeper", "Authorised by": " A. Kamara "},
            "not a row"]
    assert shared.procurement_variations(rows) == [
        {"ref": "VO-1", "date": "", "code": "DRL", "quantity_delta": 0.0, "rate_usd": None, "reason": "",
         "authorised_by": ""},
        {"ref": "VO-3", "date": "2018-05-12", "code": "CAS", "quantity_delta": 2.0, "rate_usd": 12.5,
         "reason": "Deeper", "authorised_by": "A. Kamara"}]
    assert shared.procurement_variations(None) == []


def _works_project(community: str, contractor: str, status: str, completion: str, due: str,
                   logged: float, claimed: float, failed: bool) -> bytes:
    """A saved project whose works record holds a drilled borehole: its
    contractor, the drilling checklist answered, the metres reconciled and a
    contract with one signed variation order."""
    from types import SimpleNamespace

    from groundwater.project_io import serialize_project
    from groundwater.scorecards import works_record

    log = SimpleNamespace(borehole_ref="BH-1", total_depth_m=45.0, intervals=[], water_strikes_m=[], status=status,
                          completion_date=completion,
                          site=SimpleNamespace(easting=None, northing=None, utm_zone=None, contractor=contractor))
    record = works_record(SimpleNamespace(contractor="", community=community, district="Port Loko"), log, None,
                          None, {"drl-daily-drilling-log": "no" if failed else "yes"},
                          {"logged_m": logged, "claimed_m": claimed},
                          {"lines": [{"code": "DRL", "item": "Drilling", "unit": "m", "quantity": 45.0,
                                      "rate_usd": 40.0}],
                           "variations": [{"ref": "VO-1", "code": "DRL", "quantity_delta": 5.0,
                                           "authorised_by": "A. Kamara"}], "due": due})
    return serialize_project({"project_summary": {"community": community, "district": "Port Loko", "status": status},
                              "works_record": record}, "0.3.0")


def test_the_portfolio_shows_the_contractor_scorecards(sample_data):
    """PLAN.md step 6.3: the Portfolio page sets the files' boreholes out by
    contractor, every figure with the boreholes it rests on, the checks that
    failed listed, and the whole report offered for these files only."""
    from docx import Document

    from groundwater.text import phrase

    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Portfolio"
    at.run()
    rokel = _works_project("Rokel", "Kaba Drilling", "Successful", "13th May, 2018", "2018-05-10", 60.0, 66.0, True)
    files = [("rokel.yaml", rokel, "application/x-yaml"), ("rokel_copy.yaml", rokel, "application/x-yaml"),
             ("kuntolo.yaml", _works_project("Kuntolo", "WiNGiN", "Dry hole", "05/06/2018", "2018-06-01", 50.0,
                                             50.0, False), "application/x-yaml")]
    at.file_uploader(key="portfolio_upload").set_value(files)
    at.run()
    assert not at.exception, at.exception
    assert phrase("scorecards.page_heading") in [str(h.value) for h in at.subheader]
    shown = [str(m.value) for m in at.markdown]
    # the copy of Rokel's file counts its borehole once
    assert "2 boreholes drilled by 2 contractors, from 3 project files." in shown
    assert phrase("scorecards.summary_unranked", minimum=5) in shown
    cards = next(d.value for d in at.dataframe if phrase("scorecards.col_metres") in d.value.columns)
    assert list(cards[phrase("scorecards.col_contractor")]) == ["Kaba Drilling", "WiNGiN"]
    assert list(cards[phrase("scorecards.col_metres")]) == [
        "+6 m on 1 borehole (66 m claimed, 60 m logged); 1 claimed more than 3 m over the logs",
        "0 m on 1 borehole (50 m claimed, 50 m logged)"]
    assert list(cards[phrase("scorecards.col_time")]) == ["1 of 1 finished late, 3 days in all, at most 3",
                                                         "1 date not read"]
    assert list(cards[phrase("scorecards.col_dry")]) == ["0 of 1 dry", "1 of 1 dry"]
    failed = next(d.value for d in at.dataframe if phrase("scorecards.col_failed_on") in d.value.columns)
    assert list(failed[phrase("scorecards.col_contractor")]) == ["Kaba Drilling"]

    at.button(key="build_scorecards").click()
    at.run()
    assert not at.exception, at.exception
    built = at.session_state["scorecard_report"]
    text = "\n".join(p.text for p in Document(built["path"]).paragraphs)
    assert phrase("scorecards.title") in text and phrase("scorecards.method_heading") in text
    label = "Download the contractor scorecards (.docx)"
    assert any(b.label == label for b in at.download_button)
    # other files want the report built again
    at.file_uploader(key="portfolio_upload").set_value(files[:1])
    at.run()
    assert not at.exception, at.exception
    assert not any(b.label == label for b in at.download_button)


def test_portfolio_drilldown_flow(sample_data):
    """Upload saved projects, then drill into one site for its detail and
    the one-page brief."""
    from groundwater.project_io import serialize_project

    rokel = serialize_project({"project_summary": {
        "community": "Rokel", "district": "Port Loko",
        "easting": 694667.0, "northing": 936225.0, "utm_zone": 28,
        "status": "Successful", "total_depth_m": 50.0,
        "safe_yield_m3_per_h": 2.4, "water_verdict": "pass",
        "cost_per_meter_usd": 130.0,
    }}, "0.2.0")
    kuntolo = serialize_project({"project_summary": {
        "community": "Kuntolo", "district": "Tonkolili",
        "easting": 800000.0, "northing": 950000.0, "utm_zone": 28,
        "status": "Dry hole", "total_depth_m": 45.0,
    }}, "0.2.0")

    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Portfolio"
    at.run()
    at.file_uploader(key="portfolio_upload").set_value([
        ("rokel.yaml", rokel, "application/x-yaml"),
        ("kuntolo.yaml", kuntolo, "application/x-yaml"),
    ])
    at.run()
    assert not at.exception
    # drilling into the second site shows its record and a brief to download
    at.selectbox(key="portfolio_site").select(1)
    at.run()
    assert not at.exception
    detail = str(at.table[-1].value)
    assert "Kuntolo" in detail and "Dry / failed" in detail
    assert any(b.label.startswith("Download site brief")
               for b in at.download_button)


def test_coverage_tab_csv_flow(sample_data):
    """Upload a WPDx CSV export (offline path) and get the district ranking."""
    csv_text = (
        "lat_deg,lon_deg,status_clean,water_source_clean\n"
        "7.8767,-11.1875,Functional,Borehole\n"      # Kenema
        "8.8817,-12.0442,Functional,Borehole\n"      # Bombali
        "8.8820,-12.0450,Non-Functional,Borehole\n"  # Bombali (broken)
        "8.4657,-13.2317,Functional,Borehole\n"      # Western Area Urban
        "8.0000,-14.0000,Functional,Borehole\n"      # offshore -> unassigned
    )
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Coverage gap"
    at.run()
    assert not at.exception
    # default source is the CSV upload; feed the export
    at.file_uploader(key="cov_csv").set_value(
        [("wpdx.csv", csv_text.encode(), "text/csv")]
    )
    at.run()
    assert not at.exception
    # District mode (default): ranking covers all 16 districts
    ranking = next(t.value for t in at.dataframe if "District" in t.value.columns
                   and "Rank" in t.value.columns)
    assert len(ranking) == 16
    by_district = dict(zip(ranking["District"], ranking["Functional"], strict=True))
    assert by_district["Bombali"] == 1  # one functional of the two Bombali points
    assert by_district["Kenema"] == 1
    assert any(m.label == "Districts" for m in at.metric)

    # Chiefdom mode: finer ranking with a Chiefdom column
    at.radio(key="cov_resolution").set_value("Chiefdom")
    at.run()
    assert not at.exception
    chief_ranking = next(t.value for t in at.dataframe
                         if "Chiefdom" in t.value.columns and "Rank" in t.value.columns)
    assert len(chief_ranking) == 166  # all chiefdom polygons ranked
    assert any(m.label == "Chiefdoms" for m in at.metric)


def test_coverage_says_why_when_every_row_was_discarded(sample_data):
    """An export that parsed to nothing explains itself.

    The hardest case for the page to get right is the one where the count and
    the discards disagree completely: a WPdx export whose header did not
    survive - a BOM makes the first column "\ufefflat_deg" - loses every
    coordinate, so there is nothing to rank and everything to explain.
    Reporting only the count leaves "No water points found in that source",
    which is the opposite of what the export says, and a programme officer
    reads it as "this area has no water points" rather than "this file did
    not load".
    """
    csv_text = (
        "lat_deg,lon_deg,status_clean,water_source_clean\n"
        ",,Functional,Borehole\n"
        ",,Functional,Borehole\n"
        ",,Non-Functional,Borehole\n"
    )
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Coverage gap"
    at.run()
    assert not at.exception
    at.file_uploader(key="cov_csv").set_value(
        [("wpdx.csv", csv_text.encode(), "text/csv")]
    )
    at.run()
    assert not at.exception

    # nothing survived, so there is no ranking
    assert not [t for t in at.dataframe if "Rank" in getattr(t.value, "columns", [])]

    # ...and the page says what was thrown away, not merely that it found none
    page = " ".join(
        [w.value for w in at.warning] + [e.value for e in at.error]
        + [i.value for i in at.info] + [c.value for c in at.caption]
    )
    assert "no usable latitude and longitude" in page.lower(), page


def test_waterpoints_tab_guarded(sample_data):
    """The water points tab renders and, given coordinates, shows the lookup
    control - without touching the network (no button click)."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["nav"] = "Water points"
    at.run()
    assert not at.exception
    # no coordinates yet: the tab shows its guidance, no lookup control
    assert not any(b.key == "run_waterpoints" for b in at.button)
    # with coordinates the radius slider and lookup button appear
    at.session_state["meta_easting"] = 825127.0
    at.session_state["meta_northing"] = 983069.0
    at.session_state["meta_zone"] = "28N"
    at.run()
    assert not at.exception
    assert at.slider(key="wp_radius") is not None
    assert any(b.key == "run_waterpoints" for b in at.button)


def test_recompute_on_project_load(sample_data):
    """Loading a project rebuilds the analyses from saved data sources."""
    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    assert not at.exception
    # simulate a loaded project: saved data files + the recompute flag
    at.session_state["src_ves"] = {"sample": "rokel/rokel_ves.xlsx"}
    at.session_state["src_wq"] = {"sample": "dr_timbo/dr_timbo_water_quality.xlsx"}
    at.session_state["src_log"] = {"sample": "dr_timbo/dr_timbo_drilling_log.xlsx"}
    at.session_state["src_pump"] = {"sample": "dr_timbo/dr_timbo_constant_test.xlsx"}
    at.session_state["design_swl"] = 9.44
    at.session_state["_recompute_pending"] = True
    at.run()
    assert not at.exception
    for key in ("ves_results", "wq_assessment", "borehole_design",
                "drilling_log", "pump_analysis"):
        assert key in at.session_state, f"{key} not rebuilt on load"
    # the recompute flag is consumed, so it does not run again
    assert "_recompute_pending" not in at.session_state


def test_a_saved_survey_reopens_without_inverting(sample_data, monkeypatch):
    """PLAN.md step 1.6: the inversions a project file carries are used."""
    import groundwater.recompute as recompute_module

    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    at.session_state["src_ves"] = {"sample": "rokel/rokel_ves.xlsx"}
    at.session_state["_recompute_pending"] = True
    at.run()
    assert not at.exception
    cache = at.session_state["inversion_cache"]
    first = at.session_state["ves_results"][1]
    assert len(cache) == len(first) == 2

    # a second session opening that file: the loader puts the file's cache in
    # the session beside the sources, and nothing is inverted
    def refuse(*args, **kwargs):
        raise AssertionError("a saved survey was inverted again")

    monkeypatch.setattr(recompute_module, "invert_sounding", refuse)
    again = AppTest.from_file(APP, default_timeout=600)
    again.run()
    again.session_state["src_ves"] = {"sample": "rokel/rokel_ves.xlsx"}
    again.session_state["inversion_cache"] = cache
    again.session_state["_recompute_pending"] = True
    again.run()
    assert not again.exception
    assert not again.session_state["recompute_diagnostics"]["issues"]
    restored = again.session_state["ves_results"][1]
    assert [r.fit_error_percent for r in restored] == [r.fit_error_percent for r in first]


def test_depth_spine_page(app):
    """The Depth Spine page draws whatever the project has loaded.

    It shares session state with the rest of the app, so the drilling log the
    Borehole design page parsed is the hole the spine draws, and a screen
    placed on the section becomes the project's design.
    """
    from groundwater.depth_spine.view import SpineInputs, build_view

    goto(app, "Borehole design")
    app.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    app.run()
    assert not app.exception
    log = app.session_state["drilling_log"]
    # and the page itself draws it
    goto(app, "Depth Spine")
    assert any(h.value == "Depth Spine" for h in app.header)

    view = build_view(SpineInputs(name="test", log=log))
    assert view["section"]["totalDepth"] == log.total_depth_m
    assert view["section"]["waterStrikes"] == list(log.water_strikes_m)
    # The bill of quantities comes from the same design the section draws.
    assert view["costing"]["quantityBasis"]["screenM"] == pytest.approx(
        sum(s["base"] - s["top"] for s in view["design"]["screens"])
    )


def test_the_depth_spine_prices_as_the_costing_page_does(sample_data):
    """The spine signs a price on its "Costing & BoQ" stage, so it has to be
    the Costing page's price for the design. The Streamlit spine priced at
    the catalogue rates, the default percentages and no mobilisation, so with
    the Costing page's overheads, VAT and distance changed the two disagreed.
    """
    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    assert not at.exception, at.exception
    goto(at, "Borehole design")
    at.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    at.run()
    goto(at, "Costing & BoQ")
    for key, value in (("cost_overheads", 25.0), ("cost_vat", 15.0),
                       ("cost_distance", 120.0), ("cost_samples", 2)):
        at.number_input(key=key).set_value(value)
    at.run()
    at.button(key="run_cost").click()
    at.run()
    assert not at.exception, at.exception
    estimate = at.session_state["cost_estimate"]

    goto(at, "Depth Spine")
    assert not at.exception, at.exception
    [workspace] = at.get("component_instance")
    costing = json.loads(workspace.proto.json_args)["view"]["costing"]
    assert costing["price"] == pytest.approx(estimate.price_usd, abs=0.005)
    assert costing["vat"] == pytest.approx(estimate.vat_usd, abs=0.005) and costing["vat"] > 0
    assert costing["directCost"] == pytest.approx(estimate.direct_cost_usd, abs=0.005)
    assert costing["overheadsPercent"] == 25.0


def test_the_scanned_sheets_page_says_what_it_could_not_read(sample_data):
    """An image-only PDF came back as an "unknown" sheet whose review said
    every value had been extracted, and a VES table the template could not
    take was written as a template that read back with no readings. The page
    says each, in the engine's words."""
    pytest.importorskip("pdfplumber")
    from groundwater.text import phrase
    from test_extraction import _pdf

    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    goto(at, "Scanned sheets")
    at.file_uploader(key="scan").set_value([("scan.pdf", _pdf([]), "application/pdf")])
    at.run()
    at.button(key="run_extract").click()
    at.run()
    assert not at.exception, at.exception
    assert [e.value for e in at.error] == [phrase("extraction.no_text_layer")]

    placed = [(60, 760, "SCHLUMBERGER ARRAY VES FIELD DATA"),
              (60, 650, "No."), (110, 650, "Spacing"), (200, 650, "Reading")]
    for k, (spacing, reading) in enumerate((("1.5", "120.4"), ("2", "98.1"), ("3", "80.2"))):
        placed += [(60, 630 - 18 * k, str(k + 1)), (110, 630 - 18 * k, spacing),
                   (200, 630 - 18 * k, reading)]
    at.file_uploader(key="scan").set_value([("sheet.pdf", _pdf(placed), "application/pdf")])
    at.run()
    at.button(key="run_extract").click()
    at.run()
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]
    warnings = [str(w.value) for w in at.warning]
    assert any("cannot be placed in the VES template" in w for w in warnings), warnings
    assert "1 item(s) to review" in " ".join(str(s.value) for s in at.success)


def test_the_static_level_box_shows_the_level_the_design_uses(sample_data):
    """Open the app, then load the test and the log: the box said 0.0.

    The box was prefilled only when the widget had never run, and it runs on
    the first page load, so it held 0.0 under a caption saying "Prefilled from
    the pumping test (9.44 m)" while the design used 9.44 m.
    """
    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    assert not at.exception, at.exception
    goto(at, "Pumping test")
    at.selectbox(key="sample_pump").select("dr_timbo/dr_timbo_constant_test.xlsx")
    at.run()
    goto(at, "Borehole design")
    at.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    at.run()
    assert not at.exception, at.exception

    level = at.session_state["pump_analysis"].test.static_water_level_m
    design = at.session_state["borehole_design"]
    assert design.static_water_level_m == pytest.approx(level, abs=0.01)
    assert at.number_input(key="design_swl").value == pytest.approx(level, abs=0.01)

    # a level the analyst types over it is the level, and it stays
    at.number_input(key="design_swl").set_value(12.0)
    at.run()
    assert at.number_input(key="design_swl").value == 12.0
    assert at.session_state["borehole_design"].static_water_level_m == 12.0


def test_the_design_takes_the_intake_the_pumping_report_prints(sample_data):
    """With an annual swing entered the pumping report set the intake deeper,
    for the drought year, and the design kept the test day's depth: the
    drawing and the completion report named one depth and the pumping report
    another (hydraulics-6; the browser app had the same fault)."""
    from groundwater.hydraulics import pump_intake_depth
    from groundwater.seasonal import seasonal_yield

    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    goto(at, "Pumping test")
    at.selectbox(key="sample_pump").select("dr_timbo/dr_timbo_constant_test.xlsx")
    at.run()
    at.number_input(key="seasonal_range").set_value(8.0)
    at.run()
    month = at.selectbox(key="seasonal_month").value
    goto(at, "Borehole design")
    at.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    at.run()
    assert not at.exception, at.exception

    analysis = at.session_state["pump_analysis"]
    seasonal = seasonal_yield(analysis, month=(month or None), annual_range_m=8.0)
    reported, _ = pump_intake_depth(analysis, seasonal)
    # the swing has to move the depth, or this test proves nothing
    assert reported > analysis.yield_recommendation.pump_installation_depth_m
    design = at.session_state["borehole_design"]
    assert design.pump_intake_m >= reported


def test_the_design_chooses_the_rising_main_from_the_sample(sample_data):
    """PLAN.md step 7.2: with no water quality sheet the design page says the
    rising main rests on no chemistry; with Dr Timbo's aggressive sample it
    specifies stainless steel, quoting the indices and the sample, inspected
    every 6 months, and the bill of quantities prices it over the intake."""
    from groundwater.text import phrase

    at = AppTest.from_file(APP, default_timeout=600)
    at.run()
    goto(at, "Pumping test")
    at.selectbox(key="sample_pump").select("dr_timbo/dr_timbo_constant_test.xlsx")
    at.run()
    goto(at, "Borehole design")
    at.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    at.run()
    assert not at.exception, at.exception
    design = at.session_state["borehole_design"]
    assert design.rising_main["chemistry"] == "none"
    warned = " ".join(str(w.value) for w in at.warning)
    assert phrase("rising_main.no_sample") in warned
    assert [m.value for m in at.metric if m.label == "Rising main"] == ["galvanised iron"]

    goto(at, "Water quality")
    at.selectbox(key="sample_wq").select("dr_timbo/dr_timbo_water_quality.xlsx")
    at.run()
    assert not at.exception, at.exception
    # the design follows the sample without the design page being opened
    choice = at.session_state["borehole_design"].rising_main
    assert choice["upgraded"] and choice["material"] == "stainless steel"
    goto(at, "Borehole design")
    at.run()
    assert not at.exception, at.exception
    design = at.session_state["borehole_design"]
    choice = design.rising_main
    assert choice["chemistry"] == "aggressive" and choice["inspection_months"] == 6
    warned = " ".join(str(w.value) for w in at.warning)
    assert choice["statement"] in warned
    assert "WQ-BH1-01 of 20th May, 2018" in warned and "Ryznar index 11.37" in warned
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["Rising main"] == "stainless steel"
    assert metrics["Rising main inspection"] == "every 6 months"

    goto(at, "Costing & BoQ")
    at.button(key="run_cost").click()
    at.run()
    assert not at.exception, at.exception
    estimate = at.session_state["cost_estimate"]
    line = next(i for i in estimate.items if i.code == "RSM1")
    assert line.quantity == design.pump_intake_m


def test_a_sanitary_inspection_is_recorded_and_read_with_the_sample(sample_data):
    """PLAN.md step 7.5: the form on the Asset registry page records an
    inspection in the borehole's history, with the questions left unanswered
    marked so, and the statement under it reads that inspection with the
    E. coli count of the water quality sample on the Water quality page."""
    from datetime import date

    from groundwater.ingestion import read_quality_workbook
    from groundwater.quality import assess_sample
    from groundwater.registry import Asset, AssetEvent, asset_from_dict
    from groundwater.sanitary import sanitary_risk

    asset = Asset(asset_id="SL-WAR-8FEEVKQ-T", community="Dr. Timbo's",
                  district="Western Area Rural", easting=694912.0, northing=938150.0,
                  utm_zone=28, events=[AssetEvent("2020-01-10", "commissioned")])
    quality = assess_sample(read_quality_workbook(
        REPO / "examples" / "data" / "dr_timbo" / "dr_timbo_water_quality.xlsx"))
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["asset_record"] = asset.as_dict()
    at.session_state["wq_assessment"] = quality
    goto(at, "Asset registry")
    shown = " ".join(str(i.value) for i in at.info)
    assert "No sanitary inspection with the form filled in is on record" in shown

    at.date_input(key="sanitary_when").set_value(date(2024, 5, 20))
    at.text_input(key="sanitary_by").set_value("A. Bangura")
    for number, answer in ((1, "Y"), (2, "N"), (3, "N"), (4, "Y"), (5, "N"), (6, "N"),
                           (7, "Y"), (8, "N")):
        at.radio(key=f"sanitary_q{number}").set_value(answer)
    next(b for b in at.button if b.label == "Record the inspection").click()
    at.run()
    assert not at.exception, at.exception

    recorded = asset_from_dict(at.session_state["asset_record"])
    inspections = [e for e in recorded.events if e.kind == "inspection"]
    assert [(e.when, e.by, e.answers) for e in inspections] == [
        ("2024-05-20", "A. Bangura", "YNNYNNYN--")]
    expected = sanitary_risk(recorded, date.today(), quality)
    # three risks found and two questions open: 3 to 5, intermediate either
    # way, so the matrix gives one priority even for an incomplete form
    assert (expected["action"], expected["settled"]) == ("higher", True)
    assert "between 3 and 5 of 10, which the form rates intermediate" in expected["statement"]
    written = " ".join(str(w.value) for w in at.warning)
    assert expected["statement"] in written


def test_the_service_rounds_follow_the_design_and_the_mechanic(sample_data):
    """PLAN.md step 7.3: the Asset registry page schedules this borehole's
    routine inspection at the interval the design sets for its rising main
    (Dr Timbo's aggressive sample: stainless steel, every 6 months), puts it
    on the round of the mechanic named for its chiefdom until one is named on
    its record, and builds the round sheet."""
    import docx

    from groundwater.registry import Asset, AssetEvent, asset_from_dict
    from groundwater.text import phrase, phrase_table

    asset = Asset(asset_id="SL-WAR-8FEEVKQ-T", community="Dr. Timbo's",
                  district="Western Area Rural", chiefdom="Koya", easting=694912.0,
                  northing=938150.0, utm_zone=28,
                  events=[AssetEvent("2020-01-10", "commissioned")])
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["asset_record"] = asset.as_dict()
    goto(at, "Borehole design")
    at.selectbox(key="sample_log").select("dr_timbo/dr_timbo_drilling_log.xlsx")
    at.run()
    goto(at, "Water quality")
    at.selectbox(key="sample_wq").select("dr_timbo/dr_timbo_water_quality.xlsx")
    at.run()
    assert not at.exception, at.exception
    goto(at, "Asset registry")
    basis = phrase("rounds.basis_design", months=6, material="stainless steel",
                   water=phrase_table("rounds.water")["aggressive"])
    captions = " ".join(str(c.value) for c in at.caption)
    assert phrase("rounds.interval_line", title="Routine inspection", basis=basis) in captions
    # the record saved with the project carries the design's rising main
    saved = asset_from_dict(at.session_state["asset_record"])
    assert saved.rising_main == {"material": "stainless steel", "inspection_months": 6,
                                 "chemistry": "aggressive"}

    at.text_area(key="org_area_mechanics").input("koya: M. Kolleh\nno colon")
    at.run()
    assert not at.exception, at.exception
    assert "M. Kolleh: 1 stop." in " ".join(str(m.value) for m in at.markdown)
    assert phrase("rounds.mechanics_no_colon", line=2) in " ".join(str(w.value) for w in at.warning)

    next(t for t in at.text_input if t.label == "Area mechanic").input("A. Bangura")
    next(b for b in at.button if b.label == "Name the area mechanic").click()
    at.run()
    assert not at.exception, at.exception
    assert asset_from_dict(at.session_state["asset_record"]).area_mechanic == "A. Bangura"
    assert "A. Bangura: 1 stop." in " ".join(str(m.value) for m in at.markdown)

    at.button(key="rounds_sheet").click()
    at.run()
    assert not at.exception, at.exception
    sheet = docx.Document(at.session_state["artifacts"]["Round sheets"])
    text = "\n".join(p.text for p in sheet.paragraphs)
    assert "Area mechanic: A. Bangura" in text and basis in text
    cells = [c.text for t in sheet.tables for row in t.rows for c in row.cells]
    assert "SL-WAR-8FEEVKQ-T" in cells and "BOREHOLE SL-WAR-8FEEVKQ-T" in cells


def test_how_long_boreholes_keep_working_on_the_registry_page():
    """PLAN.md step 7.4: the Asset registry page reads this borehole and the
    register of saved project files as a time to the first failure: the
    curve over them all and by pump type with its tables, a group of too few
    boreholes said to be too few, and the pump type named on this borehole's
    record joining it to its group."""
    from datetime import date

    import yaml

    from groundwater.models import SiteMetadata
    from groundwater.project_io import serialize_project
    from groundwater.registry import Asset, AssetEvent, asset_from_dict, mint_asset_id
    from groundwater.survival import survival_analysis, survival_rows, survival_text
    from groundwater.text import phrase

    files = []
    for i in range(12):
        easting = 690000.0 + 1000.0 * i
        days = 300 + 90 * i if i < 6 else 1500 + 60 * i
        when = date.fromordinal(date(2018, 1, 10).toordinal() + days).isoformat()
        record = Asset(
            asset_id=mint_asset_id(SiteMetadata(district="Western Area Rural", easting=easting,
                                                northing=938150.0, utm_zone=28)),
            community=f"Site {i}", district="Western Area Rural", easting=easting,
            northing=938150.0, utm_zone=28,
            pump_type="Afridev" if i < 6 else "India Mark II",
            events=[AssetEvent("2018-01-10", "commissioned"),
                    AssetEvent(when, "failure" if i % 2 else "routine_inspection")])
        files.append((f"site_{i}.yaml", serialize_project({"asset_record": record.as_dict()},
                                                          "0.3.0")))
    this = Asset(asset_id="SL-WAR-8FEEVKQ-T", community="Dr. Timbo's",
                 district="Western Area Rural", easting=694912.0, northing=938150.0,
                 utm_zone=28, events=[AssetEvent("2020-01-10", "commissioned"),
                                      AssetEvent("2022-03-01", "failure")])
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["asset_record"] = this.as_dict()
    at.session_state["registry_upload_kept"] = {"ids": [name for name, _ in files],
                                                "files": files}
    goto(at, "Asset registry")
    at.date_input(key="survival_date").set_value(date(2024, 6, 1))
    at.run()
    assert not at.exception, at.exception

    register = [asset_from_dict(this.as_dict())] + [
        asset_from_dict(yaml.safe_load(blob)["asset"]) for _, blob in files]
    result = survival_analysis(register, "2024-06-01")
    words = survival_text(result)
    written = " ".join(str(m.value) for m in at.markdown)
    assert words["summary"][0] in written and words["overall"]["lines"][0] in written
    pump = next(g for g in words["groupings"] if g["by"] == "pump_type")
    assert all(line in written for line in pump["lines"])
    # this borehole names no pump yet, so it is a group of one, too few
    assert phrase("survival.too_few", label=phrase("survival.not_recorded"), n=1, minimum=5,
                  failed=1, censored=0) in written
    frames = [element.value for element in at.dataframe]
    assert any(list(frame.iloc[0]) == survival_rows([result["overall"]])[0] for frame in frames)
    assert len(at.get("image")) >= 2

    # naming its pump joins it to its group
    next(t for t in at.text_input if t.label == "Pump type").input(" India Mark II ")
    next(b for b in at.button if b.label == "Save the pump and contractor").click()
    at.run()
    assert not at.exception, at.exception
    assert asset_from_dict(at.session_state["asset_record"]).pump_type == "India Mark II"
    register[0] = asset_from_dict(at.session_state["asset_record"])
    pump = next(g for g in survival_text(survival_analysis(register, "2024-06-01"))["groupings"]
                if g["by"] == "pump_type")
    written = " ".join(str(m.value) for m in at.markdown)
    assert pump["lines"][0].startswith("India Mark II: 7 followed") and pump["lines"][0] in written


def test_reports_by_sms_and_from_a_phone_on_the_registry_page(monkeypatch):
    """PLAN.md step 7.1: the Streamlit app has no report page and says so on
    its Asset registry page. It imports SMS messages and a phone's file of
    reports, shows what became of each line and why, holds a report the
    registry reads only with a note until a person ticks it, adds this
    borehole's reports to its history once, offers the rest as a file, and
    holds the plate and the round sheets to the report address."""
    import docx

    from groundwater import fault_reports as F
    from groundwater import registry
    from groundwater.models import SiteMetadata
    from groundwater.registry import Asset, AssetEvent, asset_from_dict, mint_asset_id
    from groundwater.text import phrase

    def minted(district, easting, northing, zone):
        return mint_asset_id(SiteMetadata(district=district, easting=easting,
                                          northing=northing, utm_zone=zone))

    this = minted("Western Area Rural", 694912.0, 938150.0, 28)
    other = minted("Bo", 790500.0, 875300.0, 28)
    noted = minted("Kailahun", 280400.0, 925600.0, 29)
    # the registry's check stands in for one that reads the third identifier
    # with a note, as it reads one minted with the older check character
    real = registry.validate_asset_id

    def check(text):
        ok, reason = real(text)
        if ok and registry.parse_asset_id(text) == noted:
            return True, "Compare it with the plate."
        return ok, reason

    monkeypatch.setattr(registry, "validate_asset_id", check)
    sms = "\n".join([
        f"2026-05-01 08:14 +23276123456 BH {this} BROKEN handle snapped",
        f"BH {this} SLOW",
        "hello",
        f"2026-05-02 +232 BH {other} WORKING",
        f"2026-05-02 +232 BH {noted} TASTE salty",
    ]).encode("utf-8")
    phone = json.dumps(F.reports_file([F.report_record(
        this, "slow", "2026-05-04", "slow since Monday", "Aminata")])).encode("utf-8")
    at = AppTest.from_file(APP, default_timeout=600)
    at.session_state["asset_record"] = Asset(
        asset_id=this, community="Dr. Timbo's", district="Western Area Rural",
        easting=694912.0, northing=938150.0, utm_zone=28,
        events=[AssetEvent("2020-01-10", "commissioned")]).as_dict()
    at.session_state["reports_upload_kept"] = {
        "ids": ["sms.txt", "phone.json"], "files": [("sms.txt", sms), ("phone.json", phone)]}
    goto(at, "Asset registry")
    assert not at.exception, at.exception
    assert phrase("fault_reports.browser_only") in " ".join(str(c.value) for c in at.caption)

    # each line read, a message with no day refused rather than dated today
    written = " ".join(str(m.value) for m in at.markdown)
    assert "6 read: 3 accepted, 1 held for a person to confirm, 2 refused." in written
    table = next(f.value for f in at.dataframe if list(f.value.columns) == F.report_header())
    assert list(table["Outcome"]) == ["Accepted", "Refused", "Refused", "Accepted",
                                      "Held for a person to confirm", "Accepted"]
    assert table["Why"][1] == phrase("fault_reports.no_date")
    assert table["Sender"][0] == "+23276123456"
    held = phrase("fault_reports.held", note="Compare it with the plate.")
    assert held in " ".join(str(w.value) for w in at.warning)

    # this borehole's two reports, added to its history at the press, once
    add = next(b for b in at.button if b.key == "reports_add")
    assert add.label == "Add 2 report(s) to this borehole's history"
    add.click()
    at.run()
    assert not at.exception, at.exception
    events = asset_from_dict(at.session_state["asset_record"]).events
    assert [(e.when, e.kind, e.by) for e in events] == [
        ("2020-01-10", "commissioned", ""),
        ("2026-05-01", "failure", "SMS from +23276123456"),
        ("2026-05-04", "reported_slow", "Aminata")]
    assert "2 report(s) added" in " ".join(str(s.value) for s in at.success)
    assert not [b for b in at.button if b.key == "reports_add"]
    assert "history holds the reports that name it" in " ".join(str(c.value) for c in at.caption)

    # the rest are offered as a file, the held one only once a person ticks it
    infos = " ".join(str(i.value) for i in at.info)
    assert "1 imported report(s) name boreholes" in infos
    next(c for c in at.checkbox if c.label.startswith(phrase("fault_reports.confirm"))).check()
    at.run()
    assert not at.exception, at.exception
    assert "2 imported report(s) name boreholes" in " ".join(str(i.value) for i in at.info)

    # a report address the registry refuses holds back the plate and the
    # round sheets; one it takes makes the plate's symbol a link
    at.text_input(key="org_report_url").input("http://water.example.org/")
    at.run()
    assert not at.exception, at.exception
    refused = phrase("fault_reports.url_refused", url="http://water.example.org/")
    assert sum(str(e.value) == refused for e in at.error) == 2
    assert at.button(key="asset_placard").disabled and at.button(key="rounds_sheet").disabled
    at.text_input(key="org_report_url").input("https://water.example.org/toolkit/")
    at.run()
    assert not at.exception, at.exception
    assert not at.button(key="asset_placard").disabled
    at.button(key="asset_placard").click()
    at.run()
    assert not at.exception, at.exception
    plate = docx.Document(at.session_state["artifacts"]["Borehole identification plate"])
    text = "\n".join(p.text for p in plate.paragraphs)
    assert phrase("fault_reports.placard_caption_link") in text
    assert phrase("fault_reports.placard_caption_text") not in text
