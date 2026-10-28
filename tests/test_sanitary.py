"""Sanitary inspection (PLAN.md step 7.5): WHO's form for a borehole with a
handpump, scored, and read with the E. coli count into one risk statement."""

from __future__ import annotations

import json
import random
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
import yaml

from groundwater import sanitary as SAN
from groundwater.models import SiteMetadata, WaterQualityResult, WaterQualitySample
from groundwater.quality.assess import assess_sample
from groundwater.registry import Asset, AssetEvent, asset_from_dict, clean_answers, event_id, merge_events
from groundwater.text import phrase, phrase_table

REPO = Path(__file__).resolve().parents[1]
JS = REPO / "docs" / "js"
ASSET = "SL-WAR-8FEEVKQ-T"
TODAY = "2024-07-01"


def _asset(*events) -> Asset:
    return Asset(asset_id=ASSET, community="Dr. Timbo's", district="Western Area Rural",
                 events=list(events))


def _sample(*results, when="2024-06-30") -> WaterQualitySample:
    return WaterQualitySample(site=SiteMetadata(community="Ref"), sample_date=when,
                              results=[WaterQualityResult(**r) for r in results])


_PANEL = [{"parameter": "Arsenic", "value": 0.001, "unit": "mg/L"},
          {"parameter": "Fluoride", "value": 0.3, "unit": "mg/L"},
          {"parameter": "Nitrate (as NO3)", "value": 5.0, "unit": "mg/L"}]


def _assessed(*ecoli, when="2024-06-30"):
    return assess_sample(_sample(*ecoli, *_PANEL, when=when))


def _ecoli(value=None, unit="CFU/100 mL", parameter="E. coli", **kw) -> dict:
    return {"parameter": parameter, "value": value, "unit": unit, **kw}


# ------------------------------------------------------------------- the form

def test_the_form_is_ten_questions_and_its_bands_cover_every_score():
    questions = SAN.form_questions()
    assert len(questions) == 10 and all(q.endswith("?") for q in questions)
    assert "latrine within 10 m" in questions[0]
    bands = [SAN.score_band(score) for score in range(11)]
    assert bands == ["low"] * 3 + ["intermediate"] * 3 + ["high"] * 3 + ["very_high"] * 2
    with pytest.raises(ValueError):
        SAN.score_band(11)


def test_the_score_is_the_number_of_risks_present():
    score = SAN.score_answers("YNNYNNYNYN")
    assert score == {"form": "tubewell_handpump", "n": 10, "answered": 10, "yes": 4,
                     "low": 4, "high": 4, "band_low": "intermediate",
                     "band_high": "intermediate", "complete": True}


def test_a_question_left_unanswered_is_not_a_no():
    """Three questions open: the score may be anything from the yes answers to
    those plus three, and the form says it is incomplete."""
    score = SAN.score_answers("YNN--NYN-N")
    assert (score["answered"], score["yes"], score["low"], score["high"]) == (7, 2, 2, 5)
    assert (score["band_low"], score["band_high"], score["complete"]) == ("low", "intermediate", False)


def test_answers_that_do_not_fit_the_form_score_nothing():
    assert SAN.score_answers("YNY") is None
    assert SAN.score_answers("YNNYNNYNYNN") is None
    assert SAN.score_answers("YNNYNNYNYN", form="hand_dug_well") is None


def test_an_answer_nobody_can_read_is_an_unanswered_question():
    assert clean_answers("yNnY") == "YNNY"
    assert clean_answers("Y?N x") == "Y-N--"
    assert clean_answers(None) == "" and clean_answers("  ") == ""
    # by code point, as the browser now reads it too
    assert clean_answers("Y\U0001f600N") == "Y-N"


@pytest.mark.parametrize("raw", [
    "\ufeffYNNYNNYNYN",      # a byte order mark: JavaScript's trim took it, Python's did not
    "YNNYNNYNYN\x1c",        # a file separator: Python's strip took it, JavaScript's did not
    "\x85YNNYNNYNYN\x1f",
    "\xa0YNNYNNYNYN\u3000",
    " \tYNNYNNYNYN\r\n",
])
def test_answers_are_trimmed_alike_in_both_engines(raw):
    """The ends are trimmed of what either engine's own trim takes, so the
    same stored answers score in both; a character neither trims (a zero
    width space) is still an answer nobody can read."""
    assert clean_answers(raw) == "YNNYNNYNYN"
    assert SAN.score_answers(raw)["yes"] == 4
    assert clean_answers("\u200bYNNYNNYNYN") == "-YNNYNNYNYN"


def test_the_trim_is_the_union_of_both_engines_trims():
    """trim_field takes exactly what str.strip or String.trim would take; a
    zero width space, which neither takes, is kept."""
    import sys

    from groundwater.registry import trim_field

    python = {c for c in range(sys.maxunicode + 1) if (chr(c) + "x").strip() == "x"}
    trimmed = {c for c in range(sys.maxunicode + 1) if trim_field(chr(c) + "x") == "x"}
    node = shutil.which("node")
    if node is None:
        javascript = python - set(range(0x1C, 0x20)) - {0x85} | {0xFEFF}
    else:
        script = ("const o = []; for (let c = 0; c <= 0x10FFFF; c++) {"
                  " if ((String.fromCodePoint(c) + 'x').trim() === 'x') o.push(c); }"
                  " process.stdout.write(JSON.stringify(o));")
        done = subprocess.run([node, "-e", script], capture_output=True, text=True,
                              check=True)
        javascript = set(json.loads(done.stdout))
    # each engine's own trim took a few the other did not
    assert python - javascript == set(range(0x1C, 0x20)) | {0x85}
    assert javascript - python == {0xFEFF}
    assert trimmed == python | javascript
    assert trim_field(None) == "" and trim_field("\u200b x \u200b") == "\u200b x \u200b"


def test_the_form_rows_name_each_answer():
    rows = SAN.form_rows("Y-NNNNNNNN")
    words = phrase_table("sanitary.answer")
    assert rows[0] == ["1", SAN.form_questions()[0], words["Y"]]
    assert rows[1][2] == words["-"] and rows[2][2] == words["N"] and len(rows) == 10


# ----------------------------------------------------------------- the grades

@pytest.mark.parametrize("count,grade", [
    (0, "A"), (0.5, "B"), (1, "B"), (10, "B"), (10.5, "C"), (11, "C"), (100, "C"),
    (101, "D"), (1000, "D"), (1000.5, "E"), (25000, "E")])
def test_a_count_is_graded_as_who_grades_it(count, grade):
    assert SAN.ecoli_grade(count) == grade


@pytest.mark.parametrize("bound,inclusive,grades", [
    (0, False, ("B", "E")), (0, True, ("A", "E")), (10, False, ("C", "E")),
    (10, True, ("B", "E")), (50, False, ("C", "E")), (1000, False, ("E", "E")),
    (1000, True, ("D", "E")), (5000, False, ("E", "E"))])
def test_a_count_not_put_to_a_number_allows_the_grades_above_it(bound, inclusive, grades):
    assert SAN.ecoli_grades_above(bound, inclusive) == grades


# ----------------------------------------------------------------- the matrix

def test_the_matrix_never_lowers_the_priority_as_either_side_worsens():
    tables = SAN.sanitary_tables()
    matrix = tables["forms"]["tubewell_handpump"]["matrix"]
    order = [g["grade"] for g in tables["grades"]]
    rank = SAN.PRIORITIES.index
    for i, grade in enumerate(order):
        row = [rank(cell) for cell in matrix[grade]]
        assert row == sorted(row), grade
        if i:
            above = [rank(cell) for cell in matrix[order[i - 1]]]
            assert all(a <= b for a, b in zip(above, row, strict=True)), grade


def test_each_cell_is_the_worse_of_the_two_sides_as_the_file_says():
    """sanitary_risk.yaml says its cells are a choice: the worse of the
    grade's own WHO risk word and the score band's. Held to it here, so a
    cell corrected against WHO's figure is a deliberate change to this test
    and to the file's comment together."""
    grade_level = {"A": 0, "B": 1, "C": 2, "D": 3, "E": 4}
    band_level = {"low": 1, "intermediate": 2, "high": 3, "very_high": 4}
    priority = {1: "low", 2: "higher", 3: "higher", 4: "urgent"}
    matrix = SAN.sanitary_tables()["forms"]["tubewell_handpump"]["matrix"]
    for grade, level in grade_level.items():
        for score in range(11):
            worse = max(level, band_level[SAN.score_band(score)])
            assert matrix[grade][score] == priority[worse], (grade, score)


# ------------------------------------------------------------- the microbiology

def test_no_assessment_is_no_sample():
    assert SAN.microbiology(None) is None


@pytest.mark.parametrize("result,kind,grades,count", [
    (_ecoli(0.0), "counted", ("A", "A"), 0.0),
    (_ecoli(7.0), "counted", ("B", "B"), 7.0),
    (_ecoli(2.0, unit="CFU/mL"), "counted", ("D", "D"), 200.0),
    (_ecoli(None, below_detection=True), "none_found", ("A", "A"), 0.0),
    (_ecoli(None, below_detection=True, detection_limit=1.0), "none_found", ("A", "A"), 0.0),
    (_ecoli(None, greater_than=50.0), "bound", ("C", "E"), None),
    (_ecoli(None, greater_than=0.0), "detected", ("B", "E"), None),
    (_ecoli(40.0, parameter="Thermotolerant coliforms"), "counted", ("C", "C"), 40.0),
    (_ecoli(None, below_detection=True, detection_limit=10.0), "ungraded", (None, None), None),
])
def test_the_e_coli_row_is_graded_from_the_assessment(result, kind, grades, count):
    micro = SAN.microbiology(_assessed(result))
    assert micro["date"] == "2024-06-30" and micro["kind"] == kind
    assert (micro["grade_low"], micro["grade_high"]) == grades
    assert micro["count"] == count
    if kind == "ungraded":
        assert "detection limit" in micro["reason"]


def test_a_count_whose_unit_is_not_per_100_ml_is_not_graded():
    """Only a standards table writing E. coli in another unit lets an
    evaluable row reach this, so it is held here rather than in the seeds."""
    from types import SimpleNamespace

    row = SimpleNamespace(parameter="E. coli", status="exceeds_health", evaluable=True,
                          remark="", unit="colonies per plate", value=5.0, greater_than=None,
                          greater_than_inclusive=False)
    micro = SAN.microbiology(SimpleNamespace(sample=None, rows=[row]))
    assert (micro["kind"], micro["grade_low"], micro["unit"]) == ("unit", None, "colonies per plate")
    risk = SAN.sanitary_risk(_asset(), TODAY, SimpleNamespace(sample=None, rows=[row],
                                                              verdict_state="health_fail"))
    assert "its unit, colonies per plate, is not a count per 100 mL" in risk["statement"]
    assert "date not recorded" in risk["statement"]


@pytest.mark.parametrize("result", [
    _ecoli(3.0, unit="/100 mL"),
    _ecoli(None, unreadable="lots"),
    _ecoli(None, below_detection=True, detection_limit=1.0, unit="CFU/mL"),
    _ecoli(None, below_detection=True, detection_limit=10.0),
])
def test_a_result_that_could_not_be_graded_ends_its_sentence_once(result):
    """The assessment's remark is a sentence with its own full stop, or
    not; the statement adds one, so the remark's is taken off."""
    risk = _risk(SAN.inspection_event("2024-06-30", "NNNNNNNNNN"), assessment=_assessed(result))
    micro = risk["microbiology"]
    assert micro["kind"] == "ungraded" and micro["reason"]
    assert not micro["reason"].endswith(".")
    assert ".." not in risk["statement"]
    assert f"could not be graded: {micro['reason']}." in risk["statement"]


def test_a_result_with_no_reason_is_still_one_sentence(tmp_path):
    """No assessment the toolkit makes leaves an ungraded row without a
    remark, so the row is made by hand, and the browser is given the same."""
    from types import SimpleNamespace

    fields = {"parameter": "E. coli", "status": "indeterminate", "evaluable": False,
              "remark": " . ", "unit": "CFU/100 mL", "value": None, "greater_than": None,
              "greater_than_inclusive": False}
    risk = SAN.sanitary_risk(_asset(), TODAY, SimpleNamespace(
        sample=None, rows=[SimpleNamespace(**fields)], verdict_state=None))
    assert risk["microbiology"]["reason"] == ""
    assert phrase("sanitary.ecoli_ungraded_bare", parameter="E. coli",
                  date="date not recorded") in risk["statement"]
    case = {"asset_id": ASSET, "events": [], "sample": None, "today": TODAY,
            "assessment": {"sample": None, "rows": [fields], "verdict_state": None},
            "answers": [], "counts": []}
    answer = _browser([case], tmp_path)[0]
    assert "threw" not in answer, answer.get("threw")
    assert answer["ok"]["risk"] == json.loads(json.dumps(risk))


def test_a_sample_without_e_coli_says_so_and_the_worst_row_is_used():
    micro = SAN.microbiology(assess_sample(_sample(*_PANEL)))
    assert micro["kind"] == "missing" and micro["grade_low"] is None
    worst = SAN.microbiology(_assessed(_ecoli(0.0), _ecoli(150.0, parameter="Faecal coliforms")))
    assert (worst["parameter"], worst["grade_low"]) == ("Faecal coliforms", "D")


# -------------------------------------------------------------- the statement

def _risk(*events, assessment=None, today=TODAY):
    return SAN.sanitary_risk(_asset(*events), today, assessment)


def test_a_complete_form_and_a_clean_count_read_low():
    risk = _risk(SAN.inspection_event("2024-06-30", "NNYNNNNNNN", by="A. Bangura"),
                 assessment=_assessed(_ecoli(0.0)))
    assert (risk["action"], risk["settled"]) == ("low", True)
    statement = risk["statement"]
    assert "Sanitary inspection of 2024-06-30 by A. Bangura: 1 of 10 risks present" in statement
    assert "E. coli in the water sample (2024-06-30): 0 per 100 mL, grade A" in statement
    assert phrase("sanitary.together", priority=phrase_table("sanitary.priority")["low"]) in statement
    assert statement.endswith(phrase("sanitary.verdict", verdict="Safe to drink"))


def test_a_heavy_count_is_urgent_whatever_the_inspection_says():
    """Row E of the matrix is urgent throughout, so even with no inspection
    on record the priority is read, and the statement says the sanitary
    side is not known."""
    risk = _risk(assessment=_assessed(_ecoli(2400.0)))
    assert (risk["action"], risk["settled"]) == ("urgent", True)
    assert phrase("sanitary.inspection_none") in risk["statement"]
    assert "grade E (very high risk)" in risk["statement"]


def test_with_no_sample_the_microbiological_side_is_unknown_and_no_priority_is_read():
    risk = _risk(SAN.inspection_event("2024-06-30", "NNNNNNNNNN"))
    assert risk["action"] is None and not risk["settled"]
    assert phrase("sanitary.no_sample") in risk["statement"]
    assert phrase("sanitary.together_unknown") in risk["statement"]


def test_a_sample_the_history_records_is_named_when_its_analysis_is_not_loaded():
    """The History table beside the statement lists the sample, so the
    statement does not say there is none: it names it, and says its results
    are not loaded. A sample dated after today is not named."""
    events = (AssetEvent("2024-03-01", "water_sample", "E. coli 12 per 100 mL field kit"),
              AssetEvent("2024-05-02", "water_sample", by="M. Kolleh"),
              AssetEvent("2024-09-01", "water_sample"),
              SAN.inspection_event("2024-06-01", "YNNYNNYNYN"))
    risk = _risk(*events)
    statement = risk["statement"]
    assert phrase("sanitary.no_sample") not in statement
    assert phrase("sanitary.sample_not_loaded", when="2024-05-02 by M. Kolleh") in statement
    assert risk["sample_event"]["when"] == "2024-05-02" and risk["action"] is None
    assert "No water sample is on record" not in statement
    # with no sample in the history the sentence says both are missing
    assert phrase("sanitary.no_sample") in _risk(events[3])["statement"]
    assert _risk(events[3])["sample_event"] is None


def test_a_later_sample_than_the_one_read_is_named():
    later = AssetEvent("2024-06-20", "water_sample", by="A. Bangura")
    risk = _risk(later, assessment=_assessed(_ecoli(3.0), when="2024-05-30"))
    assert phrase("sanitary.sample_later", when="2024-06-20 by A. Bangura") in risk["statement"]
    # the same day, an earlier one, or an analysis whose date cannot be read: not named
    for when, sampled in (("2024-06-20", "2024-06-20"), ("2024-06-30", "2024-06-20"),
                          ("", "2024-06-20"), ("20th June", "2024-06-20")):
        quiet = _risk(AssetEvent(sampled, "water_sample"),
                      assessment=_assessed(_ecoli(3.0), when=when))
        assert "later water sample" not in quiet["statement"], when


@pytest.mark.parametrize("written", [
    "20240520", "2024-W21-1", "2024W211", "2024-05-20T09:00", "2024-5-20",
    "\u0968\u0966\u0968\u096a-\u0966\u096b-\u0968\u0966", "0000-05-20", "2024-02-30", "+2024-05-20"])
def test_a_date_is_read_only_as_yyyy_mm_dd(written):
    """Python 3.11 reads compact and week dates that the browser and Python
    3.10 do not; both engines now read YYYY-MM-DD and nothing else, so an
    inspection on such a date is not read in either, and says so."""
    from groundwater.registry import iso_date

    assert iso_date(written) is None and AssetEvent(written, "repair").date is None
    risk = _risk(SAN.inspection_event(written, "YYYYYYYYYN"))
    assert risk["inspection"] is None and risk["action"] is None
    assert phrase("sanitary.inspection_undated_only", count=1) in risk["statement"]
    with pytest.raises(ValueError, match="today's date"):
        _risk(today=written)


def test_a_date_with_blanks_at_its_ends_is_read_and_an_early_year_is_its_own():
    from groundwater.registry import iso_date

    assert iso_date("\ufeff2024-05-20\x1c ") == date(2024, 5, 20)
    assert iso_date("0024-05-20") == date(24, 5, 20)
    risk = _risk(SAN.inspection_event("2024-06-30", "NNNNNNNNNN"),
                 AssetEvent("30/06/2024", "inspection", form="tubewell_handpump",
                            answers="YYYYYYYYYY"),
                 assessment=_assessed(_ecoli(0.0)))
    assert risk["inspection"]["when"] == "2024-06-30" and risk["action"] == "low"
    assert phrase("sanitary.inspection_undated", count=1) in risk["statement"]


def test_a_known_side_still_says_how_bad_it_is_at_least():
    # every grade agrees at 9 of 10, so the priority is read with no sample
    risk = _risk(SAN.inspection_event("2024-06-30", "YYYYYYYYYN"))
    assert (risk["action"], risk["settled"]) == ("urgent", True)
    assert phrase("sanitary.no_sample") in risk["statement"]
    floor = _risk(SAN.inspection_event("2024-06-30", "YYYYNNNNNN"))
    assert floor["action"] is None
    assert phrase("sanitary.together_floor",
                  priority=phrase_table("sanitary.priority")["higher"]) in floor["statement"]


def test_an_incomplete_form_is_taken_at_its_worse_end():
    risk = _risk(SAN.inspection_event("2024-06-30", "YNN--NYN-N"),
                 assessment=_assessed(_ecoli(0.0)))
    assert (risk["priority_low"], risk["priority_high"], risk["action"]) == ("low", "higher", "higher")
    assert "between 2 and 5 of 10, which the form rates low to intermediate" in risk["statement"]
    assert "not counted as a no" in risk["statement"]


def test_a_count_not_put_to_a_number_is_taken_at_its_worse_end():
    risk = _risk(SAN.inspection_event("2024-06-30", "NNNNNNNNNN"),
                 assessment=_assessed(_ecoli(None, greater_than=50.0)))
    assert risk["action"] == "urgent" and not risk["settled"]
    assert "more than 50 per 100 mL, not counted, so grade C to E" in risk["statement"]


def test_the_latest_scored_inspection_is_read_and_a_later_blank_one_is_named():
    risk = _risk(SAN.inspection_event("2023-01-10", "YYYYYYNNNN"),
                 SAN.inspection_event("2024-02-01", "YNNNNNNNNN"),
                 AssetEvent("2024-05-01", "inspection", "apron looked fine"),
                 SAN.inspection_event("2099-01-01", "YYYYYYYYYY"),
                 assessment=_assessed(_ecoli(0.0)))
    assert risk["inspection"]["when"] == "2024-02-01" and risk["action"] == "low"
    assert phrase("sanitary.inspection_later", when="2024-05-01") in risk["statement"]


def test_two_forms_on_one_day_are_read_at_the_one_showing_most_risk():
    risk = _risk(SAN.inspection_event("2024-06-30", "NNNNNNNNNN", by="A"),
                 SAN.inspection_event("2024-06-30", "YYYNNNNNNN", by="B"),
                 assessment=_assessed(_ecoli(0.0)))
    assert risk["inspection"]["by"] == "B" and risk["score"]["yes"] == 3
    assert phrase("sanitary.inspection_several", count=2, date="2024-06-30") in risk["statement"]


def test_a_blank_or_unreadable_form_gives_no_score():
    blank = _risk(SAN.inspection_event("2024-06-30", "----------"))
    assert blank["score"] is None
    assert phrase("sanitary.inspection_blank", when="2024-06-30") in blank["statement"]
    wrong = _risk(AssetEvent("2024-06-30", "inspection", form="tubewell_handpump", answers="YN"))
    assert phrase("sanitary.inspection_unreadable", when="2024-06-30") in wrong["statement"]


def test_an_undated_sample_is_named_as_undated():
    risk = _risk(assessment=_assessed(_ecoli(3.0), when=""))
    assert "E. coli in the water sample (date not recorded): 3 per 100 mL" in risk["statement"]


def test_the_statement_needs_a_date_for_today():
    with pytest.raises(ValueError, match="today's date"):
        _risk(today="last Tuesday")
    assert _risk(today=date(2024, 7, 1))["statement"] == _risk()["statement"]


# ---------------------------------------------------- the inspection as a record

def test_the_same_form_recorded_twice_merges_and_a_different_one_does_not():
    one = SAN.inspection_event("2024-06-30", "YNNNNNNNNN", by="A. Bangura")
    again = SAN.inspection_event("2024-06-30", "yNNNNNNNNN", by="A. Bangura")
    other = SAN.inspection_event("2024-06-30", "YYNNNNNNNN", by="A. Bangura")
    assert len(merge_events(ASSET, [one], [again])) == 1
    assert len(merge_events(ASSET, [one], [other])) == 2


def test_an_event_without_a_form_keeps_the_identifier_it_always_had():
    """The browser engine's own answers, as test_registry pins them."""
    assert event_id(ASSET.replace("8FEEVKQ-T", "8T4KQ2A-7"), "2024-01-01", "failure",
                    "pump seized", "M") == "2024-01-01/failure/8Z6E"
    plain = AssetEvent("2024-01-01", "inspection", "looked").identified_for(ASSET)
    assert "form" not in plain.as_dict() and "answers" not in plain.as_dict()


def test_an_inspection_survives_a_project_file():
    asset = _asset(SAN.inspection_event("2024-06-30", "YN--NNNNNY", by="A. Bangura"),
                   SAN.inspection_event("2024-07-01", "----------"))
    text = yaml.safe_dump({"asset": asset.as_dict()}, sort_keys=False)
    back = asset_from_dict(yaml.safe_load(text)["asset"])
    assert back is not None
    assert [(e.form, e.answers) for e in back.events] == [
        ("tubewell_handpump", "YN--NNNNNY"), ("tubewell_handpump", "----------")]
    assert back.as_dict() == asset.with_events().as_dict()


def test_the_asset_record_carries_the_statement_the_form_and_its_source(tmp_path):
    from docx import Document

    from groundwater.reporting.registry import AssetReportInputs, build_asset_record

    asset = _asset(AssetEvent("2020-01-10", "commissioned"),
                   SAN.inspection_event("2024-06-30", "YNNYNNNNNN", by="A. Bangura"))
    quality = _assessed(_ecoli(4.0))
    out = build_asset_record(
        AssetReportInputs(asset=asset, figures_dir=tmp_path, today=date(2024, 7, 1),
                          quality=quality),
        tmp_path / "record.docx")
    document = Document(str(out))
    text = "\n".join(p.text for p in document.paragraphs)
    statement = SAN.sanitary_risk(asset, "2024-07-01", quality)["statement"]
    assert statement in text and "Sanitary risk" in text
    assert phrase_table("references.citations")["who_surveillance_1997"] in text
    cells = [cell.text for t in document.tables for row in t.rows for cell in row.cells]
    assert SAN.form_questions()[9] in cells


# ------------------------------------------- the browser's engine, held to it

_DRIVER = r"""
import { readFileSync } from 'node:fs';
import vm from 'node:vm';
const sandbox = vm.createContext({ console, structuredClone, TextEncoder, TextDecoder });
for (const f of ['gwt-data.js', 'gwt-geo.js', 'gwt-core.js', 'support.js']) {
  vm.runInContext(readFileSync(process.argv[2] + '/' + f, 'utf8'), sandbox, { filename: f });
}
const C = sandbox.GWT.core;
const cases = JSON.parse(readFileSync(0, 'utf8'));
process.stdout.write(JSON.stringify(cases.map((r) => {
  try {
    // an assessment given whole is read as given, for a row no sample makes
    const assessment = r.assessment !== undefined ? r.assessment
      : (r.sample ? C.assessSample(r.sample) : null);
    const asset = { asset_id: r.asset_id, events: r.events };
    return { ok: {
      risk: C.sanitaryRisk(asset, r.today, assessment),
      merged: C.mergeEvents(r.asset_id, r.events),
      scores: r.answers.map((a) => C.scoreAnswers(a)),
      rows: r.answers.map((a) => C.formRows(a)),
      grades: r.counts.map((c) => C.ecoliGrade(c)),
      above: r.counts.map((c) => [C.ecoliGradesAbove(c, false), C.ecoliGradesAbove(c, true)]),
    } };
  } catch (e) {
    return { threw: String(e && e.stack || e) };
  }
})));
"""

_PARAMETERS = ["E. coli", "Thermotolerant coliforms", "Faecal coliforms", "E. coli"]
_BY = ["", "A. Bangura", "  M. Kolleh ", "Fatmata", "Ibrahim \U0001f600"]


def _case(seed: int) -> dict:
    """An asset's events, a sample and today, drawn from one seed."""
    rng = random.Random(seed)
    events = []
    # few days, so two forms often share one, and two of them after the
    # earlier of the two todays
    days = ["2024-01-20", "2024-02-10", "2024-03-01", "2024-03-01", "2024-05-20",
            "2024-06-30", "2024-06-30", "2024-08-15"]
    for _ in range(rng.randint(0, 6)):
        day = rng.choice(days)
        roll = rng.random()
        if roll < 0.7:
            answers = "".join(rng.choice(rng.choice(["NNNNNNNY", "YYNNN-", "YYYYYN-"]))
                              for _ in range(10))
            if rng.random() < 0.15:
                answers = rng.choice(["----------", "YN", "yNnn-xNNNN"])
            events.append({"when": day, "kind": "inspection", "by": rng.choice(_BY),
                           "note": rng.choice(["", "apron cracked"]),
                           "form": rng.choice(["tubewell_handpump"] * 9 + [""]),
                           "answers": answers})
        elif roll < 0.85:
            events.append({"when": day, "kind": "inspection", "note": "looked"})
        else:
            events.append({"when": rng.choice([day, "not written"]),
                           "kind": rng.choice(["repair", "failure", "water_sample"])})
    sample = None
    if rng.random() < 0.85:
        results = [dict(r) for r in _PANEL]
        for _ in range(rng.choice([0, 1, 1, 1, 1, 2])):
            roll = rng.random()
            row = {"parameter": rng.choice(_PARAMETERS), "unit": rng.choice(
                ["CFU/100 mL", "CFU/100 mL", "MPN/100 mL", "CFU/mL", "", "wibbles"])}
            if roll < 0.5:
                row["value"] = rng.choice([0.0, 0.5, 3.0, 10.0, 11.0, 99.0, 101.0, 1000.0, 1500.0])
            elif roll < 0.7:
                row.update(value=None, below_detection=True,
                           detection_limit=rng.choice([None, 1.0, 10.0, 10.0]))
            else:
                row.update(value=None, greater_than=rng.choice([0.0, 0.0, 10.0, 50.0, 1000.0]),
                           greater_than_inclusive=rng.random() < 0.5)
            results.insert(0, row)
        sample = {"site": {"community": "Ref"}, "flags": [], "results": results,
                  "sample_date": rng.choice(["2024-06-30", "20th May, 2018", ""])}
    # what the review found the engines reading apart, drawn from a second
    # generator so the cases above stay as they were: a compact date, which
    # neither engine reads, answers with ends the two engines' own trims
    # took differently, and a sample older than one the history records
    odd = random.Random(f"odd {seed}")
    for event in events:
        if odd.random() < 0.08:
            event["when"] = "20240520"
        if event.get("answers") and odd.random() < 0.15:
            event["answers"] = odd.choice(["\ufeff", "\x85 "]) + event["answers"] + odd.choice(
                ["\x1c", "\u3000", ""])
    if odd.random() < 0.4:
        events.append({"when": odd.choice(days), "kind": "water_sample",
                       "by": odd.choice(_BY)})
    if sample is not None and odd.random() < 0.3:
        sample["sample_date"] = "2024-02-01"
    return {"asset_id": ASSET, "events": events, "sample": sample,
            "today": rng.choice(["2024-07-01", "2024-03-15"]),
            "answers": ["YNNYNNYNYN", "YNN--NYN-N", "--", "yn" * 5, "----------"],
            "counts": [0, 0.5, 10, 10.5, 100, 1000, 1001]}


def _python(case: dict) -> dict:
    assessment = None
    if case["sample"] is not None:
        sample = WaterQualitySample(site=SiteMetadata(community="Ref"),
                                    sample_date=case["sample"]["sample_date"],
                                    results=[WaterQualityResult(**r) for r in case["sample"]["results"]])
        assessment = assess_sample(sample)
    asset = _asset(*case["events"])
    asset.events = merge_events(ASSET, case["events"])
    return json.loads(json.dumps({
        "risk": SAN.sanitary_risk(asset, case["today"], assessment),
        "merged": [e.as_dict() for e in merge_events(ASSET, case["events"])],
        "scores": [SAN.score_answers(a) for a in case["answers"]],
        "rows": [SAN.form_rows(a) for a in case["answers"]],
        "grades": [SAN.ecoli_grade(c) for c in case["counts"]],
        "above": [[list(SAN.ecoli_grades_above(c, False)), list(SAN.ecoli_grades_above(c, True))]
                  for c in case["counts"]],
    }))


SEEDS = range(40)


def _browser(cases: list, folder: Path) -> list:
    """What the browser's engine answers for each case, in a Node vm."""
    node = shutil.which("node")
    if node is None:
        pytest.skip("node is not installed; run `node --version` to check")
    driver = folder / "sanitary.mjs"
    driver.write_text(_DRIVER, encoding="utf-8")
    done = subprocess.run([node, str(driver), str(JS)], input=json.dumps(cases),
                          capture_output=True, text=True, encoding="utf-8", check=True)
    return json.loads(done.stdout)


@pytest.fixture(scope="module")
def browser_answers(tmp_path_factory):
    return _browser([_case(seed) for seed in SEEDS], tmp_path_factory.mktemp("sanitary"))


@pytest.mark.parametrize("seed", SEEDS)
def test_the_browser_reads_the_same_risk(browser_answers, seed):
    answer = browser_answers[seed]
    assert "threw" not in answer, answer.get("threw")
    assert answer["ok"] == _python(_case(seed))


def test_the_seeded_cases_reach_every_rule():
    told = [_python(_case(seed))["risk"] for seed in SEEDS]
    kinds = {t["microbiology"]["kind"] for t in told if t["microbiology"]}
    assert {"counted", "none_found", "bound", "detected", "ungraded", "missing"} <= kinds
    assert {t["action"] for t in told} == {None, "low", "higher", "urgent"}
    assert any(t["score"] and not t["score"]["complete"] for t in told)
    assert any(t["microbiology"] is None for t in told)
    statements = " ".join(t["statement"] for t in told)
    for head in ("A later inspection, of", "forms were recorded on",
                 "the action priority is between", "on the side that is known it is at least",
                 "none of the form's questions was answered",
                 "do not fit a form this toolkit has", "could not be graded: ",
                 "records a later water sample", "records a water sample of",
                 "records no water sample", "cannot be read as YYYY-MM-DD"):
        assert head in statements, head
    assert ".." not in statements


def test_no_statement_the_two_engines_are_held_to_ends_a_sentence_twice():
    """reference.json holds what parity.mjs holds the browser to; a double
    full stop there would be locked into both engines."""
    reference = json.loads((REPO / "tests" / "webapp" / "reference.json").read_text(encoding="utf-8"))
    statements = [c["expect"]["risk"]["statement"] for c in reference["sanitary"]["cases"]]
    assert statements and not [s for s in statements if ".." in s]
