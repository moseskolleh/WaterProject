"""Reading a field sheet that arrived as a PDF or a photograph.

What an extraction hands on has to be either right or marked: a filled
template that reads back a resistance as a resistivity, a misread digit
that passes as clean, or a review that says nothing needs checking when
every reading was lost is worse than no extraction at all. No test here
calls the network: the AI path is given its replies.
"""

from __future__ import annotations

import sys
import zlib
from types import ModuleType, SimpleNamespace

import pytest
from openpyxl import load_workbook

from groundwater.cli import main
from groundwater.extraction import fill_ves_template, write_review_workbook
from groundwater.extraction.claude import ClaudeExtractor, _reply_payload
from groundwater.extraction.pdf_text import _cell_confidence
from groundwater.ingestion import read_ves_workbook
from groundwater.text import phrase
from groundwater.utils import is_whole_number, parse_number

# A Schlumberger sheet as crews keep it: the half-spacing, the geometric
# factor, the measured resistance and the apparent resistivity worked out
# from them on the sheet.
_COLUMNS = ["No.", "AB/2 (m)", "MN/2 (m)", "K", "Resistance (ohm)",
            "App. Resistivity (ohm-m)"]
_ROWS = [["1", "1.5", "0.25", "13.74", "8.76", "120.4"],
         ["2", "2", "0.25", "24.74", "3.97", "98.1"],
         ["3", "3", "0.25", "56.16", "1.43", "80.2"],
         ["4", "4", "0.25", "100.14", "0.659", "66.0"],
         ["5", "6", "0.25", "225.8", "0.227", "51.3"]]


def _document(keep, kind="ves", tables=True):
    payload = {
        "document_kind": kind,
        "header": [{"name": "Community", "value": "Rokel", "confidence": 0.95},
                   {"name": "Sounding Number", "value": "VES 3", "confidence": 0.95}],
        "tables": [{"title": "Readings", "columns": [_COLUMNS[i] for i in keep],
                    "rows": [[row[i] for i in keep] for row in _ROWS]}] if tables else [],
        "uncertain_cells": [],
        "notes": "",
    }
    return ClaudeExtractor._to_document(payload, "scan.jpg")


def _read_back(document, tmp_path):
    return read_ves_workbook(fill_ves_template(document, tmp_path / "filled.xlsx"))[0]


# --------------------------------------------------------- the filled template


def test_the_template_takes_the_resistivity_and_doubles_a_half_spacing(tmp_path):
    """"Resistance (ohm)" matched the first needle, "resist", so every
    resistance went into the template's apparent-resistivity column, a factor
    K too small, and "MN/2" went into the full MN column; every cell parsed,
    so nothing was flagged."""
    document = _document([0, 1, 2, 3, 4, 5])
    sounding = _read_back(document, tmp_path)
    assert list(sounding.rho_app) == [120.4, 98.1, 80.2, 66.0, 51.3]
    assert list(sounding.mn) == [0.5] * 5
    assert list(sounding.ab2) == [1.5, 2.0, 3.0, 4.0, 6.0]
    assert document.review_items == []
    ws = load_workbook(tmp_path / "filled.xlsx").active
    assert [ws.cell(row=10, column=c).value for c in range(1, 5)] == [
        "No.", "AB/2 (m)", "MN/2 (m)", "Apparent Resistivity (ohm-m)"]


def test_a_sheet_of_resistances_is_read_back_as_resistivities(tmp_path):
    """With no worked-out column, the resistances go in as resistances, with
    the sheet's K beside them, and the reader multiplies them out as it does
    for any such sheet; without K it works K out from AB/2 and MN."""
    with_k = _read_back(_document([0, 1, 2, 3, 4]), tmp_path)
    assert list(with_k.rho_app) == pytest.approx(
        [float(r[3]) * float(r[4]) for r in _ROWS], rel=1e-12)
    assert "rho_computed_from_resistance" in [f.code for f in with_k.flags]

    without_k = _read_back(_document([0, 1, 2, 4]), tmp_path)
    # the sheet's own K was rounded; the factor worked out agrees to it
    assert list(without_k.rho_app) == pytest.approx([120.4, 98.1, 80.2, 66.0, 51.3],
                                                    rel=2e-3)
    assert list(without_k.mn) == [0.5] * 5


def test_readings_the_template_cannot_take_are_refused_and_said_so(tmp_path):
    """A table the VES reader would read back with no readings is not written
    as a template, and the review says why rather than that nothing needs
    checking."""
    reasons = {
        # a resistance needs K or MN to become a resistivity
        (0, 1, 4): "its resistances need a K or an MN column",
        # nothing the reader takes as a spacing or a reading
        (0, 3): "no column is headed as the AB/2 spacing; no column is headed "
                "as an apparent resistivity or a resistance",
    }
    for keep, reason in reasons.items():
        document = _document(list(keep))
        assert reason in document.review_items[0]
        with pytest.raises(ValueError, match="cannot be placed in the VES template"):
            fill_ves_template(document, tmp_path / "x.xlsx")
        review = write_review_workbook(document, tmp_path / "review.xlsx")
        texts = [r[0].value for r in load_workbook(review)["Review"].iter_rows()]
        assert not any("None: all values" in str(t) for t in texts)
    assert not (tmp_path / "x.xlsx").exists()

    wenner = ClaudeExtractor._to_document({
        "document_kind": "ves", "header": [], "uncertain_cells": [], "notes": "",
        "tables": [{"title": "T", "columns": ["a (m)", "Apparent Resistivity (ohm-m)"],
                    "rows": [["1", "120"]]}]}, "scan.jpg")
    assert "the Wenner spacing a" in wenner.review_items[0]
    # a table the transcription left untitled is named as the browser names it
    untitled = ClaudeExtractor._to_document({
        "document_kind": "ves", "header": [], "uncertain_cells": [], "notes": "",
        "tables": [{"title": "", "columns": ["a (m)", "Apparent Resistivity (ohm-m)"],
                    "rows": [["1", "120"]]}]}, "scan.jpg")
    assert untitled.review_items[0].startswith("table 'Table 1' cannot be placed")


def test_a_sheet_with_no_table_is_never_reviewed_as_clean(tmp_path):
    document = _document([], kind="pumping_test", tables=False)
    assert document.review_items == [phrase("extraction.no_table")]
    review = write_review_workbook(document, tmp_path / "review.xlsx")
    texts = [r[0].value for r in load_workbook(review)["Review"].iter_rows()]
    assert phrase("extraction.no_table") in texts
    assert not any("None: all values" in str(t) for t in texts)


def test_the_template_carries_nothing_the_sheet_did_not_give(tmp_path):
    """The blank template's "Syscal Junior" and sounding number 1 were read
    back as the sheet's instrument and number."""
    payload = {"document_kind": "ves", "header": [], "uncertain_cells": [],
               "notes": "", "tables": [{"title": "T", "columns": _COLUMNS, "rows": _ROWS}]}
    sounding = _read_back(ClaudeExtractor._to_document(payload, "scan.jpg"), tmp_path)
    assert sounding.instrument == ""
    assert sounding.sounding_id == "VES 1"


# ---------------------------------------------------------- the cell grader


def test_a_number_with_a_misread_letter_is_held_back():
    """parse_number takes the first number in a cell, so "12O.4" passed as
    clean and reached the template as 12."""
    column = "Apparent Resistivity (ohm-m)"
    for cell in ("12O.4", "9B.1", "8O.2", "12O", "1 234", "1/2", "<0.01", "45 ft"):
        assert parse_number(cell) is not None, cell
        assert _cell_confidence(cell, column) == (
            0.5, "contains characters other than a number"), cell
    # the ohm sign (U+2126) as well as the Greek capital omega a text layer
    # may carry for it
    for cell in ("120.4", "078,7", "2,933", "80m", "120.4 ohm-m", "120.4 Ωm",
                 "120.4 Ω·m", "5 min", "1e3"):
        assert _cell_confidence(cell, column) == (1.0, ""), cell
    assert _cell_confidence("", column)[0] == 0.6
    assert _cell_confidence("--", column)[0] == 0.4


def test_a_whole_number_is_the_number_and_a_unit_at_most():
    assert is_whole_number("1.234,5") and is_whole_number(" 0.8 l/s ")
    assert is_whole_number(12) and not is_whole_number(float("nan"))
    assert not is_whole_number(None) and not is_whole_number(True)
    assert not is_whole_number("12.") and not is_whole_number("m")


# ------------------------------------------------------------ the PDF reader


def _pdf(placed) -> bytes:
    """One page with each (x, y, text) placed in Helvetica; none draws a line
    and no text, as a scanner's page carries none."""
    ops = ["BT", "/F1 10 Tf"]
    for x, y, text in placed:
        ops.append(f"1 0 0 1 {x} {y} Tm ({text}) Tj")
    ops.append("ET")
    if not placed:
        ops = ["0.5 w", "60 600 m", "400 600 l", "S"]
    content = zlib.compress("\n".join(ops).encode("latin-1"), 9)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"),
        b"<< /Length 6 0 R /Filter /FlateDecode >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        str(len(content)).encode(),
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    return bytes(out)


def _captioned_sheet(tmp_path, cells=("120.4", "98.1", "80.2", "66.0", "51.3")):
    placed = [(60, 760, "SCHLUMBERGER ARRAY VES FIELD DATA"),
              (60, 735, "Community: Rokel"), (300, 735, "Sounding Number: VES A-1"),
              (60, 668, "Field readings \\(Schlumberger array\\)"),
              (60, 650, "No."), (110, 650, "AB/2 \\(m\\)"), (200, 650, "MN \\(m\\)"),
              (280, 650, "Apparent Resistivity \\(ohm-m\\)")]
    for i, (ab2, rho) in enumerate(zip(("1.5", "2", "3", "4", "6"), cells, strict=True)):
        y = 630 - i * 18
        placed += [(60, y, str(i + 1)), (110, y, ab2), (200, y, "0.5"), (280, y, rho)]
    placed.append((60, 630 - 5 * 18, "Notes: MN is the full potential electrode spacing."))
    path = tmp_path / "sheet.pdf"
    path.write_bytes(_pdf(placed))
    return path


def test_a_caption_and_a_note_stay_out_of_the_reading_table(tmp_path):
    """A line was taken into a table when it had two words, so a caption above
    the headings became the header, the headings the first reading, and the
    filled template came back with no readings while the review said nothing
    needed checking; a note under the table became a reading of its own."""
    pytest.importorskip("pdfplumber")
    from groundwater.extraction.pdf_text import extract_pdf_text

    document = extract_pdf_text(_captioned_sheet(tmp_path))
    [table] = document.tables
    assert table.columns == ["No.", "AB/2 (m)", "MN (m)", "Apparent Resistivity (ohm-m)"]
    assert [row[3] for row in table.rows] == ["120.4", "98.1", "80.2", "66.0", "51.3"]
    assert document.uncertain_cells == [] and document.review_items == []
    sounding = _read_back(document, tmp_path)
    assert list(sounding.rho_app) == [120.4, 98.1, 80.2, 66.0, 51.3]


def test_a_misread_cell_in_a_pdf_is_on_the_review(tmp_path):
    pytest.importorskip("pdfplumber")
    from groundwater.extraction.pdf_text import extract_pdf_text

    document = extract_pdf_text(
        _captioned_sheet(tmp_path, cells=("12O.4", "9B.1", "8O.2", "66.0", "51.3")))
    assert [(c.row, c.column, c.reason) for c in document.uncertain_cells] == [
        (r, 3, "contains characters other than a number") for r in range(3)]
    assert len(document.review_items) == 3 and "'12O.4'" in document.review_items[0]


def test_an_image_only_pdf_is_refused_as_the_browser_refuses_it(tmp_path):
    """Read as an empty document, a scanner's PDF came back as an "unknown"
    sheet whose review said every value was extracted with high confidence."""
    pytest.importorskip("pdfplumber")
    from groundwater.extraction.pdf_text import extract_pdf_text

    path = tmp_path / "scan.pdf"
    path.write_bytes(_pdf([]))
    with pytest.raises(ValueError) as refused:
        extract_pdf_text(path)
    assert str(refused.value) == phrase("extraction.no_text_layer")
    assert "AI assisted extraction" in str(refused.value)


def test_the_command_line_counts_what_needs_checking(tmp_path, capsys):
    """It looked for uncertain cells on the tables, which carry none, so a
    sheet with cells to check printed nothing."""
    pytest.importorskip("pdfplumber")
    sheet = _captioned_sheet(tmp_path, cells=("12O.4", "98.1", "80.2", "66.0", "51.3"))
    assert main(["extract", str(sheet), "--out", str(tmp_path / "out")]) == 0
    printed = capsys.readouterr().out
    assert "1 item(s) need checking" in printed
    assert "filled VES template" in printed


# ------------------------------------------------------- the AI transcription


def _reply(stop_reason, *blocks):
    return SimpleNamespace(stop_reason=stop_reason,
                           content=[SimpleNamespace(**b) for b in blocks])


def test_a_transcription_cut_off_at_the_output_limit_says_so():
    """The partial JSON went straight to the parser: "Expecting value: line 1
    column 3430"."""
    half = '{"document_kind": "pumping_test", "tables": [{"rows": [["1", "5.0'
    for reply in (_reply("max_tokens", {"type": "thinking", "thinking": "..."},
                         {"type": "text", "text": half}),
                  _reply("max_tokens", {"type": "thinking", "thinking": "..."})):
        with pytest.raises(RuntimeError) as cut:
            _reply_payload(reply)
        assert str(cut.value) == phrase("extraction.cut_off")
        assert "cut off" in str(cut.value)
    with pytest.raises(RuntimeError, match=r"not a complete transcription \(stop reason: end_turn\)"):
        _reply_payload(_reply("end_turn", {"type": "text", "text": half}))
    assert _reply_payload(_reply("end_turn", {"type": "text", "text": '{"a": 1}'})) == {"a": 1}


def test_the_extractor_reports_the_cut_off_without_the_network(tmp_path, monkeypatch):
    """The whole extract call, with the SDK stood in for: no request leaves."""
    fake = ModuleType("anthropic")

    class _Client:
        def __init__(self, *args, **kwargs):
            self.messages = SimpleNamespace(create=lambda **kw: _reply(
                "max_tokens", {"type": "text", "text": '{"document_kind": "ves", "hea'}))

    fake.Anthropic = _Client
    monkeypatch.setitem(sys.modules, "anthropic", fake)
    scan = tmp_path / "scan.png"
    scan.write_bytes(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(RuntimeError, match="The transcription was cut off"):
        ClaudeExtractor().extract(scan)
