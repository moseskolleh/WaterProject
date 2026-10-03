"""Photo evidence with provenance (PLAN.md step 2.4).

A supervision photograph carries its capture time, its position and its
SHA-256, each with where it came from; the checklist items that need one
hold the supervision record back until they have it; and the reports print
where each time and position came from. The committed fixture JPEG carries
EXIF written by hand (tests/webapp/fixtures/make_photo_fixture.py), so what
it ought to read is known without trusting another reader. That the browser
reads the same record from the same file is parity's job
(tests/webapp/parity.mjs, "photo provenance").
"""

from __future__ import annotations

import base64
import hashlib
import importlib.util
from pathlib import Path

import pytest

from groundwater.photos import describe_provenance, photo_provenance, read_exif
from groundwater.project_io import deserialize_project, serialize_project, stale_on_load
from groundwater.readiness import assess_readiness
from groundwater.supervision import load_checklists
from groundwater.text import phrase, phrase_table

FIXTURES = Path(__file__).resolve().parent / "webapp" / "fixtures"
FIXTURE = FIXTURES / "photo_exif.jpg"


def _maker():
    spec = importlib.util.spec_from_file_location(
        "make_photo_fixture", FIXTURES / "make_photo_fixture.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


MAKE = _maker()
ATTACHED = "2026-10-03T09:00:00Z"


def test_the_committed_fixture_is_the_one_the_script_writes():
    assert FIXTURE.read_bytes() == MAKE.fixture_bytes()


def test_the_fixture_reads_as_written():
    exif = read_exif(FIXTURE.read_bytes())
    assert exif["taken_at"] == "2024-03-05T14:22:10+00:00"
    gps = exif["gps"]
    assert gps["lat"] == pytest.approx(8 + 29 / 60 + 34.56 / 3600, abs=1e-12)
    assert gps["lon"] == pytest.approx(-(13 + 14 / 60 + 12.34 / 3600), abs=1e-12)
    assert gps["accuracy_m"] == 5.0


def test_a_big_endian_file_reads_the_same():
    tiff = MAKE.build_tiff(">", **MAKE.FIXTURE_EXIF)
    assert read_exif(MAKE.jpeg_with_exif(tiff)) == read_exif(FIXTURE.read_bytes())


def test_southern_and_eastern_signs():
    tiff = MAKE.build_tiff("<", taken="2024:03:05 14:22:10",
                           lat=("S", ((1, 1), (30, 1), (0, 1))),
                           lon=("E", ((2, 1), (15, 1), (0, 1))))
    gps = read_exif(MAKE.jpeg_with_exif(tiff))["gps"]
    assert (gps["lat"], gps["lon"], gps["accuracy_m"]) == (-1.5, 2.25, None)


@pytest.mark.parametrize("kwargs", [
    # a camera with no clock set
    {"taken": "0000:00:00 00:00:00"},
    {"taken": "    :  :     :  :  "},
    {"taken": "2024-03-05 14:22:10"},
    # a camera with no fix writes zeros
    {"lat": ("N", ((0, 1), (0, 1), (0, 1))), "lon": ("E", ((0, 1), (0, 1), (0, 1)))},
    # a zero denominator, and a missing reference
    {"lat": ("N", ((8, 0), (0, 1), (0, 1))), "lon": ("W", ((13, 1), (0, 1), (0, 1)))},
    {"lat": ("X", ((8, 1), (0, 1), (0, 1))), "lon": ("W", ((13, 1), (0, 1), (0, 1)))},
])
def test_a_blank_or_broken_value_is_not_read(kwargs):
    exif = read_exif(MAKE.jpeg_with_exif(MAKE.build_tiff("<", **kwargs)))
    assert exif == {"taken_at": None, "gps": None}


def test_a_value_padded_with_a_control_byte_is_not_read():
    """Only the spaces JavaScript's trim() takes are taken off a value.

    str.strip() also took 0x1C to 0x1F, so "W\\x1f" read as west here and as
    no reference in the browser, and the two engines recorded different
    positions for the same file. parity.mjs holds the browser to this case.
    """
    tiff = MAKE.build_tiff("<", taken="2024:03:05 14:22:10\x1d", offset="+05:00\x1f",
                           lat=("N", ((8, 1), (1, 1), (1, 1))),
                           lon=("W\x1f", ((13, 1), (0, 1), (0, 1))))
    assert read_exif(MAKE.jpeg_with_exif(tiff)) == {"taken_at": None, "gps": None}
    # the spaces trim() does take are still taken
    tiff = MAKE.build_tiff("<", taken=" 2024:03:05 14:22:10\t", offset="+05:00 ",
                           lat=("N ", ((8, 1), (1, 1), (1, 1))),
                           lon=("\x0bW", ((13, 1), (0, 1), (0, 1))))
    exif = read_exif(MAKE.jpeg_with_exif(tiff))
    assert exif["taken_at"] == "2024-03-05T14:22:10+05:00"
    assert exif["gps"]["lon"] == -13.0


@pytest.mark.parametrize("data", [
    b"", b"not a jpeg", MAKE.BASE_JPEG, FIXTURE.read_bytes()[:60],
    # an APP1 that claims more bytes than the file has
    MAKE.BASE_JPEG[:2] + b"\xff\xe1\xff\xff" + b"Exif\x00\x00II*\x00",
])
def test_no_exif_or_a_broken_file_reads_nothing_and_never_raises(data):
    assert read_exif(data) == {"taken_at": None, "gps": None}


def test_the_record_of_the_fixture():
    data = FIXTURE.read_bytes()
    record = photo_provenance(data, ATTACHED)
    assert record == {
        "format": 1,
        "sha256": hashlib.sha256(data).hexdigest(),
        "bytes": len(data),
        "attached_at": ATTACHED,
        "taken_at": "2024-03-05T14:22:10+00:00",
        "time_source": "exif",
        "position": read_exif(data)["gps"],
        "position_source": "exif",
        "position_note": "",
        "stored": "original",
    }
    # the file's own position wins over the device's: it is where the
    # photograph was taken, not where it was attached
    fixed = photo_provenance(data, ATTACHED, device_fix={"lat": 1.0, "lon": 2.0,
                                                         "accuracy_m": 9.0})
    assert fixed["position_source"] == "exif"


def test_without_exif_the_device_clock_and_fix_are_used_and_said_to_be():
    data = MAKE.BASE_JPEG
    record = photo_provenance(data, ATTACHED, device_fix={
        "lat": 8.48, "lon": -13.23, "accuracy_m": 12.0})
    assert record["taken_at"] == ATTACHED and record["time_source"] == "device"
    assert record["position"] == {"lat": 8.48, "lon": -13.23, "accuracy_m": 12.0}
    assert record["position_source"] == "device"
    bare = photo_provenance(data, ATTACHED, position_note="refused", stored="downscaled")
    assert bare["position"] is None and bare["position_source"] == "none"
    assert bare["position_note"] == "refused" and bare["stored"] == "downscaled"


def test_the_report_states_each_source():
    shown = describe_provenance(photo_provenance(FIXTURE.read_bytes(), ATTACHED))
    assert shown["time"] == ("2024-03-05 14:22:10+00:00 ("
                             + phrase_table("evidence.time_sources")["exif"] + ")")
    assert shown["position"].startswith("8.492933, -13.236761, within 5 m (")
    assert shown["hash"] == "SHA-256 " + hashlib.sha256(FIXTURE.read_bytes()).hexdigest()[:16]
    device = describe_provenance(photo_provenance(MAKE.BASE_JPEG, ATTACHED,
                                                  position_note="unsupported"))
    assert device["time"].startswith("2026-10-03 09:00:00 (this device's clock")
    assert device["position"] == phrase_table("evidence.position_notes")["unsupported"]


@pytest.mark.parametrize("record", [None, {}, {"taken_at": "2024-01-01T00:00:00"}])
def test_a_photograph_without_a_record_is_said_to_have_none(record):
    """Nothing is filled in: a photograph from before this step says so."""
    assert describe_provenance(record) == {
        "time": phrase("evidence.no_provenance"), "position": "", "hash": ""}


# ------------------------------------------------------------------ the gate

def _photo_items():
    return [i for i in load_checklists() if i.photo_required]


def test_the_csv_names_the_items_that_need_a_photograph():
    items = {i.item_id for i in _photo_items()}
    assert items == {"des-casing-screen-assemblage", "des-backfill-placed-6",
                     "dev-borehole-disinfected-chlorine"}
    assert all(i.critical for i in _photo_items())


def test_a_missing_photograph_holds_the_supervision_record_back():
    readiness = assess_readiness({}, "supervision")
    (req,) = [r for r in readiness.requirements if r.key == "photo_evidence"]
    assert req.state == "unmet" and not readiness.is_certifiable
    for item in _photo_items():
        assert item.text in req.detail
    assert req.detail.endswith(phrase("evidence.presence_only"))
    # the other reports make no claim about the construction photographs
    assert all(r.key != "photo_evidence"
               for r in assess_readiness({}, "completion").requirements)


def test_each_photograph_present_meets_it_and_still_says_presence_only():
    photo = {"b64": base64.b64encode(FIXTURE.read_bytes()).decode()}
    some = {_photo_items()[0].item_id: photo}
    partial = assess_readiness({"supervision": {"evidence": some}}, "supervision")
    (req,) = [r for r in partial.requirements if r.key == "photo_evidence"]
    assert req.state == "unmet" and _photo_items()[0].text not in req.detail

    full = {i.item_id: photo for i in _photo_items()}
    (req,) = [r for r in assess_readiness(
        {"supervision": {"evidence": full}}, "supervision").requirements
        if r.key == "photo_evidence"]
    assert req.state == "met"
    assert req.detail == (phrase("evidence.photos_present", n=3) + " "
                          + phrase("evidence.presence_only"))


def test_an_item_answered_na_needs_no_photograph():
    responses = {i.item_id: "na" for i in _photo_items()}
    (req,) = [r for r in assess_readiness(
        {"supervision": {"responses": responses}}, "supervision").requirements
        if r.key == "photo_evidence"]
    assert (req.state, req.detail) == ("met", phrase("evidence.none_required"))


# ------------------------------------------------------- project file, report

def _evidence():
    data = FIXTURE.read_bytes()
    return {"des-backfill-placed-6": {
        "name": "seal.jpg", "mime": "image/jpeg",
        "b64": base64.b64encode(data).decode("ascii"),
        "provenance": photo_provenance(data, ATTACHED, position_note="unsupported"),
    }}


def test_the_photographs_and_their_records_survive_the_project_file():
    evidence = _evidence()
    updates = deserialize_project(serialize_project(
        {"meta_community": "Rokel", "sup_evidence": evidence}, "0.4.0"))
    assert updates["evidence"] == evidence
    # and they are the outgoing borehole's: a load drops them
    assert "sup_evidence" in stale_on_load({"sup_evidence": evidence})


def test_an_older_project_file_loads_without_evidence_and_keeps_none_invented():
    plain = serialize_project({"meta_community": "Rokel"}, "0.3.0")
    assert b"evidence" not in plain
    assert "evidence" not in deserialize_project(plain)
    # a record-less photograph keeps no record rather than gaining one
    bare = {"x": {"name": "a.jpg", "b64": "AAAA"}}
    updates = deserialize_project(serialize_project({"sup_evidence": bare}, "0.4.0"))
    assert updates["evidence"]["x"]["provenance"] is None


def test_the_supervision_report_prints_the_provenance(tmp_path):
    from docx import Document

    from groundwater.models import SiteMetadata
    from groundwater.reporting.supervision import (
        SupervisionReportInputs,
        build_supervision_report,
    )
    from groundwater.supervision import evaluate_checklist

    items = load_checklists()
    evidence = _evidence()
    evidence["dev-borehole-disinfected-chlorine"] = {
        "name": "chlorine.jpg", "b64": base64.b64encode(MAKE.BASE_JPEG).decode()}
    path = build_supervision_report(SupervisionReportInputs(
        site=SiteMetadata(community="Rokel"), items=items, responses={},
        assessment=evaluate_checklist(items, {}), evidence=evidence,
        figures_dir=tmp_path), tmp_path / "sup.docx")
    doc = Document(str(path))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    shown = describe_provenance(evidence["des-backfill-placed-6"]["provenance"])
    assert shown["time"] in cells and shown["position"] in cells and shown["hash"] in cells
    assert phrase("evidence.no_provenance") in cells
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Photographic Evidence" in text and phrase("evidence.presence_only") in text


@pytest.mark.parametrize("notes, checks, numbers", [
    ([], False, ["3.1 Photographic Evidence"]),
    (["Grout to surface."], False,
     ["3.1 Site Notes and Instructions", "3.2 Photographic Evidence"]),
    ([], True, ["3.1 Field Acceptance Checks", "3.2 Photographic Evidence"]),
    (["Grout to surface."], True, ["3.1 Site Notes and Instructions",
                                   "3.2 Field Acceptance Checks",
                                   "3.3 Photographic Evidence"]),
])
def test_the_site_record_is_numbered_as_printed(tmp_path, notes, checks, numbers):
    """A record with photographs and nothing else printed "3.3" straight under
    "3.", with no 3.1 or 3.2 before it."""
    from docx import Document

    from groundwater.models import SiteMetadata
    from groundwater.reporting.supervision import (
        SupervisionReportInputs,
        build_supervision_report,
    )
    from groundwater.supervision import evaluate_checklist
    from groundwater.supervision.field_checks import handpump_corrosion_check

    items = load_checklists()
    path = build_supervision_report(SupervisionReportInputs(
        site=SiteMetadata(community="Rokel"), items=items, responses={},
        assessment=evaluate_checklist(items, {}), evidence=_evidence(),
        notes=notes, field_checks=[handpump_corrosion_check(6.0)] if checks else [],
        figures_dir=tmp_path), tmp_path / "sup.docx")
    headings = [p.text for p in Document(str(path)).paragraphs
                if p.style.name.startswith("Heading") and p.text.startswith("3.")]
    assert headings == ["3. Site Record"] + numbers
