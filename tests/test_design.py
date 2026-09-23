"""The borehole design reads the driller's own words.

Dr Timbo's log says "fracture zone 49-52 m" on the 45-50 m row and
"fracture zone 60-62 m" on the 55-60 m row, records grouting to 20 m, and
notes a seepage at 12 m in clayey laterite. The design used to screen the
five-metre rows (45-50 m covered one metre of the zone, 60-62 m sat behind
plain casing), screen the laterite at 14.5-17 m, draw a 6 m seal with a
gravel pack inside the grouted interval, and call itself as-built.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from groundwater.config import DesignRules
from groundwater.costing import inputs_from_design
from groundwater.design import (
    AS_BUILT_NOTE,
    DESIGN_NOTE,
    design_borehole,
    fracture_ranges,
    is_clayey,
    lithology_bands,
    lithology_class,
    logged_diameter_in,
    seal_depth_for,
)
from groundwater.ingestion import read_drilling_workbook
from groundwater.ingestion.drilling import parse_installed_screens
from groundwater.ingestion.templates import write_drilling_template
from groundwater.models import DrillingLog, LithologyInterval, SiteMetadata


def _log(intervals, strikes=(), depth=None, grout=None, installed=()):
    return DrillingLog(
        site=SiteMetadata(community="X"),
        total_depth_m=depth or max(b for _, b, _ in intervals),
        intervals=[LithologyInterval(t, b, d) for t, b, d in intervals],
        water_strikes_m=list(strikes),
        grouting_depth_m=grout,
        installed_screens_m=list(installed),
    )


def test_the_fracture_zones_the_log_names_are_screened(sample_data):
    log = read_drilling_workbook(sample_data / "dr_timbo" / "dr_timbo_drilling_log.xlsx")
    design = design_borehole(log=log, static_water_level_m=9.44)
    assert [(s.top_m, s.bottom_m) for s in design.screens] == [
        (25.0, 35.0), (48.0, 53.0), (59.0, 63.0),
    ]
    basis = " | ".join(design.design_basis)
    assert "fracture zones the log names (49-52 m, 60-62 m), with 1 m of screen" in basis
    assert "water strikes recorded in the drilling log (30 m)" in basis
    assert "the 12 m strike is not screened" in basis
    assert "clayey laterites" in basis and "20 m grouted interval" in basis
    # the recorded grout is the seal, and nothing is drawn inside it
    assert design.sanitary_seal == (0.0, 20.0)
    assert design.backfill == (20.0, 23.0)
    assert design.gravel_pack == (23.0, 70.0)
    assert "cement grout from surface to 20 m as recorded on the drilling log" in basis
    # the log's 6.5 inch hole round 5 inch casing takes no pack
    assert design.borehole_diameter_in == 6.5 and "as logged" in basis
    assert design.annular_fill == "none" and round(design.annulus_mm) == 19
    assert inputs_from_design(design).gravel_interval_m == 0.0
    assert any(f.code == "thin_annulus" and "no gravel pack is drawn or priced" in f.message
               for f in design.flags)
    assert not design.as_built and design.construction_note == DESIGN_NOTE


def test_a_seepage_in_clay_is_cased_off_not_screened():
    log = _log([(0, 6, "lateritic topsoil"), (6, 20, "clayey laterite, wet from 8 m"),
                (20, 40, "fractured granite")], strikes=[8.0, 25.0])
    design = design_borehole(log=log, static_water_level_m=2.0)
    assert all(s.top_m >= 20.0 for s in design.screens)
    assert any("the 8 m strike is not screened" in b and "seepage horizon" in b
               for b in design.design_basis)
    assert any("water strikes recorded in the drilling log (25 m)" in b
               for b in design.design_basis)
    assert is_clayey("Light yellow clayey laterites") and not is_clayey("fractured granite")


def test_the_recorded_grout_sets_the_seal():
    log = _log([(0, 10, "laterite"), (10, 40, "fractured granite")], strikes=[12.0],
               grout=15.0)
    design = design_borehole(log=log, static_water_level_m=4.0)
    assert seal_depth_for(log, DesignRules()) == 15.0
    assert design.sanitary_seal == (0.0, 15.0)
    assert design.gravel_pack[0] >= 15.0
    # the 12 m strike is inside the grout, so it is not a screen target
    assert all(s.top_m >= 15.0 for s in design.screens)
    assert any("within the 15 m grouted interval" in b for b in design.design_basis)
    # without a recorded grout the rule's minimum stands
    plain = design_borehole(log=_log([(0, 10, "laterite"), (10, 40, "fractured granite")]),
                            static_water_level_m=4.0)
    assert plain.sanitary_seal == (0.0, DesignRules().sanitary_seal_depth_m)


def test_the_annulus_decides_the_fill():
    log = _log([(0, 10, "laterite"), (10, 40, "fractured granite")], strikes=[20.0])
    none = design_borehole(log=log, static_water_level_m=4.0)
    assert none.annular_fill == "none"
    assert "no gravel pack" in dict(none.summary_rows())["Annular fill"]
    assert inputs_from_design(none).gravel_interval_m == 0.0
    stabiliser = design_borehole(log=log, static_water_level_m=4.0,
                                 rules=DesignRules(borehole_diameter_in=8.0, casing_diameter_in=4.0))
    assert stabiliser.annular_fill == "formation stabiliser"
    assert inputs_from_design(stabiliser).gravel_interval_m > 0
    pack = design_borehole(log=log, static_water_level_m=4.0,
                           rules=DesignRules(borehole_diameter_in=10.0, casing_diameter_in=4.0))
    assert pack.annular_fill == "gravel pack"
    assert "gravel pack (well sorted" in dict(pack.summary_rows())["Annular fill"]


def test_the_drilled_diameter_comes_from_the_log():
    log = _log([(0, 10, "laterite"), (10, 40, "fractured granite")], strikes=[20.0])
    for interval in log.intervals:
        interval.bit_diameter_in = 9.0
    assert logged_diameter_in(log) == 9.0
    design = design_borehole(log=log, static_water_level_m=4.0)
    assert design.borehole_diameter_in == 9.0
    assert design.annular_fill == "formation stabiliser"   # 51 mm a side round 5 inch
    assert any("9 inch hole as logged" in b for b in design.design_basis)


def test_installed_screens_make_the_drawing_as_built(tmp_path):
    assert parse_installed_screens("25-35; 48-53 m") == [(25.0, 35.0), (48.0, 53.0)]
    assert parse_installed_screens("") == []
    log = _log([(0, 10, "laterite"), (10, 40, "fractured granite")], strikes=[20.0],
               installed=[(22.0, 30.0)])
    design = design_borehole(log=log, static_water_level_m=4.0)
    assert design.as_built and design.construction_note == AS_BUILT_NOTE
    assert [(s.top_m, s.bottom_m) for s in design.screens] == [(22.0, 30.0)]
    assert design.design_basis[0].startswith("screens as installed, recorded on the drilling log")
    assert dict(design.summary_rows())["Screens (as installed)"].startswith("22-30 m")
    # and the template carries the field the crew fills in
    path = write_drilling_template(tmp_path / "log.xlsx", n_rows=3)
    wb = load_workbook(path)
    ws = wb.active
    ws["B2"] = "Somewhere"
    ws["E5"] = 40
    ws["B10"] = "22-30; 33-36"
    ws["A12"] = "0-10"
    ws["E12"] = "laterite"
    ws["A13"] = "10-40"
    ws["E13"] = "fractured granite"
    wb.save(path)
    parsed = read_drilling_workbook(path)
    assert parsed.installed_screens_m == [(22.0, 30.0), (33.0, 36.0)]
    assert design_borehole(log=parsed, static_water_level_m=4.0).as_built


def test_lithology_bands_put_a_named_zone_where_it_is(sample_data):
    log = read_drilling_workbook(sample_data / "dr_timbo" / "dr_timbo_drilling_log.xlsx")
    bands = {(t, b): c.label for t, b, c in lithology_bands(log.intervals)}
    assert bands[(45.0, 49.0)] == "Basement rock"
    assert bands[(49.0, 50.0)] == "Fracture zone"
    assert bands[(50.0, 52.0)] == "Fracture zone"
    assert bands[(52.0, 55.0)] == "Basement rock"
    # named on the 55-60 m row, drawn at 60-62 m
    assert bands[(55.0, 60.0)] == "Basement rock"
    assert bands[(60.0, 62.0)] == "Fracture zone"
    assert bands[(62.0, 65.0)] == "Basement rock"
    assert bands[(35.0, 40.0)] == "Weathered rock"      # slightly weathered, not fresh
    assert bands[(15.0, 20.0)] == "Saprolite"           # clayey saprolite is saprolite
    assert fracture_ranges("Light colour granite, fracture zone 49-52 m") == [(49.0, 52.0)]
    assert lithology_class("Weathered light colour granite, fractured").label == "Fracture zone"


def test_depth_tables_print_half_metres():
    """The table used to round 14.5 m to "14" beside a drawing that said 14.5."""
    log = _log([(0, 10, "laterite"), (10, 25, "granite"),
                (25, 40, "fractured granite, water bearing")], strikes=[13.5])
    design = design_borehole(log=log, static_water_level_m=7.0)
    rows = dict(design.summary_rows())
    assert [(s.top_m, s.bottom_m) for s in design.screens] == [(12.5, 18.5), (25.0, 38.0)]
    assert rows["Screens"].startswith("12.5-18.5 m; 25-38 m")


def test_every_zone_a_fracture_phrase_names_is_read():
    """Only the first range right after the phrase was read.

    "Granite, fractures at 30-31 m and 33-34 m" gave one screen and left
    33-34 m behind plain casing; "between 49 and 52 m", an em dash, "49 m to
    52 m", "metres" and a missing unit all read as nothing, and the design
    screened the whole logged row, the defect the named zones were for.
    """
    from groundwater.design import read_fractures
    from groundwater.design.lithology import host_description

    assert fracture_ranges("Granite, fractures at 30-31 m and 33-34 m") == [
        (30.0, 31.0), (33.0, 34.0)]
    assert fracture_ranges("fractured zones 49-52 m, 55-56 m") == [(49.0, 52.0), (55.0, 56.0)]
    for text in ("fracture zone between 49 and 52 m", "fracture zone 49—52 m",
                 "fracture zone 49-52", "fracture zone 49 m to 52 m",
                 "fracture zone 49-52 metres", "fracture zone from 49m-52m"):
        assert fracture_ranges(text, 45, 50) == [(49.0, 52.0)], text
    # an aperture or a count is not a depth, and a "zone" far from its row
    # is not taken as one
    assert fracture_ranges("fracture zone 49-52 mm") == []
    assert fracture_ranges("fractured, 1-2 per metre") == []
    assert fracture_ranges("fracture zone 12-14 m", 45, 50) == []
    # depths the phrase names but that cannot be read leave the row a target
    assert read_fractures("fractures at 49 and 52 m", 45, 50).unread
    assert not read_fractures("Weathered granite, fractured", 25, 30).unread
    assert host_description("Granite, fractures at 30-31 m and 33-34 m", 30, 35) == "Granite"

    log = _log([(0, 10, "laterite"), (10, 30, "granite"),
                (30, 35, "Granite, fractures at 30-31 m and 33-34 m"),
                (35, 45, "Granite, fracture zone between 40 and 42 m"),
                (45, 50, "Granite, fracture zone 46—48 metres"),
                (50, 55, "Granite, fractures at 51 and 53 m"), (55, 60, "granite")])
    design = design_borehole(log=log, static_water_level_m=5.0)
    assert [(s.top_m, s.bottom_m) for s in design.screens] == [
        (29.0, 35.0), (39.0, 43.0), (45.0, 55.0)]
    basis = " | ".join(design.design_basis)
    assert "(30-31 m, 33-34 m, 40-42 m, 46-48 m), with 1 m of screen" in basis
    assert "intervals logged at 50-55 m" in basis


def test_a_pump_intake_in_a_bottom_screen_is_not_lifted_above_the_test():
    """A 40-68 m screen lifted Dr Timbo's 52 m intake to 39 m, above the
    42.3 m the test drew the water to: the intake hydraulics-4 took out of the
    yield, put back by the design with only an info flag that reached no
    document."""
    from groundwater.design import pump_intake_floor
    from groundwater.hydraulics import analyse_pumping_test
    from groundwater.ingestion import read_pumping_workbook
    from groundwater.reporting.completion import _design_notes

    log = _log([(0, 10, "laterite"), (10, 45, "granite"),
                (45, 50, "granite, fracture zone 49-52 m"), (50, 70, "granite")], grout=20.0)
    held = design_borehole(log=log, static_water_level_m=9.44, pump_intake_m=52.0,
                           pump_intake_floor_m=45.26, screens_m=[(40.0, 68.0)])
    assert held.pump_intake_m == 52.0
    flag = next(f for f in held.flags if f.code == "pump_intake_in_screen")
    assert flag.level == "warning" and "39 m" in flag.message and "45.26 m" in flag.message
    assert not any("pump intake at 39 m" in b for b in held.design_basis)

    class Notes:
        def __init__(self):
            self.lines = []

        def paragraph(self, text, **kw):
            self.lines.append(text)

        def bullets(self, items):
            self.lines.extend(items)

    notes = Notes()
    _design_notes(notes, held)
    assert "Design notes:" in notes.lines and any("pump_intake_in_screen" in n for n in notes.lines)

    # plain casing above the screen and below the floor is used
    moved = design_borehole(log=log, static_water_level_m=9.44, pump_intake_m=52.0,
                            pump_intake_floor_m=45.26, screens_m=[(47.0, 68.0)])
    assert moved.pump_intake_m == 46.0
    # and with no floor to check it against, the intake is not lifted at all
    unknown = design_borehole(log=log, static_water_level_m=9.44, pump_intake_m=52.0,
                              screens_m=[(40.0, 68.0)])
    assert unknown.pump_intake_m == 52.0
    assert any(f.code == "pump_intake_in_screen" for f in unknown.flags)
    # moving down is still preferred, and needs no floor
    below = design_borehole(log=log, static_water_level_m=9.44, pump_intake_m=52.0,
                            screens_m=[(48.0, 53.0)])
    assert below.pump_intake_m == 54.0

    # the floor is the one the yield recommendation holds to
    analysis = analyse_pumping_test(read_pumping_workbook(
        Path(__file__).resolve().parents[1] / "examples" / "data" / "dr_timbo"
        / "dr_timbo_constant_test.xlsx"))
    yr = analysis.yield_recommendation
    floor = pump_intake_floor(yr, 3.0)
    assert floor == yr.deepest_pumping_level_m + 3.0
    assert yr.pump_installation_depth_m >= floor


def test_a_grout_past_the_sump_is_held_there_and_flagged():
    """A grout of 70 m in a 60 m hole printed a "0-70 m" seal, "Annular fill
    70-60 m", "Backfill 70-70 m" and blamed the static water level; a grout of
    55 m gave "no aquifer intervals identified" beside the 49-52 m zone the log
    names."""
    rows = [(0, 10, "laterite"), (10, 45, "granite"),
            (45, 50, "Granite, fracture zone 49-52 m"), (50, 60, "granite")]
    deep = design_borehole(log=_log(rows, grout=70.0), static_water_level_m=5.0)
    assert seal_depth_for(_log(rows, grout=70.0), DesignRules()) == 58.0
    assert deep.sanitary_seal == (0.0, 58.0)
    assert deep.gravel_pack[0] <= deep.gravel_pack[1] == 60.0
    summary = dict(deep.summary_rows())
    assert summary["Backfill"] == "none"
    too_deep = next(f for f in deep.flags if f.code == "grout_too_deep")
    assert too_deep.level == "error" and "70 m" in too_deep.message
    shallow = next(f for f in deep.flags if f.code == "hole_too_shallow")
    assert "grouted interval" in shallow.message
    assert "Static water level" not in shallow.message
    assert any("where the drilling log records 70 m" in b for b in deep.design_basis)

    grouted = design_borehole(log=_log(rows, grout=55.0), static_water_level_m=5.0)
    basis = grouted.design_basis
    assert ("the 49-52 m fracture zone is not screened: it lies within the 55 m "
            "grouted interval") in basis
    assert not any(b.startswith("no aquifer intervals identified") for b in basis)
    assert any(b.startswith("none of the aquifer intervals identified from the data "
                            "can be screened") for b in basis)


def test_the_basis_describes_the_screens_that_are_built():
    """Zones clipped by the grout, the sump or the hole bottom, or trimmed
    away with the shallowest screens, stayed in the basis as if screened."""
    grouted = design_borehole(
        log=_log([(0, 10, "laterite"), (10, 20, "Granite, fracture zone 16-18 m"),
                  (20, 50, "granite"), (50, 55, "Granite, fracture zone 49-52 m"),
                  (55, 60, "Granite, fracture zone 59-62 m")], grout=20.0),
        static_water_level_m=5.0)
    assert [(s.top_m, s.bottom_m) for s in grouted.screens] == [(48.0, 53.0)]
    basis = grouted.design_basis
    assert ("screens positioned against the fracture zones the log names (49-52 m), "
            "with 1 m of screen either side") in basis
    assert ("the 16-18 m fracture zone is not screened: it lies within the 20 m "
            "grouted interval") in basis
    assert ("the 59-62 m fracture zone is not screened: it lies below the top of the "
            "2 m sump, at 58 m") in basis

    sump = design_borehole(
        log=_log([(0, 10, "laterite"), (10, 55, "granite"),
                  (55, 60, "Granite, fracture zone 58-62 m")]),
        static_water_level_m=5.0)
    assert [(s.top_m, s.bottom_m) for s in sump.screens] == [(57.0, 58.0)]
    assert not any("either side)" in b or "(58-62 m)" in b for b in sump.design_basis)
    assert any(b.startswith("the 58-62 m fracture zone is screened at 57-58 m, not with "
                            "1 m of screen either side") for b in sump.design_basis)

    trimmed = design_borehole(
        log=_log([(0, 15, "granite"), (15, 40, "granite, fractured")], strikes=[9.0]),
        static_water_level_m=2.0)
    assert [(s.top_m, s.bottom_m) for s in trimmed.screens] == [(14.0, 38.0)]
    assert not any("water strikes recorded" in b for b in trimmed.design_basis)
    assert ("the 9 m strike is not screened: the screens were trimmed to 60 percent "
            "of the hole, keeping the deepest sections") in trimmed.design_basis


def test_an_as_built_record_is_kept_as_recorded():
    """Installed screens 15-25; 48-54; 54-60 m in a 60 m hole with a 20 m
    grout printed "Screens (as installed) 15-25 m; 48-58 m" with only an info
    flag, listed the clipped values as the record, said nothing of 15-20 m of
    screen inside the grout, and printed "Backfill 20-20 m"."""
    log = _log([(0, 10, "laterite"), (10, 60, "granite")], grout=20.0,
               installed=[(15.0, 25.0), (48.0, 54.0), (54.0, 60.0)])
    design = design_borehole(log=log, static_water_level_m=5.0)
    assert design.as_built
    assert [(s.top_m, s.bottom_m) for s in design.screens] == [
        (15.0, 25.0), (48.0, 54.0), (54.0, 58.0)]
    assert design.design_basis[0] == ("screens as installed, recorded on the drilling "
                                      "log (15-25 m, 48-54 m, 54-60 m)")
    clipped = next(f for f in design.flags if f.code == "screen_clipped")
    assert clipped.level == "warning" and "54-60 m" in clipped.message
    grouted = next(f for f in design.flags if f.code == "screen_in_grout")
    assert grouted.level == "warning" and "15-25 m" in grouted.message
    assert dict(design.summary_rows())["Backfill"] == "none"
    # screens the analyst places still merge where they touch, and a clip of
    # theirs stays an info flag
    placed = design_borehole(log=_log([(0, 10, "laterite"), (10, 60, "granite")]),
                             static_water_level_m=5.0,
                             screens_m=[(48.0, 54.0), (54.0, 60.0)])
    assert [(s.top_m, s.bottom_m) for s in placed.screens] == [(48.0, 58.0)]
    assert [f.level for f in placed.flags if f.code == "screen_clipped"] == ["info"]


def test_an_as_built_drawing_is_titled_and_labelled_as_one(tmp_path, monkeypatch):
    """The completion report titled an as-built drawing "Borehole design"
    under a header saying "Drawing: as built", and the drawing dropped
    "(recommended)" from a pump nobody recorded installing."""
    import matplotlib

    matplotlib.use("Agg")
    from groundwater.design import draw_borehole_design
    from groundwater.reporting import completion
    from groundwater.reporting.completion import CompletionReportInputs, build_completion_report

    log = _log([(0, 10, "laterite"), (10, 40, "fractured granite")], strikes=[20.0],
               installed=[(22.0, 30.0)])
    design = design_borehole(log=log, static_water_level_m=4.0, pump_intake_m=33.0)
    figure = draw_borehole_design(design, log)
    texts = [t.get_text() for ax in figure.axes for t in ax.texts]
    assert "pump intake 33 m (recommended)" in texts

    titles = []

    def capture(design, log, path=None, title=None, **kw):
        titles.append(title)
        return draw_borehole_design(design, log, path=path, title=title, **kw)

    monkeypatch.setattr(completion, "draw_borehole_design", capture)
    build_completion_report(CompletionReportInputs(log=log, design=design,
                                                   figures_dir=tmp_path),
                            tmp_path / "completion.docx")
    assert titles and titles[0].startswith("As-built borehole record - ")
