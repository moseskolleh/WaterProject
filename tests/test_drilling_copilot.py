"""The drilling log co-pilot's sheets, read by the Python reader (PLAN.md step 2.3).

The co-pilot is a page of the browser app (docs/js/gwt-drill-copilot.js), and
rule 6 of the plan is that a field tool writes the template the readers
already take. So the page's own code is run here, in Node, beside the
support and engine scripts it uses. The Dr Timbo log is drilled again
through it, over two days, from the clock times on its sheet: each interval
picked from the lithology classes with the driller's words as the note, the
two water strikes with an airlift reading each, and each day countersigned.
It writes its workbooks with the app's own .xlsx writer, and
``read_drilling_workbook`` reads the drilling log as it reads any other.
The driller's daily report has no reader in either app, so it is held to
the layout ``write_daily_log_template`` gives the template instead.
tests/webapp/copilot.mjs drives the same page in Chromium and reads the log
back with the browser's reader.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from groundwater.design.lithology import lithology_class
from groundwater.field_kit import airlift_yield
from groundwater.geo import geographic_to_utm
from groundwater.ingestion.drilling import read_drilling_workbook
from groundwater.ingestion.templates import write_daily_log_template

REPO = Path(__file__).resolve().parents[1]
JS = REPO / "docs" / "js"

# examples/data/dr_timbo: the interval, the class the driller picks, the
# words of the sheet as the note, and the clock times of the sheet. The
# sheet's own descriptions do not all read as what they describe ("Light
# colour granite, slightly weathered" reads as weathered rock, "Weathered
# light colour granite, fractured" as a fracture zone); here the class is
# picked and the note may only add to it.
TIMBO = [
    (0, 5, "topsoil", "reddish brown, lateritic", "13:30", "13:35"),
    (5, 10, "laterite", "light yellow, clayey", "13:36", "13:41"),
    (10, 15, "laterite", "light yellow, clayey, wet from 12 m", "13:55", "14:10"),
    (15, 20, "saprolite", "clayey, with fragments of weathered granite", "14:14", "14:26"),
    (20, 25, "saprolite", "decreasing clay content", "14:30", "14:38"),
    (25, 30, "fracture", "in weathered light colour granite", "15:15", "15:23"),
    (30, 35, "fracture", "light colour granite, slightly weathered", "08:25", "08:34"),
    (35, 40, "weathered", "light colour granite", "08:38", "08:48"),
    (40, 45, "basement", "light colour granite", "08:50", "09:01"),
    (45, 50, "basement", "light colour granite", "09:03", "09:13"),
    (50, 55, "basement", "light colour granite", "09:15", "09:17"),
    (55, 60, "basement", "light colour granite", "09:18", "09:20"),
    (60, 65, "basement", "light colour granite", "09:27", "09:31"),
    (65, 70, "basement", "light colour granite, fresh", "09:56", "10:00"),
]
DAY_ONE, DAY_TWO = "2026-10-03", "2026-10-04"

STRIKES = [  # interval index, the strike, its airlift reading
    (2, {"depth_m": 12, "method": "bucket", "volume_l": 20, "timings_s": [41, 39.5, 40.2]}),
    (5, {"depth_m": 30, "method": "vnotch", "head_mm": 62}),
]

_PLAYBACK = r"""
import { readFileSync, writeFileSync } from 'node:fs';
import vm from 'node:vm';
const [dir, input, output] = process.argv.slice(2);
const sandbox = { console, TextEncoder, TextDecoder, Blob, Response,
  CompressionStream, DecompressionStream };
sandbox.window = sandbox;
vm.createContext(sandbox);
for (const f of ['support.js', 'gwt-data.js', 'gwt-core.js', 'gwt-drill-copilot.js']) {
  vm.runInContext(readFileSync(dir + '/' + f, 'utf8'), sandbox, { filename: f });
}
const GWT = sandbox.GWT, D = GWT.drillCopilot, S = GWT.support;
const job = JSON.parse(readFileSync(input, 'utf8'));
// the app's store, as far as the co-pilot uses it
const state = { site: job.site };
GWT.app = { store: { get: (k) => state[k], set: (k, v) => { state[k] = v; }, flush() {} } };
let t = 0;
GWT.drillCopilotClock = () => t;
const at = (day, hm) => Date.parse(day + 'T' + hm + ':00Z');
state.drillCopilot = D.blankSession(job.site);
Object.entries(job.setup).forEach(([k, v]) => D.setSetup(k, v));
const out = { errors: {}, rates: [], statuses: {} };
t = at(job.days[0], '13:00');
D.setGps({ lat: 8.4657, lon: -13.2317, accuracy_m: 6, at: t });
// a class typed in as words, and a note that makes the row another class
try { D.startInterval(); D.endInterval({ bottom_m: 5, lithology: 'Reddish brown clay' }); }
catch (e) { out.errors.freeText = e.message; }
try { D.endInterval({ bottom_m: 5, lithology: 'clay', note: 'soft saprolite' }); }
catch (e) { out.errors.slip = e.message; }
try { D.endInterval({ bottom_m: 5, lithology: 'topsoil', note: 'water strike at 4 m' }); }
catch (e) { out.errors.strikeInNote = e.message; }
D.cancelCurrent();
job.intervals.forEach(([top, bottom, key, note, from, to], i) => {
  const day = job.days[i < 6 ? 0 : 1];
  t = at(day, from);
  D.startInterval();
  t = at(day, to);
  out.rates.push(D.endInterval({ bottom_m: bottom, lithology: key, note }, { stop: true }).minPerM);
  if (i === 5) {
    job.strikes.forEach(([index, reading]) => D.setStrike(index, reading));
    t = at(day, '17:00');
    D.signDay(day, 'M. Kamara');
  }
});
try { D.setStrike(9, { depth_m: 70, method: 'none', reason: 'compressor down' }); }
catch (e) { out.errors.outside = e.message; }
t = at(job.days[1], '17:00');
D.signDay(job.days[1], 'M. Kamara');
out.statuses.signed = job.days.map((d) => D.dayStatus(D.session(), d));
// a correction to a signed day needs a reason, and leaves it amended
try { D.correctInterval(13, { note: 'light colour granite, fresh, hard' }); }
catch (e) { out.errors.unreasoned = e.message; }
t = at(job.days[1], '17:30');
D.correctInterval(13, { note: 'light colour granite, fresh, hard' }, 'hardness left off');
out.statuses.amended = job.days.map((d) => D.dayStatus(D.session(), d));
const session = D.session();
writeFileSync(output + '/drilling.xlsx',
  await S.writeXlsx(D.drillingSheets(session, job.site, at(job.days[1], '18:00'))));
writeFileSync(output + '/daily.xlsx', await S.writeXlsx(D.dailySheets(session)));
out.session = session;
out.digests = job.days.map((d) => D.dayDigest(session, d));
writeFileSync(output + '/result.json', JSON.stringify(out));
"""


def _playback(tmp_path: Path) -> dict:
    node = shutil.which("node")
    if node is None:
        pytest.skip("node is not installed; run `node --version` to check")
    job = {
        "site": {"client": "Dr. Timbo", "community": "Dr. Timbo's Residence",
                 "district": "Western Area Rural",
                 "contractor": "WiNGiN Heavy Duty Machines Co. Ltd"},
        "setup": {"boreholeRef": "BH-1", "method": "DTH hammer", "rig": "PAT 401",
                  "driller": "A. Sesay", "bitIn": 6.5},
        "days": [DAY_ONE, DAY_TWO],
        "intervals": TIMBO, "strikes": STRIKES,
    }
    script = tmp_path / "playback.mjs"
    script.write_text(_PLAYBACK, encoding="utf-8")
    (tmp_path / "job.json").write_text(json.dumps(job), encoding="utf-8")
    # Freetown keeps UTC all year, so the device clock reads the same here
    # on any machine
    env = dict(os.environ, TZ="Africa/Freetown")
    subprocess.run([node, str(script), str(JS), str(tmp_path / "job.json"), str(tmp_path)],
                   check=True, timeout=300, env=env)
    result = json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))
    result["dir"] = tmp_path
    return result


@pytest.fixture(scope="module")
def playback(tmp_path_factory):
    return _playback(tmp_path_factory.mktemp("drill_copilot"))


def _minutes(hm: str) -> int:
    h, m = hm.split(":")
    return int(h) * 60 + int(m)


def test_the_drilling_log_reads_back_through_the_python_parser(playback):
    """Every interval comes back as it was logged: its depths, a description
    that reads as the class picked, its times, its rate and its bit; the two
    strikes; and the header the co-pilot wrote."""
    log = read_drilling_workbook(playback["dir"] / "drilling.xlsx")
    assert not [f for f in log.flags if f.level in ("error", "warning")], log.flags
    assert [(iv.top_m, iv.bottom_m) for iv in log.intervals] == [
        (float(a), float(b)) for a, b, *_ in TIMBO]
    for k, (iv, (_, _, key, note, start, end)) in enumerate(
            zip(log.intervals, TIMBO, strict=True)):
        assert lithology_class(iv.description).key == key
        # the last interval's note, as it was corrected after the countersign
        assert iv.description.endswith(", " + note + (", hard" if k == 13 else ""))
        assert (iv.from_time, iv.to_time) == (start, end)
        # the rate column is headed in min/m, which the reader turns over
        per_m = round((_minutes(end) - _minutes(start)) / 5, 2)
        assert iv.penetration_rate_m_per_min == pytest.approx(1 / per_m)
        assert iv.bit_diameter_in == 6.5
    assert log.water_strikes_m == [12.0, 30.0]
    assert log.borehole_ref == "BH-1"
    assert log.total_depth_m == 70
    assert log.drilling_method == "DTH hammer"
    assert log.start_date == DAY_ONE
    site = log.site
    assert site.community == "Dr. Timbo's Residence"
    assert site.client == "Dr. Timbo"
    assert site.district == "Western Area Rural"
    assert site.contractor == "WiNGiN Heavy Duty Machines Co. Ltd"
    utm = geographic_to_utm(8.4657, -13.2317)
    assert site.utm_zone == utm.zone == 28
    assert site.easting == pytest.approx(utm.easting, abs=0.05)
    assert site.northing == pytest.approx(utm.northing, abs=0.05)


def test_the_rate_is_the_clock_s_and_the_airlift_the_engine_s(playback):
    """Minutes per metre from the device clock at each end of an interval,
    and each strike's yield as groundwater.field_kit.airlift_yield gives it,
    written beside the strike."""
    for rate, (_, _, _, _, start, end) in zip(playback["rates"], TIMBO, strict=True):
        assert rate == pytest.approx((_minutes(end) - _minutes(start)) / 5)
    wb = load_workbook(playback["dir"] / "drilling.xlsx", read_only=True)
    rows = list(wb["Drilling Log"].iter_rows(values_only=True))
    header = rows[10]
    assert header[3] == "Penetration rate (min/m)"
    assert header[7:9] == ("Airlift yield (L/s)", "Airlift basis")
    for index, reading in STRIKES:
        options = {k: v for k, v in reading.items() if k not in ("depth_m", "method")}
        expected = airlift_yield(reading["method"], **options)
        row = rows[11 + index]
        assert row[6] == reading["depth_m"]
        assert row[7] == pytest.approx(round(expected["q_l_per_s"], 3))
        assert row[8] == expected["basis"]
        stored = playback["session"]["intervals"][index]["strike"]
        assert stored["q_l_per_s"] == pytest.approx(expected["q_l_per_s"], rel=1e-12)
    wb.close()


def test_only_a_class_from_the_table_is_logged(playback):
    """A class typed in as words is refused, and so is a note that would make
    the row read as another class, or name a water strike for the reader to
    take a depth from."""
    errors = playback["errors"]
    assert errors["freeText"] == "Choose the formation from the list."
    assert "would read as Saprolite, not Clay" in errors["slip"]
    assert "Enter the water strike in its own box" in errors["strikeInNote"]
    assert "within the interval" in errors["outside"]


def test_each_day_is_countersigned_and_a_later_change_shows(playback):
    """Both days signed; the change to day two after its countersign needed a
    reason, was kept, and leaves the day amended until it is signed again."""
    assert playback["statuses"]["signed"] == ["signed", "signed"]
    assert "Give the reason for the change" in playback["errors"]["unreasoned"]
    assert playback["statuses"]["amended"] == ["signed", "amended"]
    session = playback["session"]
    (amendment,) = session["amendments"]
    assert amendment["day"] == DAY_TWO and amendment["reason"] == "hardness left off"
    assert amendment["interval"] == "65-70"
    signs = session["signatures"]
    assert [s["name"] for s in signs[DAY_ONE]] == ["M. Kamara"]
    assert signs[DAY_ONE][0]["metres"] == 30 and signs[DAY_TWO][0]["metres"] == 40
    assert signs[DAY_ONE][0]["digest"] == playback["digests"][0]
    assert signs[DAY_TWO][0]["digest"] != playback["digests"][1]


def _labels(ws) -> dict:
    """Every text cell of a sheet, by its coordinate."""
    return {cell.coordinate: cell.value for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, str)}


def test_the_daily_report_is_the_template_a_day(playback, tmp_path):
    """No reader takes the daily report in either app, so each day's sheet
    is held to the template's own layout for that many rows: every label
    the template writes is where it writes it, and the day's figures sit
    beside them."""
    wb = load_workbook(playback["dir"] / "daily.xlsx")
    assert wb.sheetnames == [f"Daily {DAY_ONE}", f"Daily {DAY_TWO}"]
    for name, day, count, metres, cumulative in (
            (f"Daily {DAY_ONE}", DAY_ONE, 6, 30, 30),
            (f"Daily {DAY_TWO}", DAY_TWO, 8, 40, 70)):
        template = tmp_path / f"template_{count}.xlsx"
        write_daily_log_template(template, n_rows=count)
        expected = _labels(load_workbook(template).active)
        ws = wb[name]
        got = _labels(ws)
        for coordinate, label in expected.items():
            assert got.get(coordinate) == label, (name, coordinate, label)
        assert ws["B3"].value == day
        assert ws["E4"].value == "M. Kamara"
        assert ws["E5"].value == "A. Sesay"
        totals = 9 + count
        assert ws[f"B{totals}"].value == metres
        assert ws[f"D{totals}"].value == cumulative
        signature = ws[f"E{totals + 3}"].value
        assert signature.startswith("M. Kamara, countersigned on this device at " + day)
        assert ("amended after this countersign" in signature) == (day == DAY_TWO)
        assert ws["E8"].value.startswith(
            {"topsoil": "Topsoil", "fracture": "Fracture zone"}[
                TIMBO[0 if day == DAY_ONE else 6][2]])
    day_one = wb[f"Daily {DAY_ONE}"]
    # the strike at 12 m, with its airlift yield, on the 10-15 m row
    assert day_one["F10"].value == 12
    assert day_one["G10"].value == pytest.approx(round(20 / ((41 + 39.5 + 40.2) / 3), 3))
