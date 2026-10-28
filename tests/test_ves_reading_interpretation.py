"""The VES sheet readers and the interpretation, after an audit of both.

Each test is a defect the audit reproduced, held in Python, and the last
holds the browser's engine to the same answers in a Node vm (parity.mjs holds
the same cases in a browser). A resistive layer over conductive ground is a
cover, not basement; one resistance row K cannot convert no longer loses the
workbook; a column in ohm-metres is a resistivity whatever its name; an
IPI2Win ERR in the next cell and a table headed with the Greek rho are read; a
spacing read twice with no MN change is neither double-counted nor said to be
merged; a zero at an MN overlap no longer drops its partner; a 1 m zone is
never printed as no zone; and a layer above a short sounding's depth of
investigation is not said to lie below it.
"""

from __future__ import annotations

import dataclasses
import json
import math
import shutil
import subprocess
from pathlib import Path

import numpy as np
import pytest

from groundwater.config import VESConfig
from groundwater.ingestion.ves import _sounding_or_reason, read_ves_workbook
from groundwater.models import LayeredModel, SiteMetadata, VESSounding
from groundwater.siting import assess_siting
from groundwater.text import phrase
from groundwater.ves.interpret import (
    basement_top,
    conductive_below,
    drilling_preference_table,
    interpret_model,
)
from groundwater.ves.inversion import inversion_readings
from groundwater.ves.ipi2win import _model_from_grid, read_ipi2win_models
from groundwater.ves.model_range import _summary, reading_errors

REPO = Path(__file__).resolve().parents[1]
JS = REPO / "docs" / "js"

SPACINGS = [1, 1.5, 2, 3, 4, 6, 8, 10, 15, 20, 25, 30, 40, 50, 60, 80, 100]
COVER = phrase("ves.unit_resistive_cover")


def _sounding(ab2=SPACINGS, rho=None, mn=None, array="schlumberger") -> VESSounding:
    ab2 = np.array(ab2, dtype=float)
    return VESSounding(
        site=SiteMetadata(), sounding_id="VES 1", ab2=ab2,
        mn=np.full(len(ab2), np.nan) if mn is None else np.array(mn, dtype=float),
        rho_app=np.full(len(ab2), 100.0) if rho is None else np.array(rho, dtype=float),
        array_type=array,
    )


def _interp(rho, h, ab2=SPACINGS, err=2.0, config=None):
    model = LayeredModel(np.array(rho, dtype=float), np.array(h, dtype=float),
                         fit_error_percent=err, sounding_id="VES 1")
    return interpret_model(_sounding(ab2), model, config)


# ------------------------------------------------- a resistive cover is no basement

def test_a_resistive_layer_over_a_water_zone_is_a_cover_not_basement():
    # 5,000 ohm-m 0.8 m down, then 20 m of 60 ohm-m on 8,000 ohm-m basement
    cover = _interp([300, 5000, 60, 8000], [0.8, 2.0, 20.0])
    compact = _interp([300, 2900, 60, 8000], [0.8, 2.0, 20.0])
    assert cover.layers[1].unit == COVER
    assert cover.layers[3].unit == "fresh basement"
    # bedrock is below the water zone, not 0.8 m down above it
    assert cover.depth_to_basement_m == pytest.approx(22.8)
    assert cover.water_zones == [(3, 23)]
    assert "The depth to bedrock is estimated at about 22.8 m." in cover.narrative
    assert ("Layer 2 is as resistive as fresh basement, but conductive ground lies below it "
            "in the fitted model, so it is read as resistive cover and not as bedrock."
            in cover.narrative)
    # the suitability no longer swings on whether the cover is 2,900 or 5,000 ohm-m
    a, b = assess_siting([cover])[0], assess_siting([compact])[0]
    assert a.suitability == b.suitability == 82.7
    assert "overburden of about 1 m" not in a.rationale
    assert cover.max_drilling_depth_m == compact.max_drilling_depth_m == 35.0


@pytest.mark.parametrize(("rho", "h", "basement", "covers"), [
    # clay under a resistive layer is conductive ground too
    ([300, 5000, 10, 8000], [0.8, 2.0, 20.0], 22.8, [2]),
    # nothing but a conductive half-space below: no basement at all
    ([1100, 5000, 47], [1.0, 7.0], None, [2]),
    # two covers, basement under the second water zone
    ([300, 5000, 60, 4000, 60, 8000], [0.8, 2, 5, 2, 10], 19.8, [2, 4]),
    # a resistive layer over compact rock is still basement, as before
    ([300, 5000, 1500, 8000], [0.8, 2.0, 20.0], 0.8, []),
    # a resistive half-space under a water zone, as before
    ([300, 60, 5000], [2.0, 15.0], 17.0, []),
    # no layer reaches the edge: the half-space of 800 ohm-m or more
    ([300, 2900, 60, 2000], [0.8, 2.0, 20.0], 22.8, []),
])
def test_basement_is_resistive_ground_with_nothing_conductive_below(rho, h, basement, covers):
    interp = _interp(rho, h)
    assert interp.depth_to_basement_m == (pytest.approx(basement) if basement else None)
    assert [layer.number for layer in interp.layers if layer.unit == COVER] == covers
    tops = [layer.top_m for layer in interp.layers]
    assert basement_top(rho, tops, VESConfig()) == interp.depth_to_basement_m
    if len(covers) == 2:
        assert "Layers 2 and 4 are as resistive as fresh basement" in interp.narrative


def test_conductive_ground_is_what_the_rules_read_as_weathered():
    cfg = VESConfig()
    # water-bearing and clay-rich layers are conductive; compact rock and
    # basement are not
    assert conductive_below([300, 5000, 60, 10, 900, 3000], cfg) == [
        True, True, True, False, False, False]
    # under a fitted basement edge below 800 ohm-m, the edge decides
    low = dataclasses.replace(cfg, fresh_basement_min_rho=579.7)
    assert conductive_below([300, 1637, 700], low) == [False, False, False]
    assert conductive_below([300, 1637, 500], low) == [True, True, False]


def test_rokel_a_under_a_fitted_basement_edge_has_a_cover_and_no_basement():
    # Rokel A's fit read with the 579.7 ohm-m edge a calibration set fitted:
    # its 1,637 ohm-m layer lies over the 47 ohm-m water zone it is drilled for
    edge = dataclasses.replace(VESConfig(), fresh_basement_min_rho=579.7)
    interp = _interp([1104.9, 1637.1, 47.278], [1.018, 7.085], ab2=[1, 2, 5, 10, 20, 40, 80],
                     err=13.3, config=edge)
    assert interp.layers[1].unit == COVER
    assert interp.depth_to_basement_m is None
    assert interp.water_zones == [(8, 40)]


def test_the_range_reads_basement_by_the_same_rule():
    ves = VESConfig()
    rho, h = [300.0, 5000.0, 60.0, 8000.0], [0.8, 2.0, 20.0]
    theta = [math.log(v) for v in rho] + [math.log(v) for v in h]
    basement, weathered, deepest, _interfaces, _rho = _summary(theta, 4, 50.0, ves)
    assert basement == pytest.approx(22.8)
    assert weathered == pytest.approx(19.8)
    # over a conductive half-space: no basement in the model
    theta = [math.log(v) for v in (300.0, 5000.0, 60.0)] + [math.log(1.0), math.log(3.0)]
    assert _summary(theta, 3, 50.0, ves)[0] is None


# ------------------------------------------------ the zones and their sentences

def test_a_one_metre_zone_on_half_metres_keeps_its_metre():
    for h, zone in (([3.5, 1.0], (4, 5)), ([5.5, 1.0], (6, 7))):
        interp = _interp([500, 100, 5000], h)
        assert interp.water_zones == [zone]
        assert interp.aquifer_thickness_m == 1.0
        assert interp.score > 0
        assert f"The possible water zones are {zone[0]} m to {zone[1]} m." in interp.narrative
        assert drilling_preference_table([interp])[0]["Possible Water Zones (m)"] == (
            f"{zone[0]}-{zone[1]}")
    # a zone whose ends round apart is as it was
    assert _interp([500, 100, 40, 5000], [5.6, 1.4, 1.2]).water_zones == [(6, 8)]


def test_a_layer_above_a_short_soundings_reach_is_not_said_to_lie_below_it():
    interp = _interp([400, 90, 3000], [0.6, 8.0], ab2=[1, 1.5, 2, 3, 4, 5])
    assert interp.investigation_depth_m == 2.5
    assert interp.water_zones == []
    assert "lies below the 2.5 m" not in interp.narrative
    assert ("The water-bearing layer from 0.6 m starts within the top 3 m, which are taken "
            "as unsaturated, and the sounding resolves only 2.5 m, so it is not counted as "
            "a water zone." in interp.narrative)
    # a layer that does start below the depth of investigation keeps its sentence
    deep = _interp([1000, 1500, 100, 5000], [10, 40, 20], ab2=[1, 2, 5, 10, 20, 40, 60])
    assert ("The water-bearing layer from 50 m lies below the 30 m the sounding resolves, "
            "so it is not counted as a water zone." in deep.narrative)


# ---------------------------------------------------------------- the readers

def _grid(header, rows, number="VES 1"):
    return [["Sounding No.", number], ["Array", "Schlumberger"],
            ["No."] + header] + [[k + 1] + list(r) for k, r in enumerate(rows)]


_BAD_ROWS = [(0.3, 1.0, 12.5), (0.5, 1.0, 9.0), (2, 1.0, 4.1), (3, 1.0, 1.9),
             (5, 1.0, 0.62), (10, 1.0, 0.18), (15, 1.0, 0.09)]


def test_a_resistance_row_k_cannot_convert_is_left_out_and_flagged():
    sounding, reason = _sounding_or_reason(_grid(["AB/2 (m)", "MN (m)", "R (ohm)"], _BAD_ROWS),
                                           "x.xlsx", "VES 1")
    assert reason == "" and sounding is not None
    assert sounding.ab2.tolist() == [2, 3, 5, 10, 15]
    flag = next(f for f in sounding.flags if f.code == "k_undefined")
    assert flag.level == "warning"
    assert flag.message == (
        "The Schlumberger geometric factor K needs MN/2 between 0 and AB/2, so the "
        "resistances read at AB/2 0.3 and 0.5 m could not be turned into resistivities and "
        "are left out. Check the spacings on the sheet.")


def test_one_bad_resistance_row_no_longer_loses_the_workbook(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    sheets = [(book.active, "R1", _BAD_ROWS[:1] + _BAD_ROWS[2:]),
              (book.create_sheet(), "R2", _BAD_ROWS[2:])]
    for ws, number, rows in sheets:
        ws.title = number
        for row in _grid(["AB/2 (m)", "MN (m)", "R (ohm)"], rows, number):
            ws.append(row)
    path = tmp_path / "resistance.xlsx"
    book.save(path)
    soundings = read_ves_workbook(path)
    assert [s.sounding_id for s in soundings] == ["R1", "R2"]
    assert [len(s.ab2) for s in soundings] == [5, 5]
    assert [f.code for f in soundings[0].flags].count("k_undefined") == 1


@pytest.mark.parametrize("header", [
    "R (ohm-m)", "R (Ωm)", "R (ohm.m)", "R (Ω·m)", "R (ohmm)", "App. resistance (ohm-m)",
    "Apparent Resistance (Ωm)", "ρa = K·V/I (ohm-m)", "Resistance (ohm metres)",
])
def test_a_column_in_ohm_metres_is_a_resistivity_whatever_its_name(header):
    rows = [(1, 0.5, 120.0), (2, 0.5, 95.0), (3, 0.5, 80.0), (5, 0.5, 70.0), (10, 2, 90.0)]
    sounding, _ = _sounding_or_reason(_grid(["AB/2 (m)", "MN (m)", header], rows), "x", "V")
    assert sounding.rho_app.tolist() == [120.0, 95.0, 80.0, 70.0, 90.0]
    assert "rho_computed_from_resistance" not in [f.code for f in sounding.flags]


@pytest.mark.parametrize("header", ["R (ohm)", "V/I (ohm)", "Resistance (ohms)", "R (Ω)"])
def test_a_column_in_ohms_is_still_a_resistance(header):
    rows = [(1, 0.5, 120.0), (2, 0.5, 95.0), (3, 0.5, 80.0)]
    sounding, _ = _sounding_or_reason(_grid(["AB/2 (m)", "MN (m)", header], rows), "x", "V")
    assert "rho_computed_from_resistance" in [f.code for f in sounding.flags]
    assert sounding.rho_app[0] == pytest.approx(120.0 * math.pi * (1 - 0.0625) / 0.5)


_IPI_ROWS = [[1, 832.14, 1, "0/0"], [2, 2102.8, 7.37, 1], [3, 36.71, None, 8.37]]


@pytest.mark.parametrize("err_rows", [
    [["ERR (%):", 21.5]],          # the label, the value in the next cell
    [["ERR", "21.5"]],
    [["ERR (%)", None, 21.5]],     # past an empty cell
    [["ERR"], [21.5]],             # the value below
    [["ERR = 21.5"]],              # one cell, as before
])
def test_an_ipi2win_err_beside_or_below_its_label_is_read(err_rows):
    grid = [["Sounding Number", 1]] + err_rows + [["N", "rho", "h", "z"]] + _IPI_ROWS
    model = _model_from_grid(grid, "1")
    assert model.fit_error_percent == 21.5
    assert interpret_model(None, model).fit_quality == "unreliable"


@pytest.mark.parametrize("header", [["N", "ρ", "h", "z"], ["N", "ρ, Ohm.m", "h, m", "z, m"],
                                    ["N", "ρa (ohm.m)", "h (m)", "z (m)"]])
def test_an_ipi2win_table_headed_with_the_greek_rho_is_read(header):
    model = _model_from_grid([header] + _IPI_ROWS, "1")
    assert model.resistivities.tolist() == [832.14, 2102.8, 36.71]
    assert model.thicknesses.tolist() == [1.0, 7.37]


def test_an_ipi2win_sheet_with_no_model_table_is_reported(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "A (1)"
    for row in [["Sounding Number", "A (1)"], ["ERR (%)", 21.5], ["N", "ρ", "h", "z"]] + _IPI_ROWS:
        book.active.append(row)
    notes = book.create_sheet("Notes")
    notes.append(["Models transcribed from the survey report"])
    path = tmp_path / "models.xlsx"
    book.save(path)
    skipped = []
    models = read_ipi2win_models(path, skipped=skipped)
    assert list(models) == ["A (1)"] and models["A (1)"].fit_error_percent == 21.5
    assert [(f.level, f.code) for f in skipped] == [("warning", "sheet_skipped")]
    assert skipped[0].message.startswith("Sheet 'Notes' was skipped: no model table found")
    # without the list the reading is as it was
    assert list(read_ipi2win_models(path)) == ["A (1)"]
    # a CSV with no table is named as a file
    notes_csv = tmp_path / "notes.csv"
    notes_csv.write_text("Models transcribed from the survey report\n", encoding="utf-8")
    skipped = []
    assert read_ipi2win_models(notes_csv, skipped=skipped) == {}
    assert skipped[0].message.startswith("File 'notes.csv' was skipped: no model table found")


# ------------------------------------------------------- repeats and overlaps

_REPEAT_AB2 = [1, 2, 3, 5, 10, 10, 15, 20, 30, 40, 40, 60]
_REPEAT_RHO = [120, 95, 80, 70, 90, 120, 110, 140, 190, 230, 260, 300]


def test_a_repeat_with_no_mn_change_is_neither_widened_nor_listed_twice():
    sounding = _sounding(_REPEAT_AB2, _REPEAT_RHO)
    ab2 = inversion_readings(sounding)[0]
    assert ab2.tolist().count(10.0) == 2  # not merged: both readings are fitted
    sigma, spacings = reading_errors(sounding, ab2, 3.0)
    assert spacings == []
    assert sigma.tolist() == [0.03] * len(ab2)
    # with an MN change at those spacings the splice merges each pair, and
    # each widens its one merged reading, listed once
    mn = [1.0] * 5 + [4.0] * 5 + [10.0] * 2
    merged = _sounding(_REPEAT_AB2, _REPEAT_RHO, mn)
    ab2 = inversion_readings(merged)[0]
    sigma, spacings = reading_errors(merged, ab2, 3.0)
    assert spacings == [10.0, 40.0]
    wide = sigma[ab2.tolist().index(10.0)]
    assert wide == pytest.approx(math.hypot(0.03, 0.5 * math.log(120 / 90)))


def test_the_reader_says_merged_only_where_the_splice_merges():
    data = [(1, 120), (2, 95), (3, 80), (5, 70), (10, 90), (10, 120), (15, 110), (20, 140)]
    plain, _ = _sounding_or_reason(_grid(["AB/2 (m)", "Rho (ohm-m)"], data), "x", "V")
    messages = {f.code: f.message for f in plain.flags}
    assert messages["segment_overlap"] == (
        "1 AB/2 value(s) read more than once; the readings are kept, and the inversion fits "
        "each one as it stands rather than merging them.")
    assert "merges the pair" not in messages["segment_overlap_discrepancy"]
    assert "fits each reading as it stands" in messages["segment_overlap_discrepancy"]
    # a Wenner sheet is never spliced
    wenner_grid = [["Sounding No.", "W"], ["Array", "Wenner"], ["No.", "a (m)", "Rho (ohm-m)"]]
    wenner, _ = _sounding_or_reason(wenner_grid + [[k + 1, a, r] for k, (a, r) in enumerate(data)],
                                    "x", "W")
    assert "merges the pair" not in {f.code: f.message for f in wenner.flags}[
        "segment_overlap_discrepancy"]
    # an MN change: the pair is merged, and the flags say so as before
    with_mn = [(a, 1.0 if k < 5 else 4.0, r) for k, (a, r) in enumerate(data)]
    spliced, _ = _sounding_or_reason(_grid(["AB/2 (m)", "MN (m)", "Rho (ohm-m)"], with_mn), "x", "V")
    messages = {f.code: f.message for f in spliced.flags}
    assert messages["segment_overlap"] == (
        "1 AB/2 value(s) repeated with different MN (segment changes); both readings kept.")
    assert "the inversion merges the pair by geometric mean" in messages[
        "segment_overlap_discrepancy"]


def test_a_zero_at_an_mn_overlap_keeps_the_good_reading_beside_it():
    ab2 = [1, 2, 3, 5, 8, 10, 10, 15, 20, 30, 40, 40, 60, 80, 100]
    mn = [0.5] * 6 + [4.0] * 5 + [10.0] * 4
    rho = [300, 250, 180, 120, 90, 85, 88, 80, 90, 120, 170, 0, 260, 330, 400]
    fitted, fitted_rho, _ = inversion_readings(_sounding(ab2, rho, mn))
    assert 40.0 in fitted.tolist()
    assert fitted_rho[fitted.tolist().index(40.0)] == pytest.approx(170.0, rel=1e-14)
    # no usable reading at a spacing: that spacing alone is left out
    rho[10] = 0
    fitted, _, _ = inversion_readings(_sounding(ab2, rho, mn))
    assert 40.0 not in fitted.tolist() and len(fitted) == 12


# ------------------------------------------------- the browser's engine agrees

_DRIVER = r"""
import { readFileSync } from 'node:fs';
import vm from 'node:vm';
const sandbox = vm.createContext({ console, structuredClone, TextEncoder, TextDecoder });
for (const f of ['gwt-data.js', 'gwt-geo.js', 'gwt-core.js']) {
  vm.runInContext(readFileSync(process.argv[2] + '/' + f, 'utf8'), sandbox, { filename: f });
}
const C = sandbox.GWT.core;
const cases = JSON.parse(readFileSync(0, 'utf8'));
const sounding = (c) => ({ sounding_id: 'VES 1', array_type: 'schlumberger', ab2: c.ab2,
  mn: c.mn.map((v) => (v === null ? NaN : v)), rho_app: c.rho, flags: [] });
process.stdout.write(JSON.stringify({
  interps: cases.interps.map((c) => {
    const cfg = C.withConfig({ ves: c.ves });
    const s = sounding({ ab2: c.ab2, mn: c.ab2.map(() => null), rho: c.ab2.map(() => 100) });
    const i = C.interpretModel(s, C.layeredModel(c.rho, c.h, { fit_error_percent: c.err,
      sounding_id: 'VES 1' }), cfg);
    return { units: i.layers.map((l) => l.unit), zones: i.water_zones,
      basement: i.depth_to_basement_m, drill: i.max_drilling_depth_m,
      narrative: i.narrative, suitability: C.assessSiting([i], cfg.ves)[0].suitability };
  }),
  sheets: cases.sheets.map((g) => {
    const s = C.readVesSheets([{ name: 'V', rows: g }], 'x.xlsx')[0];
    return { ab2: s.ab2, rho: s.rho_app, flags: s.flags.map((f) => [f.level, f.code, f.message]) };
  }),
  readings: cases.readings.map((c) => {
    const s = sounding(c);
    const fitted = C.inversionReadings(s);
    const e = C.readingErrors(s, fitted.ab2, 3.0);
    return { ab2: fitted.ab2, rho: fitted.rho, sigma: e.sigma, spacings: e.spacings };
  }),
}));
"""

_CASES = {
    "interps": [
        {"rho": [300, 5000, 60, 8000], "h": [0.8, 2.0, 20.0]},
        {"rho": [300, 5000, 10, 8000], "h": [0.8, 2.0, 20.0]},
        {"rho": [1100, 5000, 47], "h": [1.0, 7.0]},
        {"rho": [300, 5000, 60, 4000, 60, 8000], "h": [0.8, 2, 5, 2, 10]},
        {"rho": [300, 5000, 1500, 8000], "h": [0.8, 2.0, 20.0]},
        {"rho": [1104.9, 1637.1, 47.278], "h": [1.018, 7.085], "ab2": [1, 2, 5, 10, 20, 40, 80],
         "err": 13.3, "ves": {"fresh_basement_min_rho": 579.7}},
        {"rho": [500, 100, 5000], "h": [3.5, 1.0]},
        {"rho": [500, 100, 5000], "h": [5.5, 1.0]},
        {"rho": [400, 90, 3000], "h": [0.6, 8.0], "ab2": [1, 1.5, 2, 3, 4, 5]},
        {"rho": [400, 90, 30, 3000], "h": [0.6, 0.8, 8.0], "ab2": [1, 1.5, 2, 3, 4, 5]},
    ],
    "sheets": [
        _grid(["AB/2 (m)", "MN (m)", "R (ohm)"], _BAD_ROWS),
        _grid(["AB/2 (m)", "MN (m)", "Apparent Resistance (Ωm)"],
              [(1, 0.5, 120.0), (2, 0.5, 95.0), (3, 0.5, 80.0), (5, 0.5, 70.0), (10, 2, 90.0)]),
        _grid(["AB/2 (m)", "Rho (ohm-m)"],
              [(1, 120), (2, 95), (3, 80), (5, 70), (10, 90), (10, 120), (15, 110), (20, 140)]),
        _grid(["AB/2 (m)", "MN (m)", "Rho (ohm-m)"],
              [(1, 1.0, 120), (2, 1.0, 95), (3, 1.0, 80), (5, 1.0, 70), (10, 1.0, 90),
               (10, 4.0, 120), (15, 4.0, 110), (20, 4.0, 140)]),
    ],
    "readings": [
        {"ab2": [1, 2, 3, 5, 8, 10, 10, 15, 20, 30, 40, 40, 60, 80, 100],
         "mn": [0.5] * 6 + [4.0] * 5 + [10.0] * 4,
         "rho": [300, 250, 180, 120, 90, 85, 88, 80, 90, 120, 170, 0, 260, 330, 400]},
        {"ab2": _REPEAT_AB2, "mn": [None] * 12, "rho": _REPEAT_RHO},
        {"ab2": _REPEAT_AB2, "mn": [1.0] * 5 + [4.0] * 5 + [10.0] * 2, "rho": _REPEAT_RHO},
    ],
}
for _case in _CASES["interps"]:
    _case.setdefault("ab2", SPACINGS)
    _case.setdefault("err", 2.0)
    _case.setdefault("ves", {})


def _python_answers() -> dict:
    out = {"interps": [], "sheets": [], "readings": []}
    for c in _CASES["interps"]:
        config = dataclasses.replace(VESConfig(), **c["ves"])
        interp = _interp(c["rho"], c["h"], c["ab2"], c["err"], config)
        out["interps"].append({
            "units": [layer.unit for layer in interp.layers],
            "zones": [list(z) for z in interp.water_zones],
            "basement": interp.depth_to_basement_m, "drill": interp.max_drilling_depth_m,
            "narrative": interp.narrative,
            "suitability": assess_siting([interp], config)[0].suitability,
        })
    for grid in _CASES["sheets"]:
        s, _ = _sounding_or_reason(grid, "x.xlsx", "V")
        out["sheets"].append({"ab2": s.ab2.tolist(), "rho": s.rho_app.tolist(),
                              "flags": [[f.level, f.code, f.message] for f in s.flags]})
    for c in _CASES["readings"]:
        mn = [np.nan if v is None else v for v in c["mn"]]
        s = _sounding(c["ab2"], c["rho"], mn)
        ab2, rho, _ = inversion_readings(s)
        sigma, spacings = reading_errors(s, ab2, 3.0)
        out["readings"].append({"ab2": ab2.tolist(), "rho": rho.tolist(),
                                "sigma": sigma.tolist(), "spacings": spacings})
    return out


def _part(js, py, path=""):
    """The first place two answers part: numbers to 1e-12 relative,
    everything else exactly."""
    if isinstance(py, (int, float)) and isinstance(js, (int, float)) and not isinstance(py, bool):
        return None if math.isclose(js, py, rel_tol=1e-12, abs_tol=1e-12) else f"{path}: {js} vs {py}"
    if isinstance(py, dict) and isinstance(js, dict):
        if set(py) != set(js):
            return f"{path}: keys {sorted(js)} vs {sorted(py)}"
        return next((d for d in (_part(js[k], py[k], f"{path}.{k}") for k in py) if d), None)
    if isinstance(py, list) and isinstance(js, list):
        if len(py) != len(js):
            return f"{path}: length {len(js)} vs {len(py)}"
        return next((d for d in (_part(a, b, f"{path}[{i}]")
                                 for i, (a, b) in enumerate(zip(js, py, strict=True))) if d), None)
    return None if js == py else f"{path}: js {js!r} vs py {py!r}"


def test_the_browser_engine_reads_and_interprets_alike(tmp_path):
    node = shutil.which("node")
    if node is None:
        pytest.skip("node is not installed; run `node --version` to check")
    driver = tmp_path / "ves_reading.mjs"
    driver.write_text(_DRIVER, encoding="utf-8")
    done = subprocess.run([node, str(driver), str(JS)], input=json.dumps(_CASES),
                          capture_output=True, text=True, encoding="utf-8", check=True)
    found = _part(json.loads(done.stdout), _python_answers())
    assert found is None, found
